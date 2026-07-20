from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from dc_engines.office_translation import (
    TranslationError,
    build_translation_options,
    choose_translation_model,
    translate_office_document,
)
from docx import Document
from openai.types.chat.chat_completion import ChatCompletion, Choice
from openai.types.chat.chat_completion_message import ChatCompletionMessage
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas

from astrbot.api.provider import LLMResponse
from astrbot.core.message.components import File, Plain
from astrbot.core.provider.sources.openai_source import ProviderOpenAIOfficial
from dc_engines import office_translation_runtime


class _TranslationProvider:
    def __init__(self) -> None:
        self.calls: list[dict[str, object]] = []

    async def text_chat(self, **kwargs: object) -> LLMResponse:
        self.calls.append(kwargs)
        prompt = str(kwargs["prompt"])
        translated = (
            prompt.replace("合同标题", "Contract title")
            .replace("付款条款", "Payment terms")
            .replace("服务费", "Service fee")
            .replace("总金额", "Total amount")
            .replace("你好，世界", "Hello, world")
            .replace("Hello world", "你好，世界")
        )
        return LLMResponse(role="assistant", completion_text=translated)


class _AmbiguousPairProvider(_TranslationProvider):
    async def text_chat(self, **kwargs: object) -> LLMResponse:
        if kwargs.get("model") == "qwen3.6-flash":
            self.calls.append(kwargs)
            return LLMResponse(
                role="assistant", completion_text='{"language":"Spanish"}'
            )
        return await super().text_chat(**kwargs)


class _RuntimeEvent:
    def __init__(self, source_path: Path) -> None:
        self.extras: dict[str, object] = {
            "assistant_workbench_workspace_token": "workspace-token"
        }
        self.message_obj = SimpleNamespace(
            message=[File(name=source_path.name, file=str(source_path))]
        )

    def get_platform_id(self) -> str:
        return "巅池-Agent小助手"

    def get_extra(self, key: str, default: object = None) -> object:
        return self.extras.get(key, default)

    def set_extra(self, key: str, value: object) -> None:
        self.extras[key] = value


def _task_data(**overrides: str) -> dict[str, str]:
    task_data = {
        "file_goal": "翻译成英文",
        "operation": "translate",
        "source_language": "auto",
        "target_language": "English",
        "direction_mode": "one_way",
        "translation_style": "faithful",
        "output_mode": "target_only",
        "professional_domain": "general",
        "preserve_layout": "preserve",
        "output_format": "docx",
        "model_choice": "auto",
    }
    task_data.update(overrides)
    return task_data


def test_translation_options_require_target_and_route_real_qwen_mt_models() -> None:
    with pytest.raises(TranslationError, match="目标语言"):
        build_translation_options(_task_data(target_language=""))

    fast = build_translation_options(_task_data(model_choice="fast"))
    professional = build_translation_options(
        _task_data(translation_style="professional", professional_domain="legal")
    )
    bidirectional = build_translation_options(
        _task_data(
            direction_mode="bidirectional",
            source_language="Chinese",
            target_language="English",
            output_mode="target_only",
        )
    )

    assert choose_translation_model(fast) == "qwen-mt-turbo"
    assert choose_translation_model(professional) == "qwen-mt-plus"
    assert choose_translation_model(bidirectional) == "qwen-mt-plus"
    assert bidirectional.output_mode == "bilingual"


@pytest.mark.asyncio
async def test_astrbot_openai_compatible_provider_keeps_qwen_translation_options() -> (
    None
):
    provider = ProviderOpenAIOfficial(
        provider_config={
            "id": "test-aihubmix",
            "type": "aihubmix_chat_completion",
            "model": "qwen-mt-turbo",
            "key": ["test-key"],
        },
        provider_settings={},
    )
    try:
        payload, _contexts = await provider._prepare_chat_payload(
            "付款条款",
            model="qwen-mt-turbo",
            translation_options={
                "source_lang": "Chinese",
                "target_lang": "English",
                "domains": "legal",
            },
        )
    finally:
        await provider.terminate()

    assert payload["translation_options"] == {
        "source_lang": "Chinese",
        "target_lang": "English",
        "domains": "legal",
    }


@pytest.mark.asyncio
async def test_astrbot_provider_forwards_qwen_options_in_openai_extra_body() -> None:
    provider = ProviderOpenAIOfficial(
        provider_config={
            "id": "test-aihubmix-extra-body",
            "type": "aihubmix_chat_completion",
            "model": "qwen-mt-plus",
            "key": ["test-key"],
        },
        provider_settings={},
    )
    create = AsyncMock(
        return_value=ChatCompletion(
            id="chatcmpl-translation",
            choices=[
                Choice(
                    finish_reason="stop",
                    index=0,
                    message=ChatCompletionMessage(
                        content="Payment terms",
                        role="assistant",
                    ),
                )
            ],
            created=0,
            model="qwen-mt-plus",
            object="chat.completion",
        )
    )
    provider.client.chat.completions.create = create
    try:
        response = await provider.text_chat(
            prompt="付款条款",
            model="qwen-mt-plus",
            translation_options={
                "source_lang": "Chinese",
                "target_lang": "English",
                "terms": [{"source": "付款条款", "target": "Payment Terms"}],
            },
            request_max_retries=1,
        )
    finally:
        await provider.terminate()

    assert response.completion_text == "Payment terms"
    assert create.await_args.kwargs["extra_body"]["translation_options"] == {
        "source_lang": "Chinese",
        "target_lang": "English",
        "terms": [{"source": "付款条款", "target": "Payment Terms"}],
    }


@pytest.mark.asyncio
async def test_astrbot_runtime_uses_aihubmix_provider_and_returns_real_file(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "contract.txt"
    source.write_text("付款条款", encoding="utf-8")
    provider = _TranslationProvider()
    provider.provider_config = {"type": "aihubmix_chat_completion"}
    context = SimpleNamespace(get_all_providers=lambda: [provider])
    event = _RuntimeEvent(source)
    monkeypatch.setattr(
        office_translation_runtime,
        "get_astrbot_data_path",
        lambda: str(tmp_path / "astrbot-data"),
    )

    result = await office_translation_runtime.run_translation_workbench(
        context,
        event,
        _task_data(output_format="txt"),
    )

    assert any(
        isinstance(component, Plain) and "翻译完成" in component.text
        for component in result.chain
    )
    files = [component for component in result.chain if isinstance(component, File)]
    assert len(files) == 1
    assert Path(files[0].file_).read_text(encoding="utf-8") == "Payment terms"
    assert provider.calls[0]["model"] == "qwen-mt-turbo"
    manifest_path = Path(str(event.extras["office_translation_manifest"]))
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["workspace_token"] == "workspace-token"
    assert event.extras["office_translation_artifacts"] == {"txt": str(files[0].file_)}


@pytest.mark.asyncio
async def test_translation_runtime_creates_real_native_feishu_document(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "contract.txt"
    source.write_text("付款条款", encoding="utf-8")
    provider = _TranslationProvider()
    provider.provider_config = {"type": "aihubmix_chat_completion"}
    create = AsyncMock(
        return_value=SimpleNamespace(
            success=lambda: True,
            data=SimpleNamespace(
                document=SimpleNamespace(document_id="doccn_translation")
            ),
        )
    )
    write = AsyncMock(return_value=SimpleNamespace(success=lambda: True))
    client = SimpleNamespace(
        docx=SimpleNamespace(
            v1=SimpleNamespace(
                document=SimpleNamespace(acreate=create),
                document_block_children=SimpleNamespace(acreate=write),
            )
        )
    )
    context = SimpleNamespace(
        get_all_providers=lambda: [provider],
        get_platform_inst=lambda _platform_id: SimpleNamespace(lark_api=client),
    )
    event = _RuntimeEvent(source)
    monkeypatch.setattr(
        office_translation_runtime,
        "get_astrbot_data_path",
        lambda: str(tmp_path / "astrbot-data"),
    )

    result = await office_translation_runtime.run_translation_workbench(
        context,
        event,
        _task_data(output_format="feishu_doc"),
    )

    assert create.await_count == 1
    assert write.await_count == 1
    assert any(
        isinstance(component, Plain)
        and "https://feishu.cn/docx/doccn_translation" in component.text
        for component in result.chain
    )
    assert event.extras["office_translation_artifacts"] == {
        "feishu_doc": "https://feishu.cn/docx/doccn_translation"
    }
    manifest = json.loads(
        Path(str(event.extras["office_translation_manifest"])).read_text(
            encoding="utf-8"
        )
    )
    assert manifest["artifacts"]["feishu_doc"].endswith("doccn_translation")


@pytest.mark.asyncio
async def test_translation_runtime_reads_selected_feishu_cloud_document(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    provider = _TranslationProvider()
    provider.provider_config = {"type": "aihubmix_chat_completion"}
    read = AsyncMock(
        return_value=SimpleNamespace(
            success=lambda: True,
            data=SimpleNamespace(content="付款条款"),
        )
    )
    client = SimpleNamespace(
        docx=SimpleNamespace(
            v1=SimpleNamespace(document=SimpleNamespace(araw_content=read))
        )
    )
    context = SimpleNamespace(
        get_all_providers=lambda: [provider],
        get_platform_inst=lambda _platform_id: SimpleNamespace(lark_api=client),
    )
    placeholder = tmp_path / "unused.txt"
    event = _RuntimeEvent(placeholder)
    event.message_obj.message = []
    monkeypatch.setattr(
        office_translation_runtime,
        "get_astrbot_data_path",
        lambda: str(tmp_path / "astrbot-data"),
    )
    task_data = _task_data(output_format="txt")
    task_data["file_source"] = (
        "已关联 1 项资料：合同（https://example.feishu.cn/docx/doccn_source）"
    )

    result = await office_translation_runtime.run_translation_workbench(
        context,
        event,
        task_data,
    )

    assert read.await_count == 1
    files = [component for component in result.chain if isinstance(component, File)]
    assert len(files) == 1
    assert Path(files[0].file_).read_text(encoding="utf-8") == "Payment terms"


@pytest.mark.asyncio
async def test_translation_runtime_resolves_explicit_file_url(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    downloaded = tmp_path / "downloaded.txt"
    downloaded.write_text("付款条款", encoding="utf-8")
    download = AsyncMock(return_value=downloaded)
    monkeypatch.setattr(office_translation_runtime, "_download_source", download)
    monkeypatch.setattr(
        office_translation_runtime,
        "get_astrbot_data_path",
        lambda: str(tmp_path / "astrbot-data"),
    )
    provider = _TranslationProvider()
    provider.provider_config = {"type": "aihubmix_chat_completion"}
    context = SimpleNamespace(get_all_providers=lambda: [provider])
    placeholder = tmp_path / "unused.txt"
    event = _RuntimeEvent(placeholder)
    event.message_obj.message = []
    task_data = _task_data(output_format="txt")
    task_data["file_source"] = "https://files.example.com/contracts/source.txt"

    result = await office_translation_runtime.run_translation_workbench(
        context,
        event,
        task_data,
    )

    download.assert_awaited_once()
    assert download.await_args.args[0] == task_data["file_source"]
    assert any(isinstance(component, File) for component in result.chain)


@pytest.mark.asyncio
async def test_docx_translation_preserves_headings_tables_and_writes_trace(
    tmp_path: Path,
) -> None:
    source = tmp_path / "contract.docx"
    document = Document()
    document.add_heading("合同标题", level=1)
    body = document.add_paragraph("付款条款")
    body.style = document.styles["Normal"]
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "服务费"
    table.cell(0, 1).text = "1000"
    document.save(source)

    provider = _TranslationProvider()
    result = await translate_office_document(
        provider,
        source,
        _task_data(output_format="docx"),
        output_dir=tmp_path / "result",
    )

    translated = Document(result.artifacts["docx"])
    assert translated.paragraphs[0].text == "Contract title"
    assert translated.paragraphs[0].style.name.startswith("Heading")
    assert translated.paragraphs[1].text == "Payment terms"
    assert translated.tables[0].cell(0, 0).text == "Service fee"
    assert translated.tables[0].cell(0, 1).text == "1000"
    manifest = json.loads(result.manifest_path.read_text(encoding="utf-8"))
    assert manifest["status"] == "completed"
    assert manifest["source"]["sha256"]
    assert manifest["translation"]["target_language"] == "English"
    assert manifest["translation"]["model"] == "aihubmix/qwen-mt-turbo"
    assert manifest["artifacts"]["docx"].endswith(".docx")
    assert provider.calls[0]["model"] == "qwen-mt-turbo"
    assert provider.calls[0]["translation_options"] == {
        "source_lang": "auto",
        "target_lang": "English",
    }


@pytest.mark.asyncio
async def test_xlsx_translation_keeps_sheets_formulas_numbers_and_basic_styles(
    tmp_path: Path,
) -> None:
    source = tmp_path / "quote.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "报价"
    sheet.freeze_panes = "B2"
    sheet["A1"] = "总金额"
    sheet["A1"].font = Font(bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="176E64")
    sheet["A2"] = "服务费"
    sheet["B2"] = 1000
    sheet["C2"] = "=B2*1.06"
    workbook.create_sheet("说明")["A1"] = "你好，世界"
    workbook.save(source)

    result = await translate_office_document(
        _TranslationProvider(),
        source,
        _task_data(output_format="xlsx", output_mode="bilingual"),
        output_dir=tmp_path / "result",
    )

    translated = load_workbook(result.artifacts["xlsx"], data_only=False)
    assert translated.sheetnames == ["报价", "说明"]
    assert translated["报价"]["A1"].value == "总金额\nTotal amount"
    assert translated["报价"]["A2"].value == "服务费\nService fee"
    assert translated["报价"]["B2"].value == 1000
    assert translated["报价"]["C2"].value == "=B2*1.06"
    assert translated["报价"].freeze_panes == "B2"
    assert translated["报价"]["A1"].font.bold is True
    assert translated["报价"]["A1"].fill.fgColor.rgb.endswith("176E64")
    assert translated["说明"]["A1"].value == "你好，世界\nHello, world"


@pytest.mark.asyncio
async def test_professional_translation_passes_domain_terms_and_tm_to_qwen_mt(
    tmp_path: Path,
) -> None:
    source = tmp_path / "legal.txt"
    source.write_text("付款条款", encoding="utf-8")
    provider = _TranslationProvider()

    await translate_office_document(
        provider,
        source,
        _task_data(
            output_format="txt",
            translation_style="professional",
            professional_domain="legal",
            glossary="付款条款=Payment Terms",
            translation_memory="本合同生效=This agreement takes effect",
        ),
        output_dir=tmp_path / "result",
    )

    call = provider.calls[0]
    assert call["model"] == "qwen-mt-plus"
    assert call["translation_options"] == {
        "source_lang": "auto",
        "target_lang": "English",
        "domains": "legal",
        "terms": [{"source": "付款条款", "target": "Payment Terms"}],
        "tm_list": [{"source": "本合同生效", "target": "This agreement takes effect"}],
    }


@pytest.mark.asyncio
async def test_large_unit_is_bounded_and_regeneration_links_previous_run(
    tmp_path: Path,
) -> None:
    source = tmp_path / "large.txt"
    source.write_text("合同标题。" * 2500, encoding="utf-8")
    first_provider = _TranslationProvider()
    first = await translate_office_document(
        first_provider,
        source,
        _task_data(output_format="txt"),
        output_dir=tmp_path / "runs" / "first",
        run_id="first-run",
    )
    second_provider = _TranslationProvider()
    second = await translate_office_document(
        second_provider,
        source,
        _task_data(output_format="txt"),
        output_dir=tmp_path / "runs" / "second",
        run_id="second-run",
    )

    assert len(first_provider.calls) > 1
    assert all(len(str(call["prompt"])) <= 8000 for call in first_provider.calls)
    assert first.artifacts["txt"].stat().st_size > 1000
    second_manifest = json.loads(second.manifest_path.read_text(encoding="utf-8"))
    assert second_manifest["regeneration_of"] == "first-run"


@pytest.mark.asyncio
async def test_ambiguous_bidirectional_pair_uses_structured_language_classifier(
    tmp_path: Path,
) -> None:
    source = tmp_path / "mixed.txt"
    source.write_text("Hola mundo", encoding="utf-8")
    provider = _AmbiguousPairProvider()

    await translate_office_document(
        provider,
        source,
        _task_data(
            output_format="txt",
            source_language="English",
            target_language="Spanish",
            direction_mode="bidirectional",
        ),
        output_dir=tmp_path / "result",
    )

    assert provider.calls[0]["model"] == "qwen3.6-flash"
    translation_call = provider.calls[1]
    assert translation_call["model"] == "qwen-mt-plus"
    assert translation_call["translation_options"] == {
        "source_lang": "Spanish",
        "target_lang": "English",
    }


@pytest.mark.asyncio
async def test_encrypted_and_scanned_pdf_return_recoverable_errors(
    tmp_path: Path,
) -> None:
    encrypted = tmp_path / "encrypted.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    writer.encrypt("secret")
    with encrypted.open("wb") as output:
        writer.write(output)

    with pytest.raises(TranslationError) as encrypted_error:
        await translate_office_document(
            _TranslationProvider(),
            encrypted,
            _task_data(output_format="pdf"),
            output_dir=tmp_path / "encrypted-result",
        )
    assert encrypted_error.value.code == "password_protected"
    assert "移除密码" in encrypted_error.value.recovery

    scanned = tmp_path / "scanned.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    with scanned.open("wb") as output:
        writer.write(output)

    with pytest.raises(TranslationError) as scan_error:
        await translate_office_document(
            _TranslationProvider(),
            scanned,
            _task_data(output_format="pdf"),
            output_dir=tmp_path / "scan-result",
        )
    assert scan_error.value.code == "scanned_pdf"
    assert "OCR" in scan_error.value.recovery


@pytest.mark.asyncio
async def test_searchable_pdf_translation_creates_real_pdf_delivery(
    tmp_path: Path,
) -> None:
    source = tmp_path / "searchable.pdf"
    pdf = canvas.Canvas(str(source))
    pdf.drawString(72, 720, "Hello world")
    pdf.save()

    result = await translate_office_document(
        _TranslationProvider(),
        source,
        _task_data(
            source_language="English",
            target_language="Chinese",
            output_format="pdf",
        ),
        output_dir=tmp_path / "result",
    )

    output = result.artifacts["pdf"]
    assert output.is_file() and output.stat().st_size > 1000
    assert len(PdfReader(str(output)).pages) >= 1


@pytest.mark.asyncio
async def test_unsupported_source_and_wrong_delivery_explain_recovery(
    tmp_path: Path,
) -> None:
    source = tmp_path / "slides.pptx"
    source.write_bytes(b"not a supported translation source")

    with pytest.raises(TranslationError) as unsupported_error:
        await translate_office_document(
            _TranslationProvider(),
            source,
            _task_data(output_format="docx"),
            output_dir=tmp_path / "unsupported",
        )
    assert unsupported_error.value.code == "unsupported_format"
    assert ".docx" in unsupported_error.value.recovery

    source = tmp_path / "quote.xlsx"
    Workbook().save(source)
    with pytest.raises(TranslationError) as delivery_error:
        await translate_office_document(
            _TranslationProvider(),
            source,
            _task_data(output_format="docx"),
            output_dir=tmp_path / "wrong-delivery",
        )
    assert delivery_error.value.code == "unsupported_delivery"
    assert "Excel" in delivery_error.value.recovery
