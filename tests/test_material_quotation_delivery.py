from __future__ import annotations

import json
from zipfile import ZipFile

from dc_engines.material_quotation import normalize_material_quotation
from dc_engines.material_quotation_delivery import (
    generate_material_quotation_deliverables,
)
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader


def _normalized_estimate() -> dict[str, str]:
    return normalize_material_quotation(
        {
            "project_name": "柳州用户共创会",
            "client_name": "五菱",
            "validity_days": "15",
            "quotation_items": json.dumps(
                [
                    {
                        "item_name": "背胶",
                        "specification": "过哑膜",
                        "remark": "含安装，正式采购前复核",
                        "quantity": "10",
                        "unit": "㎡",
                        "unit_price": "8",
                        "supplier": "柳州东成广告",
                        "source_path": "供应商库/柳州-南宁搭建物料.md",
                        "source_status": "待复核",
                    },
                    {
                        "item_name": "现场摄影",
                        "quantity": "1",
                        "unit": "项",
                        "unit_price": "",
                    },
                ],
                ensure_ascii=False,
            ),
            "transport_fee": "20",
            "profit_rate": "10",
            "tax_rate": "6",
        }
    )


def test_material_quotation_delivery_generates_real_office_files(tmp_path) -> None:
    artifacts = generate_material_quotation_deliverables(
        _normalized_estimate(),
        output_dir=tmp_path,
        version=2,
    )

    assert set(artifacts) == {"internal_xlsx", "market_docx", "market_pdf"}
    assert artifacts["internal_xlsx"].read_bytes().startswith(b"PK")
    assert artifacts["market_docx"].read_bytes().startswith(b"PK")
    assert artifacts["market_pdf"].read_bytes().startswith(b"%PDF")


def test_internal_excel_keeps_traceable_calculations(tmp_path) -> None:
    artifacts = generate_material_quotation_deliverables(
        _normalized_estimate(),
        output_dir=tmp_path,
        version=1,
    )
    workbook = load_workbook(artifacts["internal_xlsx"], data_only=False)
    assert workbook.sheetnames == ["成本清单", "测算溯源"]
    sheet = workbook["成本清单"]
    trace_sheet = workbook["测算溯源"]
    text = "\n".join(
        str(cell.value or "")
        for worksheet in (sheet, trace_sheet)
        for row in worksheet.iter_rows()
        for cell in row
    )

    assert "内部资料" in text
    assert "柳州东成广告" in text
    assert "供应商库/柳州-南宁搭建物料.md" in text
    assert "利润额" in text
    assert sheet["A1"].value == "柳州用户共创会成本清单"
    assert [sheet.cell(row=3, column=column).value for column in range(1, 10)] == [
        "序号",
        "礼品 / 物料组合",
        "规格",
        "数量",
        "单位",
        "单价",
        "小计",
        "图示",
        "备注",
    ]
    assert sheet["G4"].value == "=D4*F4"
    assert sheet["I4"].value == "含安装，正式采购前复核"
    assert sheet["H4"].value is None
    assert sheet["I10"].value == "利润率 10%"
    assert str(sheet.cell(row=sheet.max_row, column=7).value).startswith("=SUM(")
    assert trace_sheet["G6"].value == "=D6*F6"
    assert str(trace_sheet["B20"].value).startswith("=")


def test_internal_excel_embeds_uploaded_item_image(tmp_path) -> None:
    image_root = tmp_path / "nas"
    image_path = (
        image_root / "projects" / "筹备组物料报价" / "draft" / "images" / "背胶.png"
    )
    image_path.parent.mkdir(parents=True)
    Image.new("RGB", (120, 80), "#45613f").save(image_path)
    estimate = _normalized_estimate()
    items = json.loads(estimate["quotation_items"])
    items[0]["image_path"] = str(image_path.relative_to(image_root))
    estimate["quotation_items"] = json.dumps(items, ensure_ascii=False)

    artifacts = generate_material_quotation_deliverables(
        estimate,
        output_dir=tmp_path / "output",
        version=1,
        image_root=image_root,
    )

    with ZipFile(artifacts["internal_xlsx"]) as archive:
        assert any(name.startswith("xl/media/") for name in archive.namelist())
    workbook = load_workbook(artifacts["internal_xlsx"])
    sheet = workbook["成本清单"]
    assert len(sheet._images) == 1
    embedded_image = sheet._images[0]
    assert embedded_image.anchor._from.col == 7
    assert embedded_image.anchor._from.row == 3
    assert embedded_image.anchor.ext.cx <= 110 * 9525
    assert embedded_image.anchor.ext.cy <= 64 * 9525
    assert sheet.row_dimensions[4].height == 52


def test_internal_excel_lays_out_batch_images_in_one_item_row(tmp_path) -> None:
    image_root = tmp_path / "nas"
    image_dir = image_root / "projects" / "筹备组物料报价" / "draft" / "images"
    image_dir.mkdir(parents=True)
    image_paths = []
    colors = ("#45613f", "#d96b32", "#345f78", "#c4a15a", "#6b4f3a", "#7b8268")
    for index, color in enumerate(colors, start=1):
        image_path = image_dir / f"背胶-{index}.png"
        Image.new("RGB", (120 + index * 10, 80), color).save(image_path)
        image_paths.append(str(image_path.relative_to(image_root)))
    estimate = _normalized_estimate()
    items = json.loads(estimate["quotation_items"])
    items[0]["image_paths"] = image_paths
    estimate["quotation_items"] = json.dumps(items, ensure_ascii=False)

    artifacts = generate_material_quotation_deliverables(
        estimate,
        output_dir=tmp_path / "output",
        version=1,
        image_root=image_root,
    )

    with ZipFile(artifacts["internal_xlsx"]) as archive:
        assert (
            len([name for name in archive.namelist() if name.startswith("xl/media/")])
            == 6
        )
    workbook = load_workbook(artifacts["internal_xlsx"])
    sheet = workbook["成本清单"]
    assert len(sheet._images) == 6
    assert {
        (image.anchor._from.col, image.anchor._from.row) for image in sheet._images
    } == {(7, 3)}
    assert (
        len(
            {
                (image.anchor._from.colOff, image.anchor._from.rowOff)
                for image in sheet._images
            }
        )
        == 6
    )
    assert sheet.row_dimensions[4].height == 52


def test_market_word_and_pdf_exclude_internal_cost_details(tmp_path) -> None:
    artifacts = generate_material_quotation_deliverables(
        _normalized_estimate(),
        output_dir=tmp_path,
        version=1,
    )
    document = Document(artifacts["market_docx"])
    word_text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
    )
    pdf_text = "\n".join(
        page.extract_text() or "" for page in PdfReader(artifacts["market_pdf"]).pages
    )

    for text in (word_text, pdf_text):
        assert "柳州用户共创会" in text
        assert "五菱" in text
        assert "¥116.60" in text or "116.60" in text
        assert "15 天" in text
        assert "估价方案" in text
        assert "背胶" not in text
        assert "柳州东成广告" not in text
        assert "供应商库" not in text
        assert "利润" not in text
        assert "运输费" not in text
