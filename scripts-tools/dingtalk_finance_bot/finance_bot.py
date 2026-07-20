#!/usr/bin/env python3
"""DingTalk finance reimbursement bot probe and workbook builder."""

from __future__ import annotations

import argparse
import asyncio
import hashlib
import hmac
import html
import json
import logging
import mimetypes
import os
import re
import secrets
import sys
import threading
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from http.cookies import SimpleCookie
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, quote, unquote, urlencode, urlparse

import httpx
from openpyxl import Workbook
from openpyxl.drawing.image import Image as WorkbookImage
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_ENV_PATH = SCRIPT_DIR / ".env"
DEFAULT_CONFIG_PATH = SCRIPT_DIR / "config.yaml"
EXAMPLE_CONFIG_PATH = SCRIPT_DIR / "config.example.yaml"
DEFAULT_OUTPUT_DIR = SCRIPT_DIR / "output"
DEFAULT_DEDUPE_LEDGER_PATH = SCRIPT_DIR / "invoice_ledger.local.json"
API_BASE_URL = "https://api.dingtalk.com"
OAPI_BASE_URL = "https://oapi.dingtalk.com"
DINGTALK_OAUTH_AUTHORIZE_URL = "https://login.dingtalk.com/oauth2/auth"
FINANCE_TOOL_SESSION_COOKIE = "dingtalk_finance_session"
FINANCE_TOOL_FALLBACK_COOKIE = "dingtalk_finance_tool"
DEFAULT_OA_CALLBACK_PATH = "/dingtalk/finance/oa-callback"
SUMMARY_COLUMNS = [
    "序号",
    "部门",
    "姓名",
    "报销项目",
    "日期",
    "金额",
    "附件 sheet",
    "备注说明",
]
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".bmp", ".gif"}
ATTACHMENT_KEYS = {
    "downloadurl",
    "extension",
    "file_id",
    "file_name",
    "fileid",
    "filename",
    "filesize",
    "fileurl",
    "media_id",
    "mediaid",
    "space_id",
    "spaceid",
}
ATTACHMENT_NAME_KEYS = {"file_name", "filename", "name"}
ATTACHMENT_ID_KEYS = {"file_id", "fileid", "media_id", "mediaid"}
ATTACHMENT_LOCATION_KEYS = {"downloadurl", "fileurl", "space_id", "spaceid"}
USER_ID_KEYS = {
    "actioner_user_id",
    "actioneruserid",
    "originator_userid",
    "originatoruserid",
    "user_id",
    "userid",
}
TASK_ID_KEYS = {"task_id", "taskid"}
CALLBACK_INSTANCE_ID_KEYS = {
    "approval_instance_id",
    "approvalinstanceid",
    "instance_id",
    "instanceid",
    "process_instance_id",
    "processinstanceid",
    "proc_ins_id",
    "procinsid",
    "procinstid",
}
REVIEW_STATUS_LABELS = {
    "needs_manual_review": "需财务处理",
    "ready_for_ocr": "待 OCR 核对",
    "ready_for_finance_review": "可财务复审",
}
REVIEW_SUMMARY_COLUMNS = [
    "序号",
    "员工",
    "部门",
    "报销项目",
    "金额",
    "初审状态",
    "问题",
    "附件数",
    "处理建议",
    "报告文件",
]


@dataclass(slots=True)
class ExportConfig:
    """Runtime configuration for DingTalk finance exports.

    Args:
        app_key: DingTalk application AppKey or Client ID.
        app_secret: DingTalk application Client Secret.
        process_code: DingTalk approval process code.
        start_time: Inclusive approval creation start time.
        end_time: Inclusive approval creation end time.
        output_dir: Local directory for attachments and workbook output.
        app_id: Optional DingTalk application ID for operator reference.
        agent_id: Optional DingTalk AgentId for operator reference.
        operator_user_id: Optional DingTalk userId for file download authorization.
        operator_union_id: Optional DingTalk unionId for storage download URLs.
        oa_callback_token: Optional shared token for OA automation callbacks.
        local_ocr_url: Optional local OCR service endpoint.
        local_ocr_timeout_seconds: Timeout for local OCR requests.
        finance_tool_token: Optional fallback shared token for the NAS web tool.
        finance_tool_auth_mode: Optional NAS web auth mode override.
        finance_tool_base_url: Optional public or intranet base URL for OAuth.
        finance_allowed_union_ids: Optional DingTalk unionId allowlist.
        finance_reviewer_name: DingTalk reviewer who owns this finance tool.
        finance_reviewer_user_id: Optional DingTalk userId for that reviewer.
        finance_reviewer_union_id: Optional DingTalk unionId for that reviewer.
        general_manager_keywords: Keywords that identify Yang's approval step.
        finance_notify_enabled: Enable DingTalk work notifications for reports.
        finance_notify_user_ids: DingTalk userIds receiving report notifications.
        finance_notify_on: Notification policy, either ``issues`` or ``all``.
        finance_stream_enabled: Enable DingTalk event Stream listener.
    """

    app_key: str
    app_secret: str
    process_code: str
    start_time: datetime
    end_time: datetime
    output_dir: Path = DEFAULT_OUTPUT_DIR
    app_id: str = ""
    agent_id: str = ""
    channel_name: str = "财务报销机器人"
    operator_user_id: str = ""
    operator_union_id: str = ""
    oa_callback_token: str = ""
    local_ocr_url: str = ""
    local_ocr_timeout_seconds: float = 30.0
    finance_tool_token: str = ""
    finance_tool_auth_mode: str = ""
    finance_tool_base_url: str = ""
    finance_allowed_union_ids: tuple[str, ...] = ()
    finance_reviewer_name: str = "覃献芳"
    finance_reviewer_user_id: str = ""
    finance_reviewer_union_id: str = ""
    general_manager_keywords: tuple[str, ...] = ("杨国民", "杨总")
    finance_notify_enabled: bool = False
    finance_notify_user_ids: tuple[str, ...] = ()
    finance_notify_on: str = "issues"
    finance_stream_enabled: bool = False


def load_env_file(path: Path = DEFAULT_ENV_PATH) -> None:
    """Load simple KEY=VALUE pairs from a local env file if present."""
    if not path.exists():
        return
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        value = value.strip().strip('"').strip("'")
        os.environ.setdefault(key.strip(), value)


def parse_simple_yaml(path: Path) -> dict[str, Any]:
    """Parse the flat key/value YAML shape used by this tool.

    This intentionally supports only the simple template format in
    config.example.yaml so the first interface probe can run without PyYAML.
    """
    if not path.exists():
        return {}
    result: dict[str, Any] = {}
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or ":" not in line:
            continue
        key, value = line.split(":", 1)
        value = value.strip().strip('"').strip("'")
        result[key.strip()] = value
    return result


def write_json(path: Path, data: Any) -> None:
    """Write JSON with stable formatting for local probe artifacts.

    Args:
        path: Output file path.
        data: JSON-serializable data.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def file_sha256(path: Path) -> str:
    """Compute the SHA-256 digest for one local attachment.

    Args:
        path: Local file path.

    Returns:
        Hex-encoded SHA-256 digest.
    """
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_duplicate_ledger(path: Path) -> list[dict[str, Any]]:
    """Load the local duplicate invoice ledger.

    Args:
        path: JSON ledger path. Missing files are treated as an empty ledger.

    Returns:
        Ledger records with hashes such as ``sha256`` or ``file_sha256``.

    Raises:
        ValueError: The ledger JSON shape is unsupported.
    """
    if not path.exists():
        return []
    data = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(data, list):
        return [record for record in data if isinstance(record, dict)]
    if isinstance(data, dict):
        records = data.get("records") or data.get("invoices") or []
        if isinstance(records, list):
            return [record for record in records if isinstance(record, dict)]
    raise ValueError(f"Unsupported duplicate ledger shape: {path}")


def load_config(config_path: Path = DEFAULT_CONFIG_PATH) -> ExportConfig:
    """Load bot config from a local YAML file and environment variables.

    Args:
        config_path: Optional YAML file with non-secret defaults.

    Returns:
        Parsed bot configuration.

    Raises:
        ValueError: Required credentials or runtime parameters are missing.
    """
    load_env_file(DEFAULT_ENV_PATH)
    data = parse_simple_yaml(config_path)
    if not data and config_path == DEFAULT_CONFIG_PATH:
        data = parse_simple_yaml(EXAMPLE_CONFIG_PATH)

    app_key = (
        os.getenv("DINGTALK_APP_KEY")
        or os.getenv("DINGTALK_CLIENT_ID")
        or str(data.get("app_key") or data.get("client_id") or "")
    )
    app_secret = (
        os.getenv("DINGTALK_APP_SECRET")
        or os.getenv("DINGTALK_CLIENT_SECRET")
        or str(data.get("app_secret") or data.get("client_secret") or "")
    )
    process_code = os.getenv("DINGTALK_PROCESS_CODE") or str(
        data.get("process_code", "")
    )
    start_time = (
        os.getenv("DINGTALK_START_TIME")
        or os.getenv("DINGTALK_START_DATE")
        or str(data.get("start_time", ""))
    )
    end_time = (
        os.getenv("DINGTALK_END_TIME")
        or os.getenv("DINGTALK_END_DATE")
        or str(data.get("end_time", ""))
    )
    output_dir = os.getenv("DINGTALK_OUTPUT_DIR") or str(
        data.get("output_dir", DEFAULT_OUTPUT_DIR)
    )
    oa_callback_token = os.getenv("DINGTALK_OA_CALLBACK_TOKEN") or str(
        data.get("oa_callback_token", "")
    )
    if oa_callback_token.startswith("replace_with_"):
        oa_callback_token = ""
    local_ocr_url = os.getenv("DINGTALK_FINANCE_OCR_URL") or str(
        data.get("local_ocr_url", "")
    )
    if local_ocr_url.startswith("replace_with_"):
        local_ocr_url = ""
    local_ocr_timeout_raw = os.getenv("DINGTALK_FINANCE_OCR_TIMEOUT_SECONDS") or str(
        data.get("local_ocr_timeout_seconds", "30")
    )
    try:
        local_ocr_timeout_seconds = max(1.0, float(local_ocr_timeout_raw))
    except ValueError:
        local_ocr_timeout_seconds = 30.0
    finance_tool_token = (
        os.getenv("DINGTALK_FINANCE_TOOL_TOKEN")
        or str(data.get("finance_tool_token", ""))
        or oa_callback_token
    )
    if finance_tool_token.startswith("replace_with_"):
        finance_tool_token = ""
    finance_tool_auth_mode = (
        os.getenv("DINGTALK_FINANCE_TOOL_AUTH_MODE")
        or str(data.get("finance_tool_auth_mode", ""))
    ).strip()
    if finance_tool_auth_mode.startswith("replace_with_"):
        finance_tool_auth_mode = ""
    finance_tool_base_url = os.getenv("DINGTALK_FINANCE_TOOL_BASE_URL") or str(
        data.get("finance_tool_base_url", "")
    )
    if finance_tool_base_url.startswith("replace_with_"):
        finance_tool_base_url = ""
    allowed_union_ids_raw = os.getenv("DINGTALK_FINANCE_ALLOWED_UNION_IDS") or str(
        data.get("finance_allowed_union_ids", "")
    )
    finance_allowed_union_ids = tuple(
        item
        for item in re.split(r"[,;\s]+", allowed_union_ids_raw)
        if item and not item.startswith("replace_with_")
    )
    finance_reviewer_name = os.getenv("DINGTALK_FINANCE_REVIEWER_NAME") or str(
        data.get("finance_reviewer_name", "覃献芳")
    )
    finance_reviewer_user_id = os.getenv("DINGTALK_FINANCE_REVIEWER_USER_ID") or str(
        data.get("finance_reviewer_user_id", "")
    )
    if finance_reviewer_user_id.startswith("replace_with_"):
        finance_reviewer_user_id = ""
    finance_reviewer_union_id = os.getenv("DINGTALK_FINANCE_REVIEWER_UNION_ID") or str(
        data.get("finance_reviewer_union_id", "")
    )
    if finance_reviewer_union_id.startswith("replace_with_"):
        finance_reviewer_union_id = ""
    general_manager_keywords_raw = os.getenv(
        "DINGTALK_FINANCE_GENERAL_MANAGER_KEYWORDS"
    ) or str(data.get("general_manager_keywords", "杨国民,杨总"))
    general_manager_keywords = tuple(
        item
        for item in re.split(r"[,;\s]+", general_manager_keywords_raw)
        if item and not item.startswith("replace_with_")
    )
    finance_notify_enabled_raw = (
        os.getenv("DINGTALK_FINANCE_NOTIFY_ENABLED")
        or str(data.get("finance_notify_enabled", "0"))
    ).strip()
    finance_notify_enabled = finance_notify_enabled_raw.lower() in {
        "1",
        "true",
        "yes",
        "on",
    }
    finance_notify_user_ids_raw = os.getenv("DINGTALK_FINANCE_NOTIFY_USER_IDS") or str(
        data.get("finance_notify_user_ids", "")
    )
    finance_notify_user_ids = tuple(
        item
        for item in re.split(r"[,;\s]+", finance_notify_user_ids_raw)
        if item and not item.startswith("replace_with_")
    )
    finance_notify_on = (
        os.getenv("DINGTALK_FINANCE_NOTIFY_ON")
        or str(data.get("finance_notify_on", "issues"))
    ).strip()
    if finance_notify_on not in {"issues", "all", "none"}:
        finance_notify_on = "issues"
    finance_stream_enabled_raw = (
        os.getenv("DINGTALK_FINANCE_STREAM_ENABLED")
        or str(data.get("finance_stream_enabled", "0"))
    ).strip()
    finance_stream_enabled = finance_stream_enabled_raw.lower() in {
        "1",
        "true",
        "yes",
        "on",
    }

    missing = [
        name
        for name, value in {
            "DINGTALK_APP_KEY": app_key,
            "DINGTALK_APP_SECRET": app_secret,
            "DINGTALK_PROCESS_CODE": process_code,
            "DINGTALK_START_TIME": start_time,
            "DINGTALK_END_TIME": end_time,
        }.items()
        if not value or value.startswith("replace_with_")
    ]
    if missing:
        raise ValueError(f"Missing required local config values: {', '.join(missing)}")

    return ExportConfig(
        app_key=app_key,
        app_secret=app_secret,
        process_code=process_code,
        start_time=datetime.fromisoformat(start_time),
        end_time=datetime.fromisoformat(end_time),
        output_dir=Path(output_dir)
        if Path(output_dir).is_absolute()
        else SCRIPT_DIR / output_dir,
        app_id=os.getenv("DINGTALK_APP_ID") or str(data.get("app_id", "")),
        agent_id=os.getenv("DINGTALK_AGENT_ID") or str(data.get("agent_id", "")),
        channel_name=os.getenv("DINGTALK_CHANNEL_NAME")
        or str(data.get("channel_name", "财务报销机器人")),
        operator_user_id=os.getenv("DINGTALK_OPERATOR_USER_ID")
        or str(data.get("operator_user_id", "")),
        operator_union_id=os.getenv("DINGTALK_OPERATOR_UNION_ID")
        or str(data.get("operator_union_id", "")),
        oa_callback_token=oa_callback_token,
        local_ocr_url=local_ocr_url.rstrip("/"),
        local_ocr_timeout_seconds=local_ocr_timeout_seconds,
        finance_tool_token=finance_tool_token,
        finance_tool_auth_mode=finance_tool_auth_mode,
        finance_tool_base_url=finance_tool_base_url.rstrip("/"),
        finance_allowed_union_ids=finance_allowed_union_ids,
        finance_reviewer_name=finance_reviewer_name,
        finance_reviewer_user_id=finance_reviewer_user_id,
        finance_reviewer_union_id=finance_reviewer_union_id,
        general_manager_keywords=general_manager_keywords or ("杨国民", "杨总"),
        finance_notify_enabled=finance_notify_enabled,
        finance_notify_user_ids=finance_notify_user_ids,
        finance_notify_on=finance_notify_on,
        finance_stream_enabled=finance_stream_enabled,
    )


class DingTalkFinanceClient:
    """Small DingTalk Workflow API client for finance approval export."""

    def __init__(
        self,
        config: ExportConfig,
        *,
        http_client: httpx.Client | None = None,
        api_base_url: str = API_BASE_URL,
        oapi_base_url: str = OAPI_BASE_URL,
    ) -> None:
        """Create a client.

        Args:
            config: Runtime bot configuration.
            http_client: Optional injected HTTP client for offline tests.
            api_base_url: DingTalk Open Platform API base URL.
            oapi_base_url: DingTalk legacy Open API base URL.
        """
        self.config = config
        self.api_base_url = api_base_url.rstrip("/")
        self.oapi_base_url = oapi_base_url.rstrip("/")
        self.http = http_client or httpx.Client(timeout=30)

    def get_access_token(self) -> str:
        """Request a DingTalk app access token.

        Returns:
            DingTalk access token string.

        Raises:
            RuntimeError: DingTalk returns an error or an unexpected response.
        """
        errors: list[str] = []
        try:
            response = self.http.post(
                f"{self.api_base_url}/v1.0/oauth2/accessToken",
                json={
                    "appKey": self.config.app_key,
                    "appSecret": self.config.app_secret,
                },
            )
            response.raise_for_status()
            data = response.json()
            token = data.get("accessToken")
            if token:
                return str(token)
            errors.append(f"new token response missing accessToken: {data}")
        except httpx.HTTPError as exc:
            errors.append(f"new token endpoint failed: {exc}")

        response = self.http.get(
            f"{self.oapi_base_url}/gettoken",
            params={
                "appkey": self.config.app_key,
                "appsecret": self.config.app_secret,
            },
        )
        response.raise_for_status()
        data = response.json()
        if data.get("errcode") not in (0, None):
            errors.append(f"old token endpoint returned error: {data}")
        token = data.get("access_token") or data.get("accessToken")
        if token:
            return str(token)
        raise RuntimeError("Unable to obtain DingTalk token: " + "; ".join(errors))

    def list_approval_instance_ids(self, access_token: str) -> list[str]:
        """List approval process instance IDs for the configured process and range.

        Args:
            access_token: DingTalk access token.

        Returns:
            Ordered approval instance IDs.
        """
        instance_ids: list[str] = []
        next_token = ""
        cursor = 0
        while True:
            payload: dict[str, Any] = {
                "process_code": self.config.process_code,
                "start_time": int(self.config.start_time.timestamp() * 1000),
                "end_time": int(self.config.end_time.timestamp() * 1000),
                "size": 20,
                "cursor": cursor,
            }
            response = self.http.post(
                f"{self.oapi_base_url}/topapi/processinstance/listids",
                params={"access_token": access_token},
                json=payload,
            )
            response.raise_for_status()
            data = response.json()
            if data.get("errcode") not in (0, None):
                raise RuntimeError(f"DingTalk listids returned error: {data}")
            result = data.get("result") or {}
            values = result.get("list") or []
            instance_ids.extend(str(value) for value in values)
            next_token = str(result.get("next_cursor") or "")
            if not next_token or not values:
                return instance_ids
            cursor = int(next_token)

    def get_approval_detail(
        self, access_token: str, instance_id: str
    ) -> dict[str, Any]:
        """Fetch one approval process instance detail.

        Args:
            access_token: DingTalk access token.
            instance_id: DingTalk process instance ID.

        Returns:
            DingTalk approval detail payload.
        """
        response = self.http.post(
            f"{self.oapi_base_url}/topapi/processinstance/get",
            params={"access_token": access_token},
            json={"process_instance_id": instance_id},
        )
        response.raise_for_status()
        data = response.json()
        if data.get("errcode") not in (0, None):
            raise RuntimeError(f"DingTalk processinstance/get returned error: {data}")
        result = data.get("process_instance") or data.get("result")
        return result if isinstance(result, dict) else data

    def get_user_detail(self, access_token: str, user_id: str) -> dict[str, Any]:
        """Fetch DingTalk user detail by userId.

        Args:
            access_token: DingTalk access token.
            user_id: DingTalk userId from approval detail.

        Returns:
            DingTalk user detail payload, including unionid when permitted.
        """
        response = self.http.post(
            f"{self.oapi_base_url}/topapi/v2/user/get",
            params={"access_token": access_token},
            json={"userid": user_id, "language": "zh_CN"},
        )
        response.raise_for_status()
        data = response.json()
        if data.get("errcode") not in (0, None):
            raise RuntimeError(f"DingTalk v2/user/get returned error: {data}")
        result = data.get("result")
        return result if isinstance(result, dict) else data

    def get_user_by_union_id(self, access_token: str, union_id: str) -> dict[str, Any]:
        """Resolve a DingTalk unionId to the approval userId.

        Args:
            access_token: DingTalk app access token.
            union_id: DingTalk unionId from OAuth login.

        Returns:
            DingTalk user mapping payload with ``userid`` when permitted.
        """
        response = self.http.post(
            f"{self.oapi_base_url}/topapi/user/getbyunionid",
            params={"access_token": access_token},
            json={"unionid": union_id},
        )
        response.raise_for_status()
        data = response.json()
        if data.get("errcode") not in (0, None):
            raise RuntimeError(f"DingTalk getbyunionid returned error: {data}")
        result = data.get("result")
        return result if isinstance(result, dict) else data

    def send_work_notification(
        self,
        access_token: str,
        *,
        user_ids: tuple[str, ...],
        content: str,
        agent_id: str,
    ) -> dict[str, Any]:
        """Send a DingTalk enterprise work notification.

        Args:
            access_token: DingTalk app access token.
            user_ids: DingTalk userIds receiving the notification.
            content: Text message body.
            agent_id: DingTalk internal app AgentId.

        Returns:
            DingTalk async send result payload.

        Raises:
            ValueError: Required notification fields are missing.
            RuntimeError: DingTalk returns a failed response.
        """
        normalized_user_ids = tuple(user_id for user_id in user_ids if user_id)
        if not normalized_user_ids:
            raise ValueError("user_ids cannot be empty")
        if not agent_id:
            raise ValueError("agent_id is required for work notifications")
        response = self.http.post(
            f"{self.oapi_base_url}/topapi/message/corpconversation/asyncsend_v2",
            params={"access_token": access_token},
            json={
                "agent_id": int(agent_id) if str(agent_id).isdigit() else agent_id,
                "userid_list": ",".join(normalized_user_ids),
                "msg": {
                    "msgtype": "text",
                    "text": {"content": content},
                },
            },
        )
        response.raise_for_status()
        data = response.json()
        if data.get("errcode") not in (0, None):
            raise RuntimeError(f"DingTalk work notification returned error: {data}")
        result = data.get("task_id") or data.get("result") or data
        return result if isinstance(result, dict) else {"task_id": result}

    def get_oauth_user_access_token(self, code: str) -> dict[str, Any]:
        """Exchange a DingTalk OAuth code for a user access token.

        Args:
            code: Authorization code returned by DingTalk login.

        Returns:
            Token payload containing accessToken and optional corpId.

        Raises:
            RuntimeError: DingTalk returns an error or omits accessToken.
        """
        response = self.http.post(
            f"{self.api_base_url}/v1.0/oauth2/userAccessToken",
            json={
                "clientId": self.config.app_key,
                "clientSecret": self.config.app_secret,
                "code": code,
                "grantType": "authorization_code",
            },
        )
        response.raise_for_status()
        data = response.json()
        if data.get("success") is False or data.get("errcode"):
            raise RuntimeError(f"DingTalk OAuth token returned error: {data}")
        if not data.get("accessToken"):
            raise RuntimeError(f"DingTalk OAuth token missing accessToken: {data}")
        return data

    def get_oauth_user_info(self, user_access_token: str) -> dict[str, Any]:
        """Read the DingTalk user profile for an OAuth user token.

        Args:
            user_access_token: DingTalk OAuth user access token.

        Returns:
            DingTalk user profile, including unionId/openId when available.

        Raises:
            RuntimeError: DingTalk returns a failed response.
        """
        response = self.http.get(
            f"{self.api_base_url}/v1.0/contact/users/me",
            headers={"x-acs-dingtalk-access-token": user_access_token},
        )
        response.raise_for_status()
        data = response.json()
        if data.get("success") is False or data.get("errcode"):
            raise RuntimeError(f"DingTalk OAuth user info returned error: {data}")
        result = data.get("result")
        return result if isinstance(result, dict) else data

    def execute_approval_task(
        self,
        access_token: str,
        *,
        process_instance_id: str,
        task_id: int | str,
        actioner_user_id: str,
        result: str,
        remark: str = "",
        file: dict[str, Any] | None = None,
    ) -> bool:
        """Approve or reject the current DingTalk approval task.

        Args:
            access_token: DingTalk access token.
            process_instance_id: DingTalk process instance ID.
            task_id: DingTalk task ID from approval detail.
            actioner_user_id: DingTalk userId of the task operator.
            result: Approval operation, either ``agree`` or ``refuse``.
            remark: Optional approval comment.
            file: Optional DingTalk approval file object.

        Returns:
            True when DingTalk reports the task operation succeeded.

        Raises:
            ValueError: The requested result is not supported.
            RuntimeError: DingTalk returns a failed response.
        """
        if result not in {"agree", "refuse"}:
            raise ValueError("result must be agree or refuse")
        payload: dict[str, Any] = {
            "processInstanceId": process_instance_id,
            "taskId": task_id,
            "actionerUserId": actioner_user_id,
            "result": result,
            "remark": remark,
        }
        if file:
            payload["file"] = file
        response = self.http.post(
            f"{self.api_base_url}/v1.0/workflow/processInstances/execute",
            headers={"x-acs-dingtalk-access-token": access_token},
            json=payload,
        )
        response.raise_for_status()
        data = response.json()
        if data.get("success") is False or data.get("result") is False:
            raise RuntimeError(f"DingTalk task execute returned error: {data}")
        return bool(data.get("result", data.get("success", True)))

    def get_approval_attachment_space(self, access_token: str, user_id: str) -> str:
        """Get the DingTalk approval DingDrive space for one operator.

        Args:
            access_token: DingTalk access token.
            user_id: DingTalk userId used for approval file access.

        Returns:
            Approval DingDrive space ID.

        Raises:
            RuntimeError: DingTalk returns a failed or incomplete response.
        """
        payload: dict[str, Any] = {"userId": user_id}
        if self.config.agent_id:
            payload["agentId"] = self.config.agent_id
        response = self.http.post(
            f"{self.api_base_url}/v1.0/workflow/processInstances/spaces/infos/query",
            headers={"x-acs-dingtalk-access-token": access_token},
            json=payload,
        )
        response.raise_for_status()
        data = response.json()
        if data.get("success") is False:
            raise RuntimeError(f"DingTalk attachment space returned error: {data}")
        result = data.get("result") if isinstance(data.get("result"), dict) else data
        space_id = result.get("spaceId") if isinstance(result, dict) else None
        if not space_id:
            raise RuntimeError(f"DingTalk attachment space missing spaceId: {data}")
        return str(space_id)

    def authorize_approval_file_download(
        self,
        access_token: str,
        *,
        user_id: str,
        file_infos: list[dict[str, Any]],
    ) -> bool:
        """Authorize a user to download approval DingDrive files.

        Args:
            access_token: DingTalk access token.
            user_id: DingTalk userId receiving download authorization.
            file_infos: File metadata with ``fileId`` and ``spaceId``.

        Returns:
            True when DingTalk reports authorization succeeded.

        Raises:
            RuntimeError: DingTalk returns a failed response.
        """
        normalized_file_infos: list[dict[str, Any]] = []
        for file_info in file_infos:
            space_id = file_info.get("spaceId")
            if isinstance(space_id, str) and space_id.isdigit():
                space_id = int(space_id)
            normalized_file_infos.append(
                {
                    **file_info,
                    "fileId": str(file_info.get("fileId", "")),
                    "spaceId": space_id,
                }
            )
        response = self.http.post(
            f"{self.api_base_url}/v1.0/workflow/processInstances/spaces/files/authDownload",
            headers={"x-acs-dingtalk-access-token": access_token},
            json={"userId": user_id, "fileInfos": normalized_file_infos},
        )
        try:
            response.raise_for_status()
        except httpx.HTTPStatusError as exc:
            raise RuntimeError(
                "DingTalk file download auth HTTP "
                f"{response.status_code}: {response.text}"
            ) from exc
        data = response.json()
        if data.get("success") is False or data.get("result") is False:
            raise RuntimeError(f"DingTalk file download auth returned error: {data}")
        return bool(data.get("result", data.get("success", True)))

    def get_approval_file_download_info(
        self,
        access_token: str,
        *,
        process_instance_id: str,
        file_id: str,
    ) -> dict[str, Any]:
        """Get a DingTalk approval attachment download URL.

        Args:
            access_token: DingTalk access token.
            process_instance_id: DingTalk process instance ID.
            file_id: Approval attachment file ID.

        Returns:
            DingTalk approval download info with downloadUri when available.

        Raises:
            RuntimeError: DingTalk returns a failed response.
        """
        response = self.http.post(
            f"{self.api_base_url}/v1.0/workflow/processInstances/spaces/files/urls/download",
            headers={"x-acs-dingtalk-access-token": access_token},
            json={"processInstanceId": process_instance_id, "fileId": str(file_id)},
        )
        try:
            response.raise_for_status()
        except httpx.HTTPStatusError as exc:
            raise RuntimeError(
                "DingTalk approval file download URL HTTP "
                f"{response.status_code}: {response.text}"
            ) from exc
        data = response.json()
        if data.get("success") is False:
            raise RuntimeError(
                f"DingTalk approval file download URL returned error: {data}"
            )
        result = data.get("result")
        return result if isinstance(result, dict) else data

    def get_storage_file_download_info(
        self,
        access_token: str,
        *,
        space_id: str,
        dentry_id: str,
        union_id: str,
        version: int | None = None,
        prefer_intranet: bool | None = None,
    ) -> dict[str, Any]:
        """Get DingTalk storage download metadata for a file.

        Args:
            access_token: DingTalk access token.
            space_id: DingTalk storage space ID.
            dentry_id: DingTalk file/dentry ID.
            union_id: Operator unionId required by the storage API.
            version: Optional file version.
            prefer_intranet: Optional intranet transfer preference.

        Returns:
            DingTalk storage download info payload.
        """
        body: dict[str, Any] = {}
        option: dict[str, Any] = {}
        if version is not None:
            option["version"] = version
        if prefer_intranet is not None:
            option["preferIntranet"] = prefer_intranet
        if option:
            body["option"] = option
        response = self.http.post(
            f"{self.api_base_url}/v1.0/storage/spaces/"
            f"{quote(str(space_id), safe='')}/dentries/"
            f"{quote(str(dentry_id), safe='')}/downloadInfos/query",
            params={"unionId": union_id},
            headers={"x-acs-dingtalk-access-token": access_token},
            json=body,
        )
        try:
            response.raise_for_status()
        except httpx.HTTPStatusError as exc:
            raise RuntimeError(
                "DingTalk storage download info HTTP "
                f"{response.status_code}: {response.text}"
            ) from exc
        data = response.json()
        if data.get("success") is False:
            raise RuntimeError(f"DingTalk storage download info returned error: {data}")
        return data.get("result") if isinstance(data.get("result"), dict) else data

    def download_approval_attachment(
        self,
        access_token: str,
        attachment: dict[str, Any],
        output_dir: Path,
    ) -> Path:
        """Download an approval attachment when its detail includes a URL.

        Args:
            access_token: DingTalk access token for guarded download URLs.
            attachment: Attachment metadata extracted from approval detail.
            output_dir: Local directory for downloaded files.

        Returns:
            Local attachment path.

        Raises:
            ValueError: The attachment metadata does not include a download URL.
        """
        download_url = (
            attachment.get("downloadUrl")
            or attachment.get("url")
            or attachment.get("fileUrl")
        )
        if not download_url:
            raise ValueError(f"Attachment has no downloadable URL: {attachment}")
        filename = re.sub(
            r"[/\\\\:*?\"<>|]+",
            "_",
            str(
                attachment.get("fileName")
                or attachment.get("file_name")
                or attachment.get("name")
                or "attachment"
            ),
        )
        output_dir.mkdir(parents=True, exist_ok=True)
        target_path = output_dir / filename
        response = self.http.get(
            str(download_url),
            headers={"x-acs-dingtalk-access-token": access_token},
            follow_redirects=True,
        )
        response.raise_for_status()
        target_path.write_bytes(response.content)
        return target_path

    def download_approval_file_attachment(
        self,
        access_token: str,
        attachment: dict[str, Any],
        output_dir: Path,
        *,
        user_id: str,
        union_id: str,
        fallback_space_id: str = "",
        process_instance_id: str = "",
    ) -> Path:
        """Download an approval DingDrive attachment identified by fileId.

        Args:
            access_token: DingTalk access token.
            attachment: Approval attachment metadata with fileId/file_id.
            output_dir: Local directory for downloaded files.
            user_id: DingTalk userId for approval file download authorization.
            union_id: DingTalk unionId for storage download info.
            fallback_space_id: Optional approval space ID if metadata has none.
            process_instance_id: Optional process instance ID for approval download URL.

        Returns:
            Local attachment path.

        Raises:
            ValueError: The attachment metadata cannot identify a file.
            RuntimeError: DingTalk does not return a usable download URL.
        """
        file_id = (
            attachment.get("fileId")
            or attachment.get("file_id")
            or attachment.get("dentryId")
            or attachment.get("dentry_id")
        )
        space_id = (
            attachment.get("spaceId")
            or attachment.get("space_id")
            or fallback_space_id
            or self.get_approval_attachment_space(access_token, user_id)
        )
        if not file_id:
            raise ValueError(f"Attachment has no DingTalk fileId: {attachment}")
        filename = re.sub(
            r"[/\\\\:*?\"<>|]+",
            "_",
            str(
                attachment.get("fileName")
                or attachment.get("file_name")
                or attachment.get("name")
                or file_id
            ),
        )
        output_dir.mkdir(parents=True, exist_ok=True)
        target_path = output_dir / filename
        direct_download_error = ""

        if process_instance_id:
            try:
                download_info = self.get_approval_file_download_info(
                    access_token,
                    process_instance_id=process_instance_id,
                    file_id=str(file_id),
                )
                download_uri = download_info.get("downloadUri")
                if not download_uri:
                    raise RuntimeError(
                        "DingTalk approval file download info missing downloadUri: "
                        f"{download_info}"
                    )
                response = self.http.get(str(download_uri), follow_redirects=True)
                response.raise_for_status()
                target_path.write_bytes(response.content)
                return target_path
            except (httpx.HTTPError, RuntimeError) as exc:
                direct_download_error = str(exc)

        try:
            self.authorize_approval_file_download(
                access_token,
                user_id=user_id,
                file_infos=[{"fileId": str(file_id), "spaceId": str(space_id)}],
            )
            download_info = self.get_storage_file_download_info(
                access_token,
                space_id=str(space_id),
                dentry_id=str(file_id),
                union_id=union_id,
            )
            signature_info = download_info.get("headerSignatureInfo") or {}
            resource_urls = signature_info.get("resourceUrls") or download_info.get(
                "resourceUrls"
            )
            if not resource_urls:
                raise RuntimeError(
                    "DingTalk storage download info missing resourceUrls: "
                    f"{download_info}"
                )
            headers = signature_info.get("headers")
            response = self.http.get(
                str(resource_urls[0]),
                headers=headers if isinstance(headers, dict) else None,
                follow_redirects=True,
            )
            try:
                response.raise_for_status()
            except httpx.HTTPStatusError as exc:
                raise RuntimeError(
                    f"DingTalk storage file download HTTP {response.status_code}: "
                    f"{response.text}"
                ) from exc
        except (httpx.HTTPError, RuntimeError) as exc:
            if direct_download_error:
                raise RuntimeError(
                    f"{direct_download_error}; storage fallback failed: {exc}"
                ) from exc
            raise
        target_path.write_bytes(response.content)
        return target_path


def maybe_parse_json(value: Any) -> Any:
    """Parse embedded JSON strings that DingTalk form fields often use.

    Args:
        value: Any value from a DingTalk response.

    Returns:
        Parsed JSON when value is a JSON-looking string, otherwise the original value.
    """
    if not isinstance(value, str):
        return value
    text = value.strip()
    if not text or text[0] not in "[{":
        return value
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        return value


def walk_attachment_candidates(value: Any, path: str = "$") -> list[dict[str, Any]]:
    """Find attachment-like objects anywhere in a DingTalk approval payload.

    Args:
        value: DingTalk payload or nested value.
        path: JSON-ish path used in the probe report.

    Returns:
        Attachment candidate records with the source path and value.
    """
    candidates: list[dict[str, Any]] = []
    value = maybe_parse_json(value)
    if isinstance(value, dict):
        lowered_keys = {str(key).lower() for key in value}
        has_file_identity = bool(lowered_keys & ATTACHMENT_ID_KEYS) or bool(
            lowered_keys & ATTACHMENT_LOCATION_KEYS
        )
        has_file_name = bool(lowered_keys & ATTACHMENT_NAME_KEYS)
        if has_file_identity and (has_file_name or "file_type" in lowered_keys):
            candidates.append({"path": path, "value": value})
        for key, child in value.items():
            candidates.extend(walk_attachment_candidates(child, f"{path}.{key}"))
    elif isinstance(value, list):
        for index, child in enumerate(value):
            candidates.extend(walk_attachment_candidates(child, f"{path}[{index}]"))
    return candidates


def walk_user_references(value: Any, path: str = "$") -> list[dict[str, str]]:
    """Find userId references in a DingTalk approval payload.

    Args:
        value: DingTalk payload or nested value.
        path: JSON-ish path used in the user probe report.

    Returns:
        User references with source path and nearby task ID when available.
    """
    references: list[dict[str, str]] = []
    value = maybe_parse_json(value)
    if isinstance(value, dict):
        user_id = ""
        task_id = ""
        for key, child in value.items():
            lowered = str(key).lower()
            if lowered in USER_ID_KEYS and isinstance(child, (int, str)):
                user_id = str(child)
            if lowered in TASK_ID_KEYS and isinstance(child, (int, str)):
                task_id = str(child)
        if user_id:
            reference = {"path": path, "user_id": user_id}
            if task_id:
                reference["task_id"] = task_id
            references.append(reference)
        for key, child in value.items():
            references.extend(walk_user_references(child, f"{path}.{key}"))
    elif isinstance(value, list):
        for index, child in enumerate(value):
            references.extend(walk_user_references(child, f"{path}[{index}]"))
    return references


def resolve_approval_users(
    client: DingTalkFinanceClient,
    access_token: str,
    detail: dict[str, Any],
) -> list[dict[str, Any]]:
    """Resolve approval userIds to DingTalk user details.

    Args:
        client: DingTalk finance API client.
        access_token: DingTalk access token.
        detail: DingTalk approval detail payload.

    Returns:
        Unique users with task IDs, source paths, and unionId/name when readable.
    """
    users: dict[str, dict[str, Any]] = {}
    for reference in walk_user_references(detail):
        user_id = reference["user_id"]
        user = users.setdefault(
            user_id,
            {
                "user_id": user_id,
                "task_ids": [],
                "paths": [],
            },
        )
        if reference.get("task_id") and reference["task_id"] not in user["task_ids"]:
            user["task_ids"].append(reference["task_id"])
        if reference["path"] not in user["paths"]:
            user["paths"].append(reference["path"])

    for user in users.values():
        try:
            detail_payload = client.get_user_detail(access_token, user["user_id"])
            user["union_id"] = detail_payload.get("unionid", "")
            user["name"] = detail_payload.get("name", "")
            user["title"] = detail_payload.get("title", "")
        except (httpx.HTTPError, RuntimeError) as exc:
            user["resolve_error"] = str(exc)
    return sorted(users.values(), key=lambda item: item["user_id"])


def detail_matches_terms(detail: dict[str, Any], terms: list[str]) -> bool:
    """Check whether a detail payload contains all requested search terms.

    Args:
        detail: DingTalk approval detail payload.
        terms: Non-empty terms that must appear in the serialized payload.

    Returns:
        True when all terms are present.
    """
    serialized = json.dumps(detail, ensure_ascii=False).lower()
    return all(term.lower() in serialized for term in terms if term)


def probe_approval_details(
    client: DingTalkFinanceClient,
    access_token: str,
    instance_ids: list[str],
    output_dir: Path,
    *,
    terms: list[str] | None = None,
    limit: int = 20,
) -> dict[str, Any]:
    """Fetch approval details and save raw payload plus attachment candidates.

    Args:
        client: DingTalk finance API client.
        access_token: DingTalk access token.
        instance_ids: Candidate approval instance IDs.
        output_dir: Local directory for probe JSON artifacts.
        terms: Optional terms used to keep only matching details.
        limit: Maximum number of instance details to inspect.

    Returns:
        Probe summary with matched instances and attachment candidate counts.
    """
    terms = terms or []
    matched: list[dict[str, Any]] = []
    output_dir.mkdir(parents=True, exist_ok=True)
    for instance_id in instance_ids[:limit]:
        detail = client.get_approval_detail(access_token, instance_id)
        if terms and not detail_matches_terms(detail, terms):
            continue
        candidates = walk_attachment_candidates(detail)
        safe_id = re.sub(r"[^A-Za-z0-9_.-]+", "_", instance_id)
        write_json(output_dir / f"{safe_id}_detail.json", detail)
        write_json(output_dir / f"{safe_id}_attachment_candidates.json", candidates)
        matched.append(
            {
                "instance_id": instance_id,
                "attachment_candidate_count": len(candidates),
                "detail_path": str(output_dir / f"{safe_id}_detail.json"),
                "attachment_candidates_path": str(
                    output_dir / f"{safe_id}_attachment_candidates.json"
                ),
            }
        )
    return {
        "ok": True,
        "inspected_count": min(len(instance_ids), limit),
        "matched_count": len(matched),
        "terms": terms,
        "matched": matched,
        "output_dir": str(output_dir),
    }


def redact_capability_error(
    error: Exception,
    config: ExportConfig,
    access_token: str = "",
) -> str:
    """Remove credentials from a capability probe error.

    Args:
        error: Exception raised by a DingTalk or local readiness check.
        config: Runtime configuration containing credentials to redact.
        access_token: Optional temporary access token to redact.

    Returns:
        Safe error text suitable for a local JSON report.
    """
    message = str(error)
    for secret_value in (config.app_key, config.app_secret, access_token):
        if secret_value:
            message = message.replace(secret_value, "***")
    return message


def run_capability_probe(
    config: ExportConfig,
    *,
    client: DingTalkFinanceClient | None = None,
    ledger_path: Path = DEFAULT_DEDUPE_LEDGER_PATH,
) -> dict[str, Any]:
    """Run the first read-only DingTalk finance interface test.

    The probe verifies app authentication, approval instance listing, and one
    approval detail when available. It also reports local prerequisites without
    downloading attachments or executing an approval task.

    Args:
        config: DingTalk finance runtime configuration.
        client: Optional injected client for offline interface tests.
        ledger_path: Local historical invoice ledger to validate.

    Returns:
        Structured capability report with no access tokens or app secrets.
    """
    finance_client = client or DingTalkFinanceClient(config)
    checks: dict[str, dict[str, Any]] = {}
    access_token = ""
    instance_ids: list[str] = []
    interface_ok = True

    try:
        access_token = finance_client.get_access_token()
        checks["app_access_token"] = {"status": "passed"}
    except Exception as exc:
        interface_ok = False
        checks["app_access_token"] = {
            "status": "failed",
            "error": redact_capability_error(exc, config),
        }

    if access_token:
        try:
            instance_ids = finance_client.list_approval_instance_ids(access_token)
            checks["workflow_instance_read"] = {
                "status": "passed",
                "instance_count": len(instance_ids),
                "process_code": config.process_code,
                "start_time": config.start_time.isoformat(),
                "end_time": config.end_time.isoformat(),
            }
        except Exception as exc:
            interface_ok = False
            checks["workflow_instance_read"] = {
                "status": "failed",
                "error": redact_capability_error(exc, config, access_token),
            }
    else:
        checks["workflow_instance_read"] = {
            "status": "skipped",
            "reason": "app_access_token_failed",
        }

    if instance_ids and interface_ok:
        try:
            detail = finance_client.get_approval_detail(access_token, instance_ids[0])
            checks["approval_detail_read"] = {
                "status": "passed",
                "sample_instance_id": instance_ids[0],
                "attachment_candidate_count": len(walk_attachment_candidates(detail)),
            }
        except Exception as exc:
            interface_ok = False
            checks["approval_detail_read"] = {
                "status": "failed",
                "sample_instance_id": instance_ids[0],
                "error": redact_capability_error(exc, config, access_token),
            }
    else:
        checks["approval_detail_read"] = {
            "status": "skipped",
            "reason": "no_instances_in_range"
            if interface_ok
            else "workflow_instance_read_failed",
        }

    attachment_identity_ready = bool(
        config.operator_user_id and config.operator_union_id
    )
    checks["attachment_download_identity"] = {
        "status": "ready" if attachment_identity_ready else "missing_config",
        "operator_user_id_configured": bool(config.operator_user_id),
        "operator_union_id_configured": bool(config.operator_union_id),
    }
    checks["local_ocr"] = {
        "status": "ready" if config.local_ocr_url else "missing_config",
        "endpoint_configured": bool(config.local_ocr_url),
    }
    reviewer_ready = bool(
        config.finance_reviewer_user_id or config.finance_reviewer_union_id
    )
    checks["finance_reviewer"] = {
        "status": "ready" if reviewer_ready else "missing_config",
        "reviewer_name": config.finance_reviewer_name,
        "reviewer_user_id_configured": bool(config.finance_reviewer_user_id),
        "reviewer_union_id_configured": bool(config.finance_reviewer_union_id),
    }
    if not ledger_path.exists():
        ledger_ready = False
        checks["duplicate_ledger"] = {
            "status": "missing_config",
            "path": str(ledger_path),
            "record_count": 0,
        }
    else:
        try:
            ledger_records = load_duplicate_ledger(ledger_path)
            checks["duplicate_ledger"] = {
                "status": "ready",
                "path": str(ledger_path),
                "record_count": len(ledger_records),
            }
            ledger_ready = True
        except Exception as exc:
            ledger_ready = False
            checks["duplicate_ledger"] = {
                "status": "invalid",
                "path": str(ledger_path),
                "error": redact_capability_error(exc, config, access_token),
            }
    checks["stream_events"] = {
        "status": "enabled" if config.finance_stream_enabled else "disabled"
    }
    checks["approval_mutation"] = {
        "status": "not_tested",
        "reason": "capabilities_is_read_only",
    }

    return {
        "ok": interface_ok,
        "mode": "read_only",
        "mutating_actions_tested": False,
        "ready_for_initial_review": (
            interface_ok
            and attachment_identity_ready
            and bool(config.local_ocr_url)
            and reviewer_ready
            and ledger_ready
        ),
        "checks": checks,
    }


def extract_approval_row(detail: dict[str, Any], index: int) -> dict[str, Any]:
    """Extract finance summary fields and attachment metadata from a detail payload.

    Args:
        detail: DingTalk approval detail payload.
        index: One-based row index in the export.

    Returns:
        Normalized summary row data.
    """
    form_values = (
        detail.get("formComponentValues") or detail.get("form_component_values") or []
    )
    flat_values: dict[str, Any] = {}
    attachments: list[dict[str, Any]] = []

    for item in form_values:
        if not isinstance(item, dict):
            continue
        name = str(
            item.get("name") or item.get("componentName") or item.get("label") or ""
        )
        value = item.get("value")
        flat_values[name] = value
        parsed_value = maybe_parse_json(value)
        if isinstance(parsed_value, list):
            for candidate in parsed_value:
                if isinstance(candidate, dict) and (
                    candidate.get("fileName")
                    or candidate.get("file_name")
                    or candidate.get("fileId")
                    or candidate.get("file_id")
                    or candidate.get("downloadUrl")
                    or candidate.get("url")
                ):
                    attachments.append(candidate)
        elif isinstance(parsed_value, dict) and (
            parsed_value.get("fileName")
            or parsed_value.get("file_name")
            or parsed_value.get("fileId")
            or parsed_value.get("file_id")
            or parsed_value.get("downloadUrl")
            or parsed_value.get("url")
        ):
            attachments.append(parsed_value)

    seen_attachments = {
        json.dumps(attachment, ensure_ascii=False, sort_keys=True)
        for attachment in attachments
    }
    for candidate in walk_attachment_candidates(detail):
        attachment = candidate["value"]
        key = json.dumps(attachment, ensure_ascii=False, sort_keys=True)
        if key not in seen_attachments:
            attachments.append(attachment)
            seen_attachments.add(key)

    attachment_sheet = f"附件{index}"
    return {
        "序号": index,
        "部门": detail.get("originatorDeptName")
        or detail.get("originator_dept_name")
        or flat_values.get("部门")
        or detail.get("department")
        or "",
        "姓名": detail.get("originatorUserName")
        or flat_values.get("报账人员")
        or flat_values.get("姓名")
        or flat_values.get("申请人")
        or detail.get("originator")
        or "",
        "报销项目": flat_values.get("报销项目")
        or flat_values.get("付款项目")
        or flat_values.get("付款事由")
        or flat_values.get("费用项目")
        or flat_values.get("项目")
        or "",
        "日期": flat_values.get("日期")
        or flat_values.get("报销日期")
        or flat_values.get("付款日期（费用发生日期）")
        or detail.get("createTime")
        or detail.get("create_time")
        or "",
        "金额": flat_values.get("金额")
        or flat_values.get("报销金额")
        or flat_values.get("报销总金额（元）")
        or flat_values.get("付款总金额")
        or flat_values.get("总金额")
        or "",
        "附件 sheet": attachment_sheet,
        "备注说明": flat_values.get("备注说明")
        or flat_values.get("备注")
        or flat_values.get("说明")
        or "",
        "attachments": attachments,
        "instance_id": detail.get("processInstanceId")
        or detail.get("process_instance_id")
        or detail.get("id")
        or "",
    }


def build_excel_package(rows: list[dict[str, Any]], workbook_path: Path) -> Path:
    """Build the finance Excel package with summary and attachment sheets.

    Args:
        rows: Normalized approval rows from DingTalk details.
        workbook_path: Target workbook path.

    Returns:
        Saved workbook path.
    """
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "目录"
    summary.append(SUMMARY_COLUMNS)
    for cell in summary[1]:
        cell.font = Font(bold=True)

    for row_index, row in enumerate(rows, start=2):
        summary.append([row.get(column, "") for column in SUMMARY_COLUMNS])
        sheet_name = str(row["附件 sheet"])
        summary.cell(row=row_index, column=7).hyperlink = f"#{sheet_name}!A1"
        summary.cell(row=row_index, column=7).style = "Hyperlink"
        sheet = workbook.create_sheet(sheet_name)
        sheet["A1"] = "返回目录"
        sheet["A1"].hyperlink = "#目录!A1"
        sheet["A1"].style = "Hyperlink"
        sheet.append(["文件名", "本地路径", "下载状态", "错误", "预览"])
        for attachment_index, attachment in enumerate(
            row.get("attachments", []), start=3
        ):
            filename = str(
                attachment.get("fileName")
                or attachment.get("file_name")
                or attachment.get("name")
                or ""
            )
            local_path = attachment.get("local_path")
            sheet.cell(row=attachment_index, column=1, value=filename)
            sheet.cell(
                row=attachment_index,
                column=3,
                value="已下载" if local_path else "未下载",
            )
            if attachment.get("download_error"):
                sheet.cell(
                    row=attachment_index,
                    column=4,
                    value=str(attachment.get("download_error")),
                )
            if local_path:
                local_path = Path(local_path)
                sheet.cell(row=attachment_index, column=2, value=str(local_path))
                sheet.cell(
                    row=attachment_index, column=2
                ).hyperlink = local_path.resolve().as_uri()
                sheet.cell(row=attachment_index, column=2).style = "Hyperlink"
                if local_path.suffix.lower() in IMAGE_SUFFIXES and local_path.exists():
                    image = WorkbookImage(str(local_path))
                    image.width = min(image.width, 320)
                    image.height = min(image.height, 240)
                    sheet.add_image(image, f"E{attachment_index}")
                    sheet.row_dimensions[attachment_index].height = 140
        for column in range(1, 6):
            sheet.column_dimensions[get_column_letter(column)].width = 32

    for column in range(1, len(SUMMARY_COLUMNS) + 1):
        summary.column_dimensions[get_column_letter(column)].width = 18
    workbook.save(workbook_path)
    return workbook_path


def download_row_attachments(
    client: DingTalkFinanceClient,
    access_token: str,
    row: dict[str, Any],
    attachment_root: Path,
    config: ExportConfig,
    process_instance_id: str,
    *,
    fallback_space_id: str = "",
) -> str:
    """Download every attachment for one approval row.

    Args:
        client: DingTalk finance API client.
        access_token: DingTalk access token.
        row: Normalized approval row whose attachments will be updated in place.
        attachment_root: Root directory for downloaded attachments.
        config: Runtime bot configuration.
        process_instance_id: DingTalk approval instance ID.
        fallback_space_id: Reused approval DingDrive space ID when available.

    Returns:
        Latest known fallback space ID.
    """
    instance_attachment_dir = attachment_root / str(process_instance_id)
    for attachment in row.get("attachments", []):
        try:
            if (
                attachment.get("downloadUrl")
                or attachment.get("url")
                or attachment.get("fileUrl")
            ):
                attachment["local_path"] = str(
                    client.download_approval_attachment(
                        access_token,
                        attachment,
                        instance_attachment_dir,
                    )
                )
            elif config.operator_user_id and config.operator_union_id:
                attachment["local_path"] = str(
                    client.download_approval_file_attachment(
                        access_token,
                        attachment,
                        instance_attachment_dir,
                        user_id=config.operator_user_id,
                        union_id=config.operator_union_id,
                        fallback_space_id=fallback_space_id,
                        process_instance_id=process_instance_id,
                    )
                )
                fallback_space_id = str(
                    attachment.get("spaceId")
                    or attachment.get("space_id")
                    or fallback_space_id
                )
            else:
                raise ValueError(
                    "Attachment needs DINGTALK_OPERATOR_USER_ID and "
                    "DINGTALK_OPERATOR_UNION_ID for DingDrive download"
                )
            if config.local_ocr_url and not read_attachment_ocr_text(attachment):
                try:
                    attachment["ocr_text"] = recognize_file_with_local_ocr(
                        config,
                        Path(str(attachment["local_path"])),
                    )
                    attachment["ocr_source"] = "local_ocr_service"
                except (OSError, RuntimeError, httpx.HTTPError) as exc:
                    attachment["ocr_error"] = str(exc)
        except (httpx.HTTPError, RuntimeError, ValueError) as exc:
            attachment["download_error"] = str(exc)
    return fallback_space_id


def extract_local_ocr_text(payload: Any) -> str:
    """Extract OCR text from supported local service response shapes.

    Args:
        payload: JSON payload returned by the local OCR service.

    Returns:
        OCR text.

    Raises:
        RuntimeError: The payload does not contain text.
    """
    if isinstance(payload, str) and payload.strip():
        return payload.strip()
    if isinstance(payload, dict):
        for key in ("text", "ocr_text", "ocrText", "recognized_text"):
            value = payload.get(key)
            if isinstance(value, str) and value.strip():
                return value.strip()
        for key in ("result", "data"):
            value = payload.get(key)
            if isinstance(value, (dict, str)):
                try:
                    return extract_local_ocr_text(value)
                except RuntimeError:
                    pass
    raise RuntimeError("Local OCR response did not contain text")


def recognize_file_with_local_ocr(config: ExportConfig, path: Path) -> str:
    """Send one local attachment to the controlled OCR service.

    Args:
        config: Runtime bot configuration with ``local_ocr_url``.
        path: Local attachment file.

    Returns:
        OCR text from the local service.

    Raises:
        RuntimeError: The OCR service is not configured or returns no text.
        OSError: The local file cannot be opened.
        httpx.HTTPError: The local OCR HTTP request fails.
    """
    if not config.local_ocr_url:
        raise RuntimeError("DINGTALK_FINANCE_OCR_URL is not configured")
    content_type = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    with path.open("rb") as file:
        response = httpx.post(
            config.local_ocr_url,
            files={"file": (path.name, file, content_type)},
            timeout=config.local_ocr_timeout_seconds,
        )
    response.raise_for_status()
    try:
        payload: Any = response.json()
    except json.JSONDecodeError:
        payload = response.text
    return extract_local_ocr_text(payload)


def parse_money_amount(value: Any) -> Decimal | None:
    """Parse a money-like value into a two-decimal Decimal.

    Args:
        value: Raw amount string or number.

    Returns:
        Decimal amount rounded to cents, or None when parsing fails.
    """
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    match = re.search(r"-?\d[\d,]*(?:\.\d{1,4})?", text)
    if not match:
        return None
    try:
        return Decimal(match.group(0).replace(",", "")).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return None


def read_attachment_ocr_text(attachment: dict[str, Any]) -> str:
    """Read OCR text from attachment metadata or a local sidecar file.

    Args:
        attachment: Attachment metadata enriched during download.

    Returns:
        OCR text, or an empty string when none is available.
    """
    for key in (
        "ocr_text",
        "ocrText",
        "recognized_text",
        "recognizedText",
        "text",
    ):
        value = attachment.get(key)
        if isinstance(value, str) and value.strip():
            return value.strip()

    local_path_value = attachment.get("local_path")
    if not local_path_value:
        return ""
    local_path = Path(str(local_path_value))
    sidecar_paths = [
        local_path.with_suffix(local_path.suffix + ".txt"),
        local_path.with_suffix(".ocr.txt"),
        local_path.with_suffix(".txt"),
    ]
    for sidecar_path in sidecar_paths:
        if sidecar_path.is_file():
            return sidecar_path.read_text(encoding="utf-8").strip()
    return ""


def extract_invoice_amount_candidates(ocr_text: str) -> list[Decimal]:
    """Extract likely invoice total amounts from OCR text.

    Args:
        ocr_text: Recognized invoice text.

    Returns:
        Distinct candidate amounts in appearance order.
    """
    candidates: list[Decimal] = []
    seen: set[Decimal] = set()
    total_patterns = [
        r"(?:价税合计|价税合计小写|小写|合计金额|金额合计|总金额|付款金额)"
        r"[^0-9￥¥-]{0,24}[￥¥]?\s*(-?\d[\d,]*(?:\.\d{1,4})?)",
    ]
    fallback_patterns = [
        r"[￥¥]\s*(-?\d[\d,]*(?:\.\d{1,4})?)",
    ]
    for pattern in total_patterns:
        for match in re.finditer(pattern, ocr_text, flags=re.IGNORECASE):
            amount = parse_money_amount(match.group(1))
            if amount is not None and amount not in seen:
                candidates.append(amount)
                seen.add(amount)
    if candidates:
        return candidates

    for pattern in fallback_patterns:
        for match in re.finditer(pattern, ocr_text, flags=re.IGNORECASE):
            amount = parse_money_amount(match.group(1))
            if amount is not None and amount not in seen:
                candidates.append(amount)
                seen.add(amount)
    return candidates


def build_initial_review_report(
    row: dict[str, Any],
    ledger_records: list[dict[str, Any]],
) -> dict[str, Any]:
    """Build a non-mutating initial review decision report.

    Args:
        row: Normalized approval row after attachment download attempts.
        ledger_records: Local invoice ledger records for duplicate hash checks.

    Returns:
        Review report showing whether the bot can safely auto-approve later.
    """
    ledger_by_hash: dict[str, dict[str, Any]] = {}
    for record in ledger_records:
        digest = str(
            record.get("sha256")
            or record.get("file_sha256")
            or record.get("attachment_sha256")
            or ""
        )
        if digest:
            ledger_by_hash[digest] = record

    issues: list[dict[str, Any]] = []
    attachment_reports: list[dict[str, Any]] = []
    downloaded_count = 0
    ocr_text_count = 0
    ocr_error_count = 0
    invoice_amount_candidates: list[Decimal] = []
    invoice_amount_best_by_attachment: list[Decimal] = []

    for index, attachment in enumerate(row.get("attachments", []), start=1):
        filename = str(
            attachment.get("fileName")
            or attachment.get("file_name")
            or attachment.get("name")
            or attachment.get("fileId")
            or attachment.get("file_id")
            or f"attachment-{index}"
        )
        attachment_report: dict[str, Any] = {
            "index": index,
            "file_name": filename,
            "download_status": "not_downloaded",
        }
        if attachment.get("download_error"):
            error_text = str(attachment["download_error"])
            attachment_report["download_error"] = error_text
            issues.append(
                {
                    "code": "attachment_download_failed",
                    "severity": "blocking",
                    "file_name": filename,
                    "message": error_text,
                }
            )
            lowered_error = error_text.lower()
            if (
                "orgauthlevelnotenough" in lowered_error
                or "nopermission" in lowered_error
                or "premium" in lowered_error
                or "comment" in lowered_error
            ):
                issues.append(
                    {
                        "code": "comment_attachment_download_limited",
                        "severity": "blocking",
                        "file_name": filename,
                        "message": "评论区补传附件元数据可读，但当前普通接口下载受限。",
                    }
                )
        if attachment.get("ocr_error"):
            ocr_error_count += 1
            error_text = str(attachment["ocr_error"])
            attachment_report["ocr_error"] = error_text
            issues.append(
                {
                    "code": "invoice_ocr_failed",
                    "severity": "blocking",
                    "file_name": filename,
                    "message": f"本地 OCR 服务识别失败：{error_text}",
                }
            )

        local_path_value = attachment.get("local_path")
        if local_path_value:
            local_path = Path(str(local_path_value))
            attachment_report["local_path"] = str(local_path)
            if local_path.exists():
                downloaded_count += 1
                digest = file_sha256(local_path)
                attachment_report["sha256"] = digest
                attachment_report["download_status"] = "downloaded"
                duplicate_record = ledger_by_hash.get(digest)
                if duplicate_record:
                    attachment_report["duplicate_record"] = duplicate_record
                    issues.append(
                        {
                            "code": "duplicate_attachment_hash",
                            "severity": "blocking",
                            "file_name": filename,
                            "sha256": digest,
                            "message": "本地发票台账中已存在相同文件哈希。",
                        }
                    )
                ocr_text = read_attachment_ocr_text(attachment)
                if ocr_text:
                    ocr_text_count += 1
                    candidates = extract_invoice_amount_candidates(ocr_text)
                    attachment_report["ocr_text_available"] = True
                    attachment_report["ocr_text_length"] = len(ocr_text)
                    attachment_report["invoice_amount_candidates"] = [
                        str(amount) for amount in candidates
                    ]
                    if candidates:
                        best_amount = max(candidates)
                        attachment_report["invoice_amount"] = str(best_amount)
                        if len(candidates) > 1:
                            attachment_report["invoice_amount_total"] = str(
                                sum(candidates, Decimal("0.00")).quantize(
                                    Decimal("0.01")
                                )
                            )
                        invoice_amount_candidates.extend(candidates)
                        invoice_amount_best_by_attachment.append(best_amount)
            else:
                attachment_report["download_status"] = "missing_local_file"
                issues.append(
                    {
                        "code": "downloaded_file_missing",
                        "severity": "blocking",
                        "file_name": filename,
                        "message": f"本地附件路径不存在：{local_path}",
                    }
                )
        attachment_reports.append(attachment_report)

    reimbursement_amount = parse_money_amount(row.get("金额"))
    amount_verification: dict[str, Any] = {
        "status": "not_available",
        "reimbursement_amount": str(reimbursement_amount)
        if reimbursement_amount is not None
        else "",
        "ocr_text_attachment_count": ocr_text_count,
        "candidate_amounts": [str(amount) for amount in invoice_amount_candidates],
    }
    if invoice_amount_best_by_attachment:
        best_total = sum(invoice_amount_best_by_attachment, Decimal("0.00")).quantize(
            Decimal("0.01")
        )
        amount_verification["best_attachment_total"] = str(best_total)
    if invoice_amount_candidates:
        candidate_total = sum(invoice_amount_candidates, Decimal("0.00")).quantize(
            Decimal("0.01")
        )
        amount_verification["candidate_total"] = str(candidate_total)

    if not attachment_reports:
        issues.append(
            {
                "code": "missing_invoice_attachment",
                "severity": "blocking",
                "message": "审批详情里没有可识别的发票或报销附件。",
            }
        )
    if reimbursement_amount is None:
        issues.append(
            {
                "code": "missing_reimbursement_amount",
                "severity": "blocking",
                "message": "审批表单里没有识别到报销金额。",
            }
        )
    if downloaded_count:
        if ocr_error_count and not ocr_text_count:
            amount_verification["status"] = "ocr_failed"
        elif not ocr_text_count:
            amount_verification["status"] = "waiting_for_ocr"
            issues.append(
                {
                    "code": "invoice_ocr_not_configured",
                    "severity": "pending_automation",
                    "message": "附件已下载，但还未接入发票 OCR。",
                }
            )
            issues.append(
                {
                    "code": "amount_unverified",
                    "severity": "pending_automation",
                    "message": "发票金额尚未与报销金额自动核对。",
                }
            )
        elif reimbursement_amount is not None:
            best_total = (
                sum(invoice_amount_best_by_attachment, Decimal("0.00")).quantize(
                    Decimal("0.01")
                )
                if invoice_amount_best_by_attachment
                else None
            )
            matched_amounts = [
                amount
                for amount in invoice_amount_candidates
                if amount == reimbursement_amount
            ]
            if best_total == reimbursement_amount:
                matched_amounts.append(best_total)
            candidate_total = (
                sum(invoice_amount_candidates, Decimal("0.00")).quantize(
                    Decimal("0.01")
                )
                if invoice_amount_candidates
                else None
            )
            if candidate_total == reimbursement_amount:
                matched_amounts.append(candidate_total)
            if matched_amounts:
                matched_unique: list[Decimal] = []
                for amount in matched_amounts:
                    if amount not in matched_unique:
                        matched_unique.append(amount)
                amount_verification["status"] = "matched"
                amount_verification["matched_amounts"] = [
                    str(amount) for amount in matched_unique
                ]
            elif invoice_amount_candidates:
                amount_verification["status"] = "mismatch"
                issues.append(
                    {
                        "code": "invoice_amount_mismatch",
                        "severity": "blocking",
                        "message": ("OCR 识别到的发票金额与审批报销金额不一致。"),
                        "reimbursement_amount": str(reimbursement_amount),
                        "candidate_amounts": [
                            str(amount) for amount in invoice_amount_candidates
                        ],
                    }
                )
            else:
                amount_verification["status"] = "amount_not_found"
                issues.append(
                    {
                        "code": "invoice_amount_not_found",
                        "severity": "blocking",
                        "message": "OCR 文本中未识别到可用于核对的发票金额。",
                    }
                )

    has_blocking_issue = any(issue["severity"] == "blocking" for issue in issues)
    issue_codes = {issue["code"] for issue in issues}
    if "comment_attachment_download_limited" in issue_codes:
        recommendation = (
            "评论区补传附件已经能被识别，但普通审批附件接口下载受限；"
            "需要开通 OA 高级版接口或让员工改从表单附件/可下载来源补传。"
        )
    elif has_blocking_issue:
        recommendation = "先处理阻塞项，不要自动通过当前审批。"
    elif amount_verification.get("status") == "matched":
        recommendation = "发票金额已与报销金额匹配，可进入财务复审。"
    elif downloaded_count:
        recommendation = "附件已下载并完成本地哈希查重，下一步接入 OCR 金额校验。"
    else:
        recommendation = "未获得可自动校验的附件，保持人工复审。"

    status = "needs_manual_review" if has_blocking_issue else "ready_for_ocr"
    if not has_blocking_issue and amount_verification.get("status") == "matched":
        status = "ready_for_finance_review"

    return {
        "ok": True,
        "instance_id": row.get("instance_id", ""),
        "employee": row.get("姓名", ""),
        "department": row.get("部门", ""),
        "project": row.get("报销项目", ""),
        "reimbursement_amount": row.get("金额", ""),
        "status": status,
        "can_auto_approve": False,
        "recommendation": recommendation,
        "amount_verification": amount_verification,
        "attachments": attachment_reports,
        "issues": issues,
    }


def compact_text(value: Any) -> str:
    """Normalize text for loose DingTalk payload matching.

    Args:
        value: Value to normalize.

    Returns:
        Lowercase string without whitespace.
    """
    return re.sub(r"\s+", "", str(value or "")).lower()


def find_pending_tasks_for_user(
    detail: dict[str, Any],
    user_id: str,
) -> list[dict[str, str]]:
    """Find active approval tasks assigned to one DingTalk userId.

    Args:
        detail: DingTalk approval detail payload.
        user_id: DingTalk userId of the finance reviewer.

    Returns:
        Pending task records with task_id and path.
    """
    target_user_id = str(user_id)
    pending_statuses = {
        "",
        "new",
        "pending",
        "running",
        "todo",
        "wait",
        "waiting",
        "待处理",
        "审批中",
    }
    pending_tasks: list[dict[str, str]] = []

    def visit(value: Any, path: str = "$") -> None:
        value = maybe_parse_json(value)
        if isinstance(value, dict):
            lowered = {str(key).lower(): child for key, child in value.items()}
            task_id = lowered.get("taskid") or lowered.get("task_id")
            current_user_id = (
                lowered.get("userid")
                or lowered.get("user_id")
                or lowered.get("actioneruserid")
                or lowered.get("actioner_user_id")
            )
            if task_id and str(current_user_id or "") == target_user_id:
                raw_status = (
                    lowered.get("task_status")
                    or lowered.get("taskstatus")
                    or lowered.get("status")
                    or lowered.get("result")
                    or ""
                )
                status = compact_text(raw_status)
                if status in pending_statuses:
                    pending_tasks.append(
                        {
                            "task_id": str(task_id),
                            "user_id": target_user_id,
                            "status": str(raw_status or ""),
                            "path": path,
                        }
                    )
            for key, child in value.items():
                visit(child, f"{path}.{key}")
        elif isinstance(value, list):
            for index, child in enumerate(value):
                visit(child, f"{path}[{index}]")

    visit(detail)
    return pending_tasks


def approval_has_general_manager_agreed(
    detail: dict[str, Any],
    keywords: tuple[str, ...],
) -> bool:
    """Check whether Yang's previous approval is present in the workflow records.

    Args:
        detail: DingTalk approval detail payload.
        keywords: Names or aliases that identify Yang's approval node.

    Returns:
        True when a matching operation record looks approved.
    """
    normalized_keywords = [compact_text(keyword) for keyword in keywords if keyword]
    if not normalized_keywords:
        return True
    approval_words = ("agree", "agreed", "同意", "通过", "已同意")
    records = (
        detail.get("operationRecords")
        or detail.get("operation_records")
        or detail.get("operation_record")
        or []
    )
    if not isinstance(records, list):
        records = []
    for record in records:
        if not isinstance(record, dict):
            continue
        serialized = compact_text(json.dumps(record, ensure_ascii=False))
        if not any(keyword in serialized for keyword in normalized_keywords):
            continue
        if any(word in serialized for word in approval_words):
            return True
    return False


def notify_initial_review_report(
    config: ExportConfig,
    client: DingTalkFinanceClient,
    access_token: str,
    report: dict[str, Any],
    marker_path: Path,
) -> dict[str, Any]:
    """Send one DingTalk work notification for a generated review report.

    Args:
        config: Runtime bot configuration.
        client: DingTalk API client.
        access_token: Reused DingTalk app access token.
        report: Initial review report.
        marker_path: Local notification marker used to avoid duplicate sends.

    Returns:
        Notification status for the scan summary.
    """
    status: dict[str, Any] = {
        "enabled": config.finance_notify_enabled,
        "sent": False,
        "marker_path": str(marker_path),
    }
    if not config.finance_notify_enabled:
        status["reason"] = "notification_disabled"
        return status
    if config.finance_notify_on == "none":
        status["reason"] = "notification_policy_none"
        return status
    if marker_path.exists():
        status["reason"] = "notification_already_sent"
        status["already_sent"] = True
        return status
    issues = [issue for issue in report.get("issues", []) if isinstance(issue, dict)]
    if config.finance_notify_on == "issues" and not issues:
        status["reason"] = "no_issues"
        return status
    user_ids = config.finance_notify_user_ids or (
        (config.finance_reviewer_user_id,) if config.finance_reviewer_user_id else ()
    )
    if not user_ids:
        status["reason"] = "missing_notify_user_ids"
        return status
    if not config.agent_id:
        status["reason"] = "missing_agent_id"
        return status

    issue_codes = [
        str(issue.get("code", "")) for issue in issues[:6] if str(issue.get("code", ""))
    ]
    issue_text = "、".join(issue_codes) if issue_codes else "无阻塞问题"
    report_path = str(report.get("report_path", ""))
    report_href = ""
    if config.finance_tool_base_url and report_path:
        try:
            report_href = config.finance_tool_base_url + output_file_href(
                config,
                report_path,
            )
        except ValueError:
            report_href = ""
    dashboard_href = config.finance_tool_base_url or ""
    content_lines = [
        "财务报销机器人初审报告",
        f"员工：{report.get('employee', '') or '-'}",
        f"项目：{report.get('project', '') or '-'}",
        f"金额：{report.get('reimbursement_amount', '') or '-'}",
        f"状态：{REVIEW_STATUS_LABELS.get(str(report.get('status', '')), str(report.get('status', '')))}",
        f"问题：{issue_text}",
        f"建议：{report.get('recommendation', '') or '-'}",
    ]
    if report_href:
        content_lines.append(f"报告：{report_href}")
    elif dashboard_href:
        content_lines.append(f"工具：{dashboard_href}")
    try:
        response = client.send_work_notification(
            access_token,
            user_ids=user_ids,
            content="\n".join(content_lines),
            agent_id=config.agent_id,
        )
    except Exception as exc:
        status["reason"] = "notification_send_failed"
        status["error"] = str(exc)
        return status
    status.update(
        {
            "sent": True,
            "sent_at": datetime.now().isoformat(timespec="seconds"),
            "user_ids": list(user_ids),
            "response": response,
        }
    )
    write_json(
        marker_path,
        {
            "sent_at": status["sent_at"],
            "instance_id": report.get("instance_id", ""),
            "report_path": report_path,
            "notify_on": config.finance_notify_on,
            "user_ids": list(user_ids),
            "response": response,
        },
    )
    return status


def run_reviewer_queue_once(
    config: ExportConfig,
    *,
    reviewer_user_id: str,
    reviewer_union_id: str = "",
    ledger_path: Path = DEFAULT_DEDUPE_LEDGER_PATH,
    terms: list[str] | None = None,
    limit: int = 20,
    skip_reviewed: bool = False,
    client: DingTalkFinanceClient | None = None,
) -> dict[str, Any]:
    """Scan approvals currently waiting for the finance reviewer.

    Args:
        config: Runtime bot configuration.
        reviewer_user_id: DingTalk userId of the reviewer account.
        reviewer_union_id: DingTalk unionId for diagnostics.
        ledger_path: Local duplicate invoice ledger path.
        terms: Optional detail filter terms.
        limit: Maximum approval details to inspect.
        skip_reviewed: Skip approvals with an existing local review report.
        client: Optional injected DingTalk client.

    Returns:
        Queue scan summary saved under the configured output directory.
    """
    terms = terms or []
    current_client = client or DingTalkFinanceClient(config)
    token = current_client.get_access_token()
    instance_ids = current_client.list_approval_instance_ids(token)
    reviewed: list[dict[str, Any]] = []
    skipped: list[dict[str, Any]] = []

    for instance_id in instance_ids[:limit]:
        safe_id = re.sub(r"[^A-Za-z0-9_.-]+", "_", instance_id)
        report_path = config.output_dir / f"review_{safe_id}.json"
        notification_marker_path = config.output_dir / f"notify_{safe_id}.json"
        if skip_reviewed and report_path.exists():
            skipped.append(
                {
                    "instance_id": instance_id,
                    "reason": "review_report_exists",
                    "report_path": str(report_path),
                }
            )
            continue
        detail = current_client.get_approval_detail(token, instance_id)
        if terms and not detail_matches_terms(detail, terms):
            skipped.append({"instance_id": instance_id, "reason": "terms_not_matched"})
            continue
        pending_tasks = find_pending_tasks_for_user(detail, reviewer_user_id)
        if not pending_tasks:
            skipped.append(
                {"instance_id": instance_id, "reason": "not_waiting_for_reviewer"}
            )
            continue
        if not approval_has_general_manager_agreed(
            detail,
            config.general_manager_keywords,
        ):
            skipped.append(
                {"instance_id": instance_id, "reason": "general_manager_not_approved"}
            )
            continue
        report = run_initial_review(
            config,
            instance_id,
            ledger_path,
            client=current_client,
            access_token=token,
            detail=detail,
        )
        notification_status = notify_initial_review_report(
            config,
            current_client,
            token,
            report,
            notification_marker_path,
        )
        reviewed.append(
            {
                "instance_id": instance_id,
                "status": report.get("status", ""),
                "can_auto_approve": report.get("can_auto_approve", False),
                "task_ids": [task["task_id"] for task in pending_tasks],
                "issue_codes": [
                    str(issue.get("code", ""))
                    for issue in report.get("issues", [])
                    if issue.get("code")
                ],
                "report_path": report.get("report_path", ""),
                "notification": notification_status,
            }
        )

    summary = {
        "ok": True,
        "mode": "reviewer_queue",
        "reviewer_name": config.finance_reviewer_name,
        "reviewer_user_id": reviewer_user_id,
        "reviewer_union_id": reviewer_union_id,
        "inspected_count": min(len(instance_ids), limit),
        "reviewed_count": len(reviewed),
        "skipped_count": len(skipped),
        "terms": terms,
        "reviewed": reviewed,
        "skipped": skipped,
    }
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    summary_path = config.output_dir / f"reviewer_queue_{stamp}.json"
    summary["summary_path"] = str(summary_path)
    write_json(summary_path, summary)
    return summary


def run_export(config: ExportConfig, instance_id: str | None = None) -> Path:
    """Run token, approval detail, attachment download, and workbook export.

    Args:
        config: Runtime bot configuration.
        instance_id: Optional single DingTalk approval instance ID.

    Returns:
        Generated workbook path.
    """
    client = DingTalkFinanceClient(config)
    token = client.get_access_token()
    instance_ids = (
        [instance_id] if instance_id else client.list_approval_instance_ids(token)
    )
    rows: list[dict[str, Any]] = []
    attachment_root = config.output_dir / "attachments"
    fallback_space_id = ""
    for index, current_instance_id in enumerate(instance_ids, start=1):
        detail = client.get_approval_detail(token, current_instance_id)
        row = extract_approval_row(detail, index)
        row["instance_id"] = row.get("instance_id") or current_instance_id
        fallback_space_id = download_row_attachments(
            client,
            token,
            row,
            attachment_root,
            config,
            current_instance_id,
            fallback_space_id=fallback_space_id,
        )
        rows.append(row)
    return build_excel_package(rows, config.output_dir / "dingtalk_finance_export.xlsx")


def run_initial_review(
    config: ExportConfig,
    instance_id: str,
    ledger_path: Path = DEFAULT_DEDUPE_LEDGER_PATH,
    *,
    client: DingTalkFinanceClient | None = None,
    access_token: str = "",
    detail: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Run a non-destructive initial review for one approval instance.

    Args:
        config: Runtime bot configuration.
        instance_id: DingTalk approval instance ID.
        ledger_path: Local duplicate invoice ledger path.
        client: Optional injected DingTalk client.
        access_token: Optional reused DingTalk access token.
        detail: Optional already-fetched approval detail payload.

    Returns:
        Review report saved under the configured output directory.
    """
    current_client = client or DingTalkFinanceClient(config)
    token = access_token or current_client.get_access_token()
    detail = detail or current_client.get_approval_detail(token, instance_id)
    row = extract_approval_row(detail, 1)
    row["instance_id"] = row.get("instance_id") or instance_id
    download_row_attachments(
        current_client,
        token,
        row,
        config.output_dir / "attachments",
        config,
        instance_id,
    )
    report = build_initial_review_report(row, load_duplicate_ledger(ledger_path))
    safe_id = re.sub(r"[^A-Za-z0-9_.-]+", "_", instance_id)
    report_path = config.output_dir / f"review_{safe_id}.json"
    report["report_path"] = str(report_path)
    write_json(report_path, report)
    return report


def run_poll_once(
    config: ExportConfig,
    *,
    ledger_path: Path = DEFAULT_DEDUPE_LEDGER_PATH,
    terms: list[str] | None = None,
    limit: int = 20,
    skip_reviewed: bool = False,
    client: DingTalkFinanceClient | None = None,
) -> dict[str, Any]:
    """Scan recent reimbursement approvals and run local initial reviews.

    Args:
        config: Runtime bot configuration.
        ledger_path: Local duplicate invoice ledger path.
        terms: Optional detail filter terms.
        limit: Maximum approval details to inspect.
        skip_reviewed: Skip approvals with an existing local review report.
        client: Optional injected DingTalk client.

    Returns:
        Poll summary saved under the configured output directory.
    """
    terms = terms or []
    current_client = client or DingTalkFinanceClient(config)
    token = current_client.get_access_token()
    instance_ids = current_client.list_approval_instance_ids(token)
    reviewed: list[dict[str, Any]] = []
    skipped: list[dict[str, str]] = []

    for instance_id in instance_ids[:limit]:
        safe_id = re.sub(r"[^A-Za-z0-9_.-]+", "_", instance_id)
        report_path = config.output_dir / f"review_{safe_id}.json"
        notification_marker_path = config.output_dir / f"notify_{safe_id}.json"
        if skip_reviewed and report_path.exists():
            skipped.append(
                {
                    "instance_id": instance_id,
                    "reason": "review_report_exists",
                    "report_path": str(report_path),
                }
            )
            continue
        detail = current_client.get_approval_detail(token, instance_id)
        if terms and not detail_matches_terms(detail, terms):
            skipped.append({"instance_id": instance_id, "reason": "terms_not_matched"})
            continue
        report = run_initial_review(
            config,
            instance_id,
            ledger_path,
            client=current_client,
            access_token=token,
            detail=detail,
        )
        notification_status = notify_initial_review_report(
            config,
            current_client,
            token,
            report,
            notification_marker_path,
        )
        reviewed.append(
            {
                "instance_id": instance_id,
                "status": report.get("status", ""),
                "can_auto_approve": report.get("can_auto_approve", False),
                "issue_codes": [
                    str(issue.get("code", ""))
                    for issue in report.get("issues", [])
                    if issue.get("code")
                ],
                "report_path": report.get("report_path", ""),
                "notification": notification_status,
            }
        )

    summary = {
        "ok": True,
        "mode": "poll_once",
        "inspected_count": min(len(instance_ids), limit),
        "reviewed_count": len(reviewed),
        "skipped_count": len(skipped),
        "terms": terms,
        "reviewed": reviewed,
        "skipped": skipped,
    }
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    summary_path = config.output_dir / f"poll_{stamp}.json"
    summary["summary_path"] = str(summary_path)
    write_json(summary_path, summary)
    return summary


def extract_callback_process_instance_id(payload: Any) -> str:
    """Extract a DingTalk approval instance ID from an OA automation payload.

    Args:
        payload: JSON body sent by DingTalk OA automation.

    Returns:
        The normalized approval process instance ID.

    Raises:
        ValueError: No usable approval instance ID is present.
    """
    candidates: list[str] = []

    def collect(value: Any) -> None:
        value = maybe_parse_json(value)
        if isinstance(value, dict):
            for key, child in value.items():
                lowered = str(key).replace("-", "_").lower()
                normalized_key = lowered.replace("_", "")
                if (
                    lowered in CALLBACK_INSTANCE_ID_KEYS
                    or normalized_key in CALLBACK_INSTANCE_ID_KEYS
                ) and isinstance(child, (int, str)):
                    candidates.append(str(child))
                collect(child)
        elif isinstance(value, list):
            for child in value:
                collect(child)
        elif isinstance(value, str):
            for match in re.finditer(
                r"(?:processInstanceId|process_instance_id|procInsId|proc_ins_id)"
                r"=([^&\s]+)",
                value,
                flags=re.IGNORECASE,
            ):
                candidates.append(match.group(1))
            for match in re.finditer(r"\bPROC-[A-Za-z0-9_.:-]+\b", value):
                candidates.append(match.group(0))
            for match in re.finditer(r"\b\d{12,}\b", value.replace(" ", "")):
                candidates.append(match.group(0))

    collect(payload)
    for candidate in candidates:
        normalized = re.sub(r"\s+", "", str(candidate).strip())
        if normalized:
            return normalized
    raise ValueError("OA callback payload does not include an approval instance ID")


def build_oa_callback_review_response(
    config: ExportConfig,
    payload: dict[str, Any],
    ledger_path: Path = DEFAULT_DEDUPE_LEDGER_PATH,
) -> dict[str, Any]:
    """Run the non-destructive finance review for one OA automation callback.

    Args:
        config: Runtime bot configuration.
        payload: JSON body sent by DingTalk OA automation.
        ledger_path: Local duplicate invoice ledger path.

    Returns:
        Compact callback response with the saved review report path.
    """
    instance_id = extract_callback_process_instance_id(payload)
    client = DingTalkFinanceClient(config)
    token = client.get_access_token()
    detail = client.get_approval_detail(token, instance_id)
    detail_process_code = str(
        detail.get("processCode")
        or detail.get("process_code")
        or detail.get("processcode")
        or ""
    )
    if detail_process_code and detail_process_code != config.process_code:
        return {
            "ok": True,
            "mode": "initial_review_skipped",
            "instance_id": instance_id,
            "reason": "process_code_mismatch",
            "report_path": "",
        }
    if config.finance_reviewer_user_id and not find_pending_tasks_for_user(
        detail,
        config.finance_reviewer_user_id,
    ):
        return {
            "ok": True,
            "mode": "initial_review_skipped",
            "instance_id": instance_id,
            "reason": "not_waiting_for_reviewer",
            "report_path": "",
        }
    if not approval_has_general_manager_agreed(detail, config.general_manager_keywords):
        return {
            "ok": True,
            "mode": "initial_review_skipped",
            "instance_id": instance_id,
            "reason": "general_manager_not_approved",
            "report_path": "",
        }
    report = run_initial_review(
        config,
        instance_id,
        ledger_path,
        client=client,
        access_token=token,
        detail=detail,
    )
    issue_codes = [str(issue.get("code", "")) for issue in report.get("issues", [])]
    response: dict[str, Any] = {
        "ok": True,
        "mode": "initial_review",
        "instance_id": instance_id,
        "status": report.get("status", ""),
        "can_auto_approve": report.get("can_auto_approve", False),
        "issue_codes": [code for code in issue_codes if code],
        "recommendation": report.get("recommendation", ""),
        "report_path": report.get("report_path", ""),
    }
    if config.finance_notify_enabled:
        safe_id = re.sub(r"[^A-Za-z0-9_.-]+", "_", instance_id)
        response["notification"] = notify_initial_review_report(
            config,
            client,
            token,
            report,
            config.output_dir / f"notify_{safe_id}.json",
        )
    return response


class DingTalkFinanceStreamEventHandler:
    """Process DingTalk Stream events for finance reimbursement review.

    Args:
        config: Runtime bot configuration.
        ledger_path: Local duplicate invoice ledger path.
        state: Shared stream status dictionary exposed by the web tool.
        state_lock: Lock protecting state updates.
    """

    def __init__(
        self,
        config: ExportConfig,
        ledger_path: Path,
        state: dict[str, Any],
        state_lock: threading.Lock,
    ) -> None:
        self.config = config
        self.ledger_path = ledger_path
        self.state = state
        self.state_lock = state_lock
        self.dingtalk_client: Any = None
        self._seen_event_ids: set[str] = set()

    def pre_start(self) -> None:
        """Match the dingtalk-stream handler interface."""

    async def process(self, event: Any) -> tuple[int, str]:
        """Handle one Stream event and acknowledge it.

        Args:
            event: dingtalk-stream EventMessage.

        Returns:
            DingTalk ack status code and message.
        """
        headers = getattr(event, "headers", None)
        data = getattr(event, "data", {}) or {}
        extensions = getattr(event, "extensions", {}) or {}
        topic = str(getattr(headers, "topic", "") or "")
        event_type = str(getattr(headers, "event_type", "") or "")
        event_id = str(getattr(headers, "event_id", "") or "")
        message_id = str(getattr(headers, "message_id", "") or "")
        dedupe_id = event_id or message_id
        now = datetime.now().isoformat(timespec="seconds")

        with self.state_lock:
            self.state["last_event_at"] = now
            self.state["last_topic"] = topic
            self.state["last_event_type"] = event_type
            self.state["last_error"] = ""

        if dedupe_id and dedupe_id in self._seen_event_ids:
            with self.state_lock:
                self.state["skipped_count"] += 1
                self.state["last_ok"] = True
                self.state["last_reason"] = "duplicate_event"
            return 200, "OK"

        payload = {
            "source": "dingtalk_stream_event",
            "headers": {
                "topic": topic,
                "eventType": event_type,
                "eventId": event_id,
                "messageId": message_id,
            },
            "data": data,
            "extensions": extensions,
        }
        try:
            instance_id = extract_callback_process_instance_id(payload)
        except ValueError:
            if dedupe_id:
                self._seen_event_ids.add(dedupe_id)
            with self.state_lock:
                self.state["skipped_count"] += 1
                self.state["last_ok"] = True
                self.state["last_reason"] = "no_process_instance_id"
                self.state["last_instance_id"] = ""
                self.state["last_mode"] = "stream_event_skipped"
            return 200, "OK"

        try:
            response = await asyncio.to_thread(
                build_oa_callback_review_response,
                self.config,
                payload,
                self.ledger_path,
            )
        except Exception as exc:
            with self.state_lock:
                self.state["error_count"] += 1
                self.state["last_ok"] = False
                self.state["last_error"] = str(exc)
                self.state["last_instance_id"] = instance_id
                self.state["last_mode"] = "stream_event_error"
            return 500, "ERROR"

        if dedupe_id:
            self._seen_event_ids.add(dedupe_id)
            if len(self._seen_event_ids) > 1000:
                self._seen_event_ids = set(list(self._seen_event_ids)[-500:])

        with self.state_lock:
            if response.get("mode") == "initial_review":
                self.state["processed_count"] += 1
            else:
                self.state["skipped_count"] += 1
            self.state["last_ok"] = True
            self.state["last_error"] = ""
            self.state["last_instance_id"] = instance_id
            self.state["last_mode"] = response.get("mode", "")
            self.state["last_reason"] = response.get("reason", "")
            self.state["last_report_path"] = response.get("report_path", "")
        return 200, "OK"

    async def raw_process(self, event_message: Any) -> Any:
        """Return the AckMessage expected by dingtalk-stream.

        Args:
            event_message: dingtalk-stream EventMessage.

        Returns:
            AckMessage for the Stream gateway.
        """
        import dingtalk_stream

        code, message = await self.process(event_message)
        ack_message = dingtalk_stream.AckMessage()
        ack_message.code = code
        ack_message.headers.message_id = event_message.headers.message_id
        ack_message.headers.content_type = "application/json"
        ack_message.message = message
        ack_message.data = event_message.data
        return ack_message


def start_finance_stream_listener(
    config: ExportConfig,
    ledger_path: Path,
    state: dict[str, Any],
    state_lock: threading.Lock,
) -> None:
    """Run the DingTalk Stream listener in the current thread.

    Args:
        config: Runtime bot configuration.
        ledger_path: Local duplicate invoice ledger path.
        state: Shared stream status dictionary exposed by the web tool.
        state_lock: Lock protecting state updates.
    """
    with state_lock:
        state["running"] = True
        state["started_at"] = datetime.now().isoformat(timespec="seconds")
        state["last_error"] = ""
    try:
        import dingtalk_stream

        logger = logging.getLogger("dingtalk_finance_stream")
        credential = dingtalk_stream.Credential(config.app_key, config.app_secret)
        client = dingtalk_stream.DingTalkStreamClient(credential, logger=logger)
        client.register_all_event_handler(
            DingTalkFinanceStreamEventHandler(
                config,
                ledger_path,
                state,
                state_lock,
            )
        )
        client.start_forever()
    except Exception as exc:
        with state_lock:
            state["last_ok"] = False
            state["last_error"] = str(exc)
    finally:
        with state_lock:
            state["running"] = False
            state["finished_at"] = datetime.now().isoformat(timespec="seconds")


def run_oa_callback_server(
    config: ExportConfig,
    *,
    host: str,
    port: int,
    path: str,
    ledger_path: Path,
    callback_token: str = "",
) -> None:
    """Serve the DingTalk OA automation callback endpoint.

    Args:
        config: Runtime bot configuration.
        host: Bind host for the local callback server.
        port: Bind port for the local callback server.
        path: HTTP path that accepts DingTalk OA automation POST requests.
        ledger_path: Local duplicate invoice ledger path.
        callback_token: Optional shared token accepted via query or header.
    """

    class OARequestHandler(BaseHTTPRequestHandler):
        def _json_response(self, status_code: int, data: dict[str, Any]) -> None:
            """Write a JSON HTTP response.

            Args:
                status_code: HTTP status code.
                data: JSON response body.
            """
            body = json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")
            self.send_response(status_code)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def _authorized(self) -> bool:
            """Check the optional shared callback token.

            Returns:
                True when no token is configured or the request token matches.
            """
            if not callback_token:
                return True
            parsed = urlparse(self.path)
            supplied = (
                self.headers.get("X-Dingtalk-Finance-Token")
                or (parse_qs(parsed.query).get("token", [""])[0])
            )
            return supplied == callback_token

        def do_GET(self) -> None:
            """Return a small health response for the callback endpoint."""
            parsed = urlparse(self.path)
            if parsed.path != path:
                self._json_response(404, {"ok": False, "error": "Not found"})
                return
            self._json_response(
                200,
                {
                    "ok": True,
                    "service": "dingtalk_finance_oa_callback",
                    "path": path,
                },
            )

        def do_POST(self) -> None:
            """Process one DingTalk OA automation callback."""
            parsed = urlparse(self.path)
            if parsed.path != path:
                self._json_response(404, {"ok": False, "error": "Not found"})
                return
            if not self._authorized():
                self._json_response(401, {"ok": False, "error": "Unauthorized"})
                return
            try:
                length = int(self.headers.get("Content-Length") or "0")
                raw_body = self.rfile.read(length) if length else b"{}"
                payload = json.loads(raw_body.decode("utf-8") or "{}")
                if not isinstance(payload, dict):
                    raise ValueError("Callback JSON body must be an object")
                result = build_oa_callback_review_response(
                    config,
                    payload,
                    ledger_path,
                )
                self._json_response(200, result)
            except ValueError as exc:
                self._json_response(400, {"ok": False, "error": str(exc)})
            except Exception as exc:
                self._json_response(500, {"ok": False, "error": str(exc)})

        def log_message(self, format_string: str, *args: Any) -> None:
            """Write callback server access logs to stderr.

            Args:
                format_string: BaseHTTPRequestHandler log format string.
                *args: Format values.
            """
            sys.stderr.write(f"[dingtalk-finance-callback] {format_string % args}\n")

    server = ThreadingHTTPServer((host, port), OARequestHandler)
    print(f"Serving DingTalk OA callback on http://{host}:{port}{path}")
    server.serve_forever()


def output_file_href(config: ExportConfig, path_value: str) -> str:
    """Build a web href for a file under the configured output directory.

    Args:
        config: Runtime bot configuration.
        path_value: Absolute or relative output file path.

    Returns:
        URL path served by the NAS tool.

    Raises:
        ValueError: The path is outside the output directory.
    """
    output_dir = config.output_dir.resolve()
    path = Path(path_value)
    if not path.is_absolute():
        path = config.output_dir / path
    resolved = path.resolve()
    relative_path = resolved.relative_to(output_dir)
    return "/output/" + quote(relative_path.as_posix())


def finance_tool_token_matches(
    config: ExportConfig, supplied_tokens: list[str]
) -> bool:
    """Validate a supplied NAS tool token without leaking timing detail.

    Args:
        config: Runtime bot configuration.
        supplied_tokens: Candidate tokens from cookie, header, query, or form.

    Returns:
        True when the NAS tool is open or one supplied token matches.
    """
    if not config.finance_tool_token:
        return True
    expected = config.finance_tool_token
    return any(
        hmac.compare_digest(expected, token) for token in supplied_tokens if token
    )


def finance_tool_auth_mode(config: ExportConfig) -> str:
    """Return the configured NAS tool authentication mode.

    Args:
        config: Runtime bot configuration.

    Returns:
        Authentication mode: dingtalk_oauth, local_token, or open.
    """
    if config.finance_tool_auth_mode:
        normalized = config.finance_tool_auth_mode.strip().lower().replace("-", "_")
        if normalized in {"dingtalk_oauth", "local_token", "open"}:
            return normalized
    if config.app_key and config.app_secret:
        return "dingtalk_oauth"
    if config.finance_tool_token:
        return "local_token"
    return "open"


def build_dingtalk_oauth_authorize_url(
    config: ExportConfig,
    *,
    redirect_uri: str,
    state: str,
) -> str:
    """Build a DingTalk OAuth login URL for the NAS tool.

    Args:
        config: Runtime bot configuration.
        redirect_uri: Tool callback URL registered in DingTalk.
        state: CSRF state value stored locally by the tool.

    Returns:
        DingTalk OAuth authorization URL.
    """
    return (
        DINGTALK_OAUTH_AUTHORIZE_URL
        + "?"
        + urlencode(
            {
                "redirect_uri": redirect_uri,
                "response_type": "code",
                "client_id": config.app_key,
                "scope": "openid",
                "state": state,
                "prompt": "consent",
            }
        )
    )


def list_review_report_summaries(
    output_dir: Path,
    *,
    limit: int = 30,
) -> list[dict[str, Any]]:
    """Read recent local review reports for the NAS tool table.

    Args:
        output_dir: Directory containing review JSON files.
        limit: Maximum reports to return.

    Returns:
        Compact review rows sorted newest first.
    """
    if not output_dir.exists():
        return []
    rows: list[dict[str, Any]] = []
    for path in sorted(
        output_dir.glob("review_*.json"),
        key=lambda item: item.stat().st_mtime,
        reverse=True,
    )[:limit]:
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        issues = data.get("issues") if isinstance(data.get("issues"), list) else []
        rows.append(
            {
                "path": str(path),
                "file_name": path.name,
                "instance_id": str(data.get("instance_id", "")),
                "employee": str(data.get("employee", "")),
                "department": str(data.get("department", "")),
                "project": str(data.get("project", "")),
                "amount": str(data.get("reimbursement_amount", "")),
                "status": str(data.get("status", "")),
                "status_label": REVIEW_STATUS_LABELS.get(
                    str(data.get("status", "")),
                    str(data.get("status", "")),
                ),
                "can_auto_approve": bool(data.get("can_auto_approve", False)),
                "issue_codes": [
                    str(issue.get("code", ""))
                    for issue in issues
                    if isinstance(issue, dict) and issue.get("code")
                ],
                "attachment_count": len(data.get("attachments") or []),
                "recommendation": str(data.get("recommendation", "")),
                "updated_at": datetime.fromtimestamp(path.stat().st_mtime).isoformat(
                    timespec="seconds"
                ),
            }
        )
    return rows


def build_review_summary_workbook(
    config: ExportConfig,
    workbook_path: Path | None = None,
    *,
    limit: int = 5000,
) -> Path:
    """Build a finance-facing workbook from local initial review reports.

    Args:
        config: Runtime bot configuration.
        workbook_path: Optional target workbook path.
        limit: Maximum recent reports to include.

    Returns:
        Saved workbook path.
    """
    rows = list_review_report_summaries(config.output_dir, limit=limit)
    if workbook_path is None:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        workbook_path = config.output_dir / f"finance_review_summary_{stamp}.xlsx"

    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "初审汇总"
    sheet.freeze_panes = "A2"
    sheet.append(REVIEW_SUMMARY_COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for index, row in enumerate(rows, start=1):
        report_path = Path(row["path"])
        issue_text = ", ".join(row["issue_codes"]) or "-"
        sheet.append(
            [
                index,
                row["employee"],
                row["department"],
                row["project"],
                row["amount"],
                row["status_label"],
                issue_text,
                row["attachment_count"],
                row["recommendation"],
                report_path.name,
            ]
        )
        report_cell = sheet.cell(row=index + 1, column=10)
        if report_path.exists():
            report_cell.hyperlink = report_path.resolve().as_uri()
            report_cell.style = "Hyperlink"

    widths = [8, 14, 24, 28, 12, 14, 30, 10, 46, 34]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    workbook.save(workbook_path)
    return workbook_path


def latest_review_summary_workbook(output_dir: Path) -> Path | None:
    """Return the newest generated finance review summary workbook.

    Args:
        output_dir: Directory containing generated workbook files.

    Returns:
        Newest summary workbook path, or None when absent.
    """
    if not output_dir.exists():
        return None
    paths = sorted(
        output_dir.glob("finance_review_summary_*.xlsx"),
        key=lambda item: item.stat().st_mtime,
        reverse=True,
    )
    return paths[0] if paths else None


def latest_poll_summary(output_dir: Path) -> dict[str, Any]:
    """Read the newest poll summary from the local output directory.

    Args:
        output_dir: Directory containing poll summary JSON files.

    Returns:
        Parsed summary data, or an empty dict when none exists.
    """
    if not output_dir.exists():
        return {}
    paths = sorted(
        [*output_dir.glob("poll_*.json"), *output_dir.glob("reviewer_queue_*.json")],
        key=lambda item: item.stat().st_mtime,
        reverse=True,
    )
    if not paths:
        return {}
    try:
        return json.loads(paths[0].read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}


def build_finance_tool_status(
    config: ExportConfig,
    auto_scan_state: dict[str, Any] | None = None,
    stream_state: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Build the NAS tool status payload.

    Args:
        config: Runtime bot configuration.
        auto_scan_state: Optional in-memory background scan status.
        stream_state: Optional in-memory DingTalk Stream listener status.

    Returns:
        Dashboard status with recent review rows and latest poll summary.
    """
    reviews = list_review_report_summaries(config.output_dir)
    latest_poll = latest_poll_summary(config.output_dir)
    latest_summary = latest_review_summary_workbook(config.output_dir)
    auth_mode = finance_tool_auth_mode(config)
    if auto_scan_state is None:
        auto_scan_state = {
            "enabled": False,
            "interval_seconds": 0,
            "running": False,
            "last_started_at": "",
            "last_finished_at": "",
            "last_ok": None,
            "last_error": "",
            "last_summary_path": "",
        }
    if stream_state is None:
        stream_state = {
            "enabled": False,
            "running": False,
            "started_at": "",
            "finished_at": "",
            "last_event_at": "",
            "last_ok": None,
            "last_error": "",
        }
    return {
        "ok": True,
        "channel_name": config.channel_name,
        "process_code": config.process_code,
        "output_dir": str(config.output_dir),
        "tool_auth_enabled": auth_mode != "open",
        "auth_mode": auth_mode,
        "dingtalk_oauth_enabled": auth_mode == "dingtalk_oauth",
        "local_ocr_enabled": bool(config.local_ocr_url),
        "auto_scan": dict(auto_scan_state),
        "stream_events": dict(stream_state),
        "latest_poll": latest_poll,
        "latest_review_summary": str(latest_summary) if latest_summary else "",
        "reviews": reviews,
        "counts": {
            "reviews": len(reviews),
            "needs_manual_review": sum(
                1 for row in reviews if row["status"] == "needs_manual_review"
            ),
            "ready_for_ocr": sum(
                1 for row in reviews if row["status"] == "ready_for_ocr"
            ),
            "ready_for_finance_review": sum(
                1 for row in reviews if row["status"] == "ready_for_finance_review"
            ),
        },
    }


def render_finance_tool_login_html(
    config: ExportConfig,
    *,
    oauth_start_url: str = "/auth/dingtalk/start",
    error: str = "",
) -> str:
    """Render the NAS finance tool login page.

    Args:
        config: Runtime bot configuration.
        oauth_start_url: Local URL that starts DingTalk OAuth.
        error: Optional login error.

    Returns:
        Complete HTML document.
    """
    notice_html = (
        f'<div class="notice error">{html.escape(error)}</div>' if error else ""
    )
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>{html.escape(config.channel_name)}｜登录</title>
  <style>
    :root {{
      color-scheme: light;
      --bg: #f6f7f9;
      --panel: #ffffff;
      --line: #d9dee7;
      --text: #1e2633;
      --muted: #687385;
      --blue: #1677ff;
      --red: #c93535;
    }}
    * {{ box-sizing: border-box; }}
    body {{
      margin: 0;
      min-height: 100vh;
      display: grid;
      place-items: center;
      background: var(--bg);
      color: var(--text);
      font: 14px/1.5 -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif;
    }}
    main {{
      width: min(420px, calc(100vw - 32px));
      background: var(--panel);
      border: 1px solid var(--line);
      border-radius: 6px;
      padding: 24px;
    }}
    h1 {{
      margin: 0 0 18px;
      font-size: 22px;
      letter-spacing: 0;
    }}
    label {{
      display: grid;
      gap: 6px;
      color: var(--muted);
      font-size: 12px;
    }}
    input {{
      width: 100%;
      height: 38px;
      border: 1px solid var(--line);
      border-radius: 4px;
      padding: 0 10px;
      font: inherit;
      color: var(--text);
      background: #fff;
    }}
    button {{
      width: 100%;
      height: 38px;
      margin-top: 14px;
      border: 0;
      border-radius: 4px;
      background: var(--blue);
      color: #fff;
      font-weight: 700;
      cursor: pointer;
    }}
    a.button {{
      width: 100%;
      height: 40px;
      margin-top: 4px;
      border-radius: 4px;
      background: var(--blue);
      color: #fff;
      font-weight: 700;
      text-decoration: none;
      display: inline-flex;
      align-items: center;
      justify-content: center;
    }}
    details {{
      margin-top: 18px;
      color: var(--muted);
    }}
    summary {{ cursor: pointer; }}
    .notice {{
      margin-bottom: 14px;
      border-radius: 6px;
      padding: 10px 12px;
      border: 1px solid #ebb1b1;
      color: var(--red);
      background: #fff;
    }}
  </style>
</head>
<body>
  <main>
    <h1>{html.escape(config.channel_name)}</h1>
    {notice_html}
    <a class="button" href="{html.escape(oauth_start_url)}">使用钉钉账号授权登录</a>
  </main>
</body>
</html>"""


def render_finance_tool_html(
    config: ExportConfig,
    *,
    message: str = "",
    error: str = "",
    auto_scan_state: dict[str, Any] | None = None,
    current_user: dict[str, Any] | None = None,
) -> str:
    """Render the NAS finance tool HTML page.

    Args:
        config: Runtime bot configuration.
        message: Optional success message.
        error: Optional error message.
        auto_scan_state: Optional in-memory background scan status.
        current_user: Optional DingTalk OAuth user session.

    Returns:
        Complete HTML document.
    """
    status = build_finance_tool_status(config, auto_scan_state)
    counts = status["counts"]
    latest_poll = status["latest_poll"]
    auto_scan = status["auto_scan"]
    reviews = status["reviews"]
    poll_path = str(latest_poll.get("summary_path", "")) if latest_poll else ""
    poll_href = ""
    if poll_path:
        try:
            poll_href = output_file_href(config, poll_path)
        except ValueError:
            poll_href = ""

    rows_html: list[str] = []
    for row in reviews:
        try:
            report_href = output_file_href(config, row["path"])
        except ValueError:
            report_href = ""
        issue_text = ", ".join(row["issue_codes"]) or "-"
        rows_html.append(
            "<tr>"
            f"<td>{html.escape(row['employee'] or '-')}</td>"
            f"<td>{html.escape(row['department'] or '-')}</td>"
            f"<td>{html.escape(row['project'] or '-')}</td>"
            f"<td>{html.escape(row['amount'] or '-')}</td>"
            f'<td><span class="status-pill status-{html.escape(row["status"])}">'
            f"{html.escape(row['status_label'] or '-')}</span></td>"
            f"<td>{html.escape(issue_text)}</td>"
            f"<td>{html.escape(str(row['attachment_count']))}</td>"
            f"<td>{html.escape(row['recommendation'] or '-')}</td>"
            f'<td><a href="{html.escape(report_href)}">JSON</a></td>'
            "</tr>"
        )
    if not rows_html:
        rows_html.append('<tr><td colspan="9" class="empty-row">暂无初审报告</td></tr>')

    notice_html = ""
    if message:
        notice_html = f'<div class="notice success">{html.escape(message)}</div>'
    if error:
        notice_html = f'<div class="notice error">{html.escape(error)}</div>'

    poll_link_html = (
        f'<a href="{html.escape(poll_href)}">扫描汇总</a>' if poll_href else "-"
    )
    latest_poll_text = "-"
    if latest_poll:
        mode_label = (
            "待审队列" if latest_poll.get("mode") == "reviewer_queue" else "审批扫描"
        )
        latest_poll_text = (
            f"{mode_label}：检查 {latest_poll.get('inspected_count', 0)} 条，"
            f"生成 {latest_poll.get('reviewed_count', 0)} 条报告，"
            f"跳过 {latest_poll.get('skipped_count', 0)} 条"
        )
    latest_summary_href = ""
    if status["latest_review_summary"]:
        try:
            latest_summary_href = output_file_href(
                config,
                str(status["latest_review_summary"]),
            )
        except ValueError:
            latest_summary_href = ""
    summary_link_html = (
        f'<a href="{html.escape(latest_summary_href)}">最近汇总表</a>'
        if latest_summary_href
        else "-"
    )
    if auto_scan.get("enabled"):
        auto_scan_text = (
            f"自动轮询每 {auto_scan.get('interval_seconds', 0)} 秒 · "
            f"{'运行中' if auto_scan.get('running') else '待命'}"
        )
        if auto_scan.get("last_finished_at"):
            auto_scan_text += f" · 上次 {auto_scan['last_finished_at']}"
        if auto_scan.get("last_error"):
            auto_scan_text += f" · {auto_scan['last_error']}"
    else:
        auto_scan_text = "自动轮询未启用"
    auth_label = {
        "dingtalk_oauth": "钉钉",
        "local_token": "口令",
        "open": "开放",
    }.get(str(status["auth_mode"]), str(status["auth_mode"]))
    ocr_label = "本地" if status["local_ocr_enabled"] else "未启用"
    user_name = ""
    if current_user:
        user_name = str(
            current_user.get("nick")
            or current_user.get("name")
            or current_user.get("union_id")
            or current_user.get("open_id")
            or ""
        )
    if status["auth_mode"] == "open":
        auth_action = '<span class="meta">开放访问</span>'
    elif user_name:
        auth_action = (
            f'<span class="meta">{html.escape(user_name)}</span>'
            '<a class="header-link" href="/logout">退出</a>'
        )
    else:
        auth_action = '<a class="header-link" href="/logout">退出</a>'

    return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>{html.escape(config.channel_name)}｜NAS 工具</title>
  <style>
    :root {{
      color-scheme: light;
      --bg: #f6f7f9;
      --panel: #ffffff;
      --line: #d9dee7;
      --text: #1e2633;
      --muted: #687385;
      --blue: #1677ff;
      --green: #168a45;
      --amber: #a86500;
      --red: #c93535;
    }}
    * {{ box-sizing: border-box; }}
    body {{
      margin: 0;
      background: var(--bg);
      color: var(--text);
      font: 14px/1.5 -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif;
    }}
    header {{
      display: flex;
      align-items: center;
      justify-content: space-between;
      padding: 18px 28px;
      background: var(--panel);
      border-bottom: 1px solid var(--line);
    }}
    h1 {{
      margin: 0;
      font-size: 22px;
      font-weight: 700;
      letter-spacing: 0;
    }}
    .header-actions {{
      display: flex;
      align-items: center;
      gap: 14px;
      flex-wrap: wrap;
    }}
    .header-link {{
      color: var(--blue);
      font-weight: 700;
      text-decoration: none;
    }}
    main {{
      max-width: 1280px;
      margin: 0 auto;
      padding: 24px 28px 40px;
    }}
    .toolbar {{
      display: grid;
      grid-template-columns: minmax(360px, 1.25fr) minmax(320px, 1fr);
      gap: 18px;
      margin-bottom: 20px;
    }}
    section {{
      background: var(--panel);
      border: 1px solid var(--line);
      border-radius: 6px;
      padding: 18px;
    }}
    h2 {{
      margin: 0 0 14px;
      font-size: 16px;
      letter-spacing: 0;
    }}
    .metrics {{
      display: grid;
      grid-template-columns: repeat(4, minmax(0, 1fr));
      gap: 10px;
    }}
    .metric {{
      border: 1px solid var(--line);
      border-radius: 6px;
      padding: 12px;
      min-height: 70px;
    }}
    .metric strong {{
      display: block;
      font-size: 24px;
      line-height: 1.1;
    }}
    .metric span {{ color: var(--muted); }}
    form {{
      display: grid;
      grid-template-columns: repeat(4, minmax(0, 1fr));
      gap: 10px;
      align-items: end;
    }}
    label {{
      display: grid;
      gap: 6px;
      color: var(--muted);
      font-size: 12px;
    }}
    input {{
      width: 100%;
      height: 36px;
      border: 1px solid var(--line);
      border-radius: 4px;
      padding: 0 10px;
      font: inherit;
      color: var(--text);
      background: #fff;
    }}
    button, .button-link {{
      height: 36px;
      border: 0;
      border-radius: 4px;
      padding: 0 14px;
      background: var(--blue);
      color: #fff;
      font-weight: 600;
      cursor: pointer;
      text-decoration: none;
      display: inline-flex;
      align-items: center;
      justify-content: center;
    }}
    .button-secondary {{ background: #303846; }}
    .actions {{
      display: flex;
      gap: 10px;
      flex-wrap: wrap;
      margin-top: 12px;
    }}
    .actions form {{
      display: flex;
      gap: 10px;
      align-items: end;
    }}
    .actions input {{ min-width: 220px; }}
    .notice {{
      margin-bottom: 16px;
      border-radius: 6px;
      padding: 10px 12px;
      border: 1px solid var(--line);
      background: var(--panel);
    }}
    .notice.success {{ border-color: #9bd6b3; color: var(--green); }}
    .notice.error {{ border-color: #ebb1b1; color: var(--red); }}
    table {{
      width: 100%;
      border-collapse: collapse;
      background: var(--panel);
      border: 1px solid var(--line);
      border-radius: 6px;
      overflow: hidden;
    }}
    th, td {{
      padding: 10px 12px;
      border-bottom: 1px solid var(--line);
      text-align: left;
      vertical-align: top;
    }}
    th {{
      color: var(--muted);
      font-size: 12px;
      font-weight: 600;
      background: #f0f3f7;
      white-space: nowrap;
    }}
    td a {{ color: var(--blue); font-weight: 600; }}
    .status-pill {{
      display: inline-flex;
      align-items: center;
      min-height: 24px;
      border-radius: 12px;
      padding: 0 9px;
      font-size: 12px;
      font-weight: 700;
      white-space: nowrap;
      background: #eef2f7;
      color: var(--muted);
    }}
    .status-needs_manual_review {{
      background: #fff1f0;
      color: var(--red);
    }}
    .status-ready_for_ocr {{
      background: #fff7e6;
      color: var(--amber);
    }}
    .empty-row {{
      color: var(--muted);
      text-align: center;
      padding: 36px 12px;
    }}
    .meta {{
      color: var(--muted);
      overflow-wrap: anywhere;
    }}
    @media (max-width: 900px) {{
      header {{ align-items: flex-start; flex-direction: column; gap: 8px; }}
      main {{ padding: 18px 14px; }}
      .toolbar {{ grid-template-columns: 1fr; }}
      form {{ grid-template-columns: 1fr 1fr; }}
      .metrics {{ grid-template-columns: 1fr; }}
      table {{ display: block; overflow-x: auto; }}
    }}
  </style>
</head>
<body>
  <header>
    <h1>{html.escape(config.channel_name)}</h1>
    <div class="header-actions">
      <span class="meta">NAS 工具 · {html.escape(str(config.output_dir))}</span>
      {auth_action}
    </div>
  </header>
  <main>
    {notice_html}
    <div class="toolbar">
      <section>
        <h2>{html.escape(config.finance_reviewer_name)}待审单据</h2>
        <form method="post" action="/scan">
          <label>报账人员<input name="employee" placeholder="可留空" /></label>
          <label>关键词<input name="keyword" placeholder="采购网络设备" /></label>
          <label>金额<input name="amount" placeholder="1315.98" /></label>
          <label>数量<input name="limit" type="number" min="1" max="100" value="20" /></label>
          <button type="submit">扫描待审</button>
        </form>
        <div class="actions">
          <form method="post" action="/export">
            <input name="instance_id" placeholder="单据 ID，可留空" />
            <button class="button-secondary" type="submit">导出 Excel</button>
          </form>
          <form method="post" action="/summary">
            <button class="button-secondary" type="submit">导出初审汇总</button>
          </form>
        </div>
      </section>
      <section>
        <h2>最近状态</h2>
        <div class="metrics">
          <div class="metric"><strong>{counts["reviews"]}</strong><span>报告</span></div>
          <div class="metric"><strong>{counts["needs_manual_review"]}</strong><span>需处理</span></div>
          <div class="metric"><strong>{counts["ready_for_ocr"]}</strong><span>待 OCR</span></div>
          <div class="metric"><strong>{counts["ready_for_finance_review"]}</strong><span>可复审</span></div>
          <div class="metric"><strong>{html.escape(ocr_label)}</strong><span>OCR</span></div>
          <div class="metric"><strong>{html.escape(auth_label)}</strong><span>授权</span></div>
        </div>
        <p class="meta">{html.escape(latest_poll_text)} · {poll_link_html}</p>
        <p class="meta">{html.escape(auto_scan_text)} · {summary_link_html}</p>
      </section>
    </div>
    <table>
      <thead>
        <tr>
          <th>员工</th>
          <th>部门</th>
          <th>项目</th>
          <th>金额</th>
          <th>状态</th>
          <th>问题</th>
          <th>附件</th>
          <th>建议</th>
          <th>报告</th>
        </tr>
      </thead>
      <tbody>
        {"".join(rows_html)}
      </tbody>
    </table>
  </main>
</body>
</html>"""


def run_finance_tool_server(
    config: ExportConfig,
    *,
    host: str,
    port: int,
    ledger_path: Path,
    auto_scan_interval: int = 0,
    auto_scan_limit: int = 20,
    auto_scan_skip_reviewed: bool = False,
    stream_events: bool = False,
) -> None:
    """Serve the NAS finance review web tool.

    Args:
        config: Runtime bot configuration.
        host: Bind host for the NAS web tool.
        port: Bind port for the NAS web tool.
        ledger_path: Local duplicate invoice ledger path.
        auto_scan_interval: Background poll interval in seconds; 0 disables it.
        auto_scan_limit: Maximum approval details per automatic scan.
        auto_scan_skip_reviewed: Skip instances that already have reports.
        stream_events: Start DingTalk event Stream listener in the background.
    """
    auto_scan_interval = max(0, auto_scan_interval)
    auto_scan_limit = max(1, min(auto_scan_limit, 100))
    auto_scan_state: dict[str, Any] = {
        "enabled": auto_scan_interval > 0,
        "interval_seconds": auto_scan_interval,
        "limit": auto_scan_limit,
        "skip_reviewed": auto_scan_skip_reviewed,
        "running": False,
        "last_started_at": "",
        "last_finished_at": "",
        "last_ok": None,
        "last_error": "",
        "last_summary_path": "",
        "last_reviewed_count": 0,
    }
    auto_scan_lock = threading.Lock()
    auto_scan_stop = threading.Event()
    stream_state: dict[str, Any] = {
        "enabled": stream_events,
        "running": False,
        "started_at": "",
        "finished_at": "",
        "last_event_at": "",
        "last_topic": "",
        "last_event_type": "",
        "last_instance_id": "",
        "last_mode": "",
        "last_reason": "",
        "last_report_path": "",
        "last_ok": None,
        "last_error": "",
        "processed_count": 0,
        "skipped_count": 0,
        "error_count": 0,
    }
    stream_lock = threading.Lock()
    oauth_states: dict[str, dict[str, str]] = {}
    sessions: dict[str, dict[str, Any]] = {}
    auth_lock = threading.Lock()

    def snapshot_auto_scan_state() -> dict[str, Any]:
        """Copy the current background scan state.

        Returns:
            Stable state snapshot for rendering or JSON responses.
        """
        with auto_scan_lock:
            return dict(auto_scan_state)

    def snapshot_stream_state() -> dict[str, Any]:
        """Copy the current Stream listener state.

        Returns:
            Stable state snapshot for rendering or JSON responses.
        """
        with stream_lock:
            return dict(stream_state)

    def run_background_scan_once() -> None:
        """Run one background scan and update in-memory status."""
        with auto_scan_lock:
            if auto_scan_state["running"]:
                return
            auto_scan_state["running"] = True
            auto_scan_state["last_started_at"] = datetime.now().isoformat(
                timespec="seconds"
            )
            auto_scan_state["last_error"] = ""
        try:
            if not config.finance_reviewer_user_id:
                raise ValueError("Auto scan requires DINGTALK_FINANCE_REVIEWER_USER_ID")
            summary = run_reviewer_queue_once(
                config,
                reviewer_user_id=config.finance_reviewer_user_id,
                reviewer_union_id=config.finance_reviewer_union_id,
                ledger_path=ledger_path,
                limit=auto_scan_limit,
                skip_reviewed=auto_scan_skip_reviewed,
            )
            with auto_scan_lock:
                auto_scan_state["last_ok"] = True
                auto_scan_state["last_summary_path"] = str(
                    summary.get("summary_path", "")
                )
                auto_scan_state["last_reviewed_count"] = int(
                    summary.get("reviewed_count", 0)
                )
        except Exception as exc:
            with auto_scan_lock:
                auto_scan_state["last_ok"] = False
                auto_scan_state["last_error"] = str(exc)
        finally:
            with auto_scan_lock:
                auto_scan_state["running"] = False
                auto_scan_state["last_finished_at"] = datetime.now().isoformat(
                    timespec="seconds"
                )

    def auto_scan_loop() -> None:
        """Run periodic background scans until the server stops."""
        run_background_scan_once()
        while not auto_scan_stop.wait(auto_scan_interval):
            run_background_scan_once()

    if auto_scan_interval:
        threading.Thread(
            target=auto_scan_loop,
            name="dingtalk-finance-auto-scan",
            daemon=True,
        ).start()
    if stream_events:
        threading.Thread(
            target=start_finance_stream_listener,
            args=(config, ledger_path, stream_state, stream_lock),
            name="dingtalk-finance-stream",
            daemon=True,
        ).start()

    class FinanceToolHandler(BaseHTTPRequestHandler):
        def _send_bytes(
            self,
            status_code: int,
            body: bytes,
            *,
            content_type: str,
            headers: dict[str, str | list[str]] | None = None,
        ) -> None:
            """Write an HTTP response.

            Args:
                status_code: HTTP status code.
                body: Response body bytes.
                content_type: Response content type.
                headers: Optional extra response headers.
            """
            self.send_response(status_code)
            self.send_header("Content-Type", content_type)
            self.send_header("Content-Length", str(len(body)))
            for key, value in (headers or {}).items():
                if isinstance(value, list):
                    for current_value in value:
                        self.send_header(key, current_value)
                else:
                    self.send_header(key, value)
            self.end_headers()
            self.wfile.write(body)

        def _send_json(self, status_code: int, data: dict[str, Any]) -> None:
            """Write a JSON response.

            Args:
                status_code: HTTP status code.
                data: JSON response body.
            """
            self._send_bytes(
                status_code,
                json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8"),
                content_type="application/json; charset=utf-8",
            )

        def _send_html(self, message: str = "", error: str = "") -> None:
            """Write the dashboard HTML.

            Args:
                message: Optional success message.
                error: Optional error message.
            """
            self._send_bytes(
                200,
                render_finance_tool_html(
                    config,
                    message=message,
                    error=error,
                    auto_scan_state=snapshot_auto_scan_state(),
                    current_user=self._current_user(),
                ).encode("utf-8"),
                content_type="text/html; charset=utf-8",
            )

        def _send_login(self, error: str = "") -> None:
            """Write the login HTML.

            Args:
                error: Optional login error.
            """
            parsed = urlparse(self.path)
            next_path = parse_qs(parsed.query).get("next", ["/"])[0] or "/"
            if not next_path.startswith("/") or next_path.startswith("//"):
                next_path = "/"
            self._send_bytes(
                200,
                render_finance_tool_login_html(
                    config,
                    oauth_start_url=(
                        "/auth/dingtalk/start?next=" + quote(next_path, safe="/")
                    ),
                    error=error,
                ).encode("utf-8"),
                content_type="text/html; charset=utf-8",
            )

        def _send_redirect(
            self,
            location: str,
            *,
            headers: dict[str, str | list[str]] | None = None,
        ) -> None:
            """Redirect the browser to another path.

            Args:
                location: Redirect target.
                headers: Optional extra response headers.
            """
            redirect_headers = {"Location": location}
            redirect_headers.update(headers or {})
            self._send_bytes(
                303,
                b"",
                content_type="text/plain; charset=utf-8",
                headers=redirect_headers,
            )

        def _request_base_url(self) -> str:
            """Build the external base URL used for OAuth callbacks.

            Returns:
                Base URL without a trailing slash.
            """
            if config.finance_tool_base_url:
                return config.finance_tool_base_url.rstrip("/")
            proto = (
                self.headers.get("X-Forwarded-Proto", "http").split(",", 1)[0].strip()
                or "http"
            )
            host_value = (
                self.headers.get("X-Forwarded-Host")
                or self.headers.get("Host")
                or f"{host}:{port}"
            )
            return f"{proto}://{host_value}".rstrip("/")

        def _start_dingtalk_oauth(self) -> None:
            """Redirect the browser to DingTalk OAuth authorization."""
            parsed = urlparse(self.path)
            next_path = parse_qs(parsed.query).get("next", ["/"])[0] or "/"
            if not next_path.startswith("/") or next_path.startswith("//"):
                next_path = "/"
            state = secrets.token_urlsafe(24)
            redirect_uri = self._request_base_url() + "/auth/dingtalk/callback"
            with auth_lock:
                oauth_states[state] = {
                    "next": next_path,
                    "redirect_uri": redirect_uri,
                    "created_at": datetime.now().isoformat(timespec="seconds"),
                }
            self._send_redirect(
                build_dingtalk_oauth_authorize_url(
                    config,
                    redirect_uri=redirect_uri,
                    state=state,
                )
            )

        def _handle_dingtalk_callback(self) -> None:
            """Complete DingTalk OAuth login and create a local session."""
            parsed = urlparse(self.path)
            query = parse_qs(parsed.query)
            if query.get("error"):
                self._send_login(
                    error=query.get("error_description", query["error"])[0]
                )
                return
            code = query.get("code", [""])[0]
            state = query.get("state", [""])[0]
            if not code or not state:
                self._send_login(error="钉钉授权回调缺少 code 或 state")
                return
            with auth_lock:
                state_payload = oauth_states.pop(state, None)
            if not state_payload:
                self._send_login(error="钉钉授权状态已失效，请重新登录")
                return
            try:
                client = DingTalkFinanceClient(config)
                token_payload = client.get_oauth_user_access_token(code)
                user_payload = client.get_oauth_user_info(
                    str(token_payload["accessToken"])
                )
                union_id = str(
                    user_payload.get("unionId") or user_payload.get("unionid") or ""
                )
                nick = str(
                    user_payload.get("nick")
                    or user_payload.get("name")
                    or user_payload.get("userName")
                    or ""
                )
                if (
                    config.finance_allowed_union_ids
                    and union_id not in config.finance_allowed_union_ids
                ):
                    self._send_login(error="当前钉钉账号不在财务工具授权名单内")
                    return
                if config.finance_reviewer_union_id:
                    if union_id != config.finance_reviewer_union_id:
                        self._send_login(
                            error=f"请使用{config.finance_reviewer_name}的钉钉账号登录"
                        )
                        return
                elif config.finance_reviewer_name and (
                    compact_text(config.finance_reviewer_name) not in compact_text(nick)
                ):
                    self._send_login(
                        error=f"请使用{config.finance_reviewer_name}的钉钉账号登录"
                    )
                    return
                reviewer_user_id = config.finance_reviewer_user_id
                if not reviewer_user_id:
                    if not union_id:
                        self._send_login(
                            error="钉钉授权没有返回 unionId，无法定位审批账号"
                        )
                        return
                    app_token = client.get_access_token()
                    union_mapping = client.get_user_by_union_id(app_token, union_id)
                    reviewer_user_id = str(
                        union_mapping.get("userid")
                        or union_mapping.get("userId")
                        or union_mapping.get("user_id")
                        or ""
                    )
                    if not reviewer_user_id:
                        self._send_login(
                            error="已登录钉钉，但无法把 unionId 转成审批 userId"
                        )
                        return
                session_id = secrets.token_urlsafe(32)
                session_user = {
                    "nick": nick,
                    "name": config.finance_reviewer_name,
                    "user_id": reviewer_user_id,
                    "union_id": union_id,
                    "open_id": str(
                        user_payload.get("openId") or user_payload.get("openid") or ""
                    ),
                    "corp_id": str(token_payload.get("corpId", "")),
                    "login_at": datetime.now().isoformat(timespec="seconds"),
                }
                with auth_lock:
                    sessions[session_id] = session_user
                self._send_redirect(
                    state_payload.get("next", "/") or "/",
                    headers={
                        "Set-Cookie": (
                            f"{FINANCE_TOOL_SESSION_COOKIE}={quote(session_id)}; "
                            "HttpOnly; SameSite=Lax; Path=/"
                        )
                    },
                )
            except Exception as exc:
                self._send_login(error=f"钉钉授权失败：{exc}")

        def _current_user(self) -> dict[str, Any] | None:
            """Read the current DingTalk OAuth session from cookies.

            Returns:
                Session user payload, or None when not logged in.
            """
            cookie = SimpleCookie()
            cookie.load(self.headers.get("Cookie", ""))
            if FINANCE_TOOL_SESSION_COOKIE not in cookie:
                return None
            session_id = unquote(cookie[FINANCE_TOOL_SESSION_COOKIE].value)
            with auth_lock:
                return sessions.get(session_id)

        def _read_form(self) -> dict[str, str]:
            """Read a URL-encoded form body.

            Returns:
                Flat form values.
            """
            length = int(self.headers.get("Content-Length") or "0")
            raw_body = self.rfile.read(length).decode("utf-8") if length else ""
            parsed = parse_qs(raw_body)
            return {key: values[0] for key, values in parsed.items() if values}

        def _authorized(self, form: dict[str, str] | None = None) -> bool:
            """Check DingTalk OAuth session or the optional fallback token.

            Args:
                form: Already parsed form data for POST requests.

            Returns:
                True when the request can access the tool.
            """
            if self._current_user():
                return True
            auth_mode = finance_tool_auth_mode(config)
            if auth_mode == "open":
                return True
            if not config.finance_tool_token:
                return False
            parsed = urlparse(self.path)
            cookie = SimpleCookie()
            cookie.load(self.headers.get("Cookie", ""))
            cookie_token = ""
            if FINANCE_TOOL_FALLBACK_COOKIE in cookie:
                cookie_token = unquote(cookie[FINANCE_TOOL_FALLBACK_COOKIE].value)
            supplied_tokens = [
                cookie_token,
                self.headers.get("X-Dingtalk-Finance-Tool-Token", ""),
                parse_qs(parsed.query).get("token", [""])[0],
            ]
            if form:
                supplied_tokens.append(form.get("token", ""))
            return finance_tool_token_matches(config, supplied_tokens)

        def _serve_output_file(self, raw_path: str) -> None:
            """Serve one file under the configured output directory.

            Args:
                raw_path: URL-decoded path suffix under /output/.
            """
            relative_path = Path(unquote(raw_path).lstrip("/"))
            output_dir = config.output_dir.resolve()
            path = (output_dir / relative_path).resolve()
            try:
                path.relative_to(output_dir)
            except ValueError:
                self._send_json(403, {"ok": False, "error": "Forbidden"})
                return
            if not path.is_file():
                self._send_json(404, {"ok": False, "error": "Not found"})
                return
            content_type = (
                mimetypes.guess_type(path.name)[0] or "application/octet-stream"
            )
            self._send_bytes(
                200,
                path.read_bytes(),
                content_type=content_type,
                headers={
                    "Content-Disposition": (
                        f"inline; filename*=UTF-8''{quote(path.name)}"
                    )
                },
            )

        def do_GET(self) -> None:
            """Handle dashboard, status API, and output file requests."""
            parsed = urlparse(self.path)
            if parsed.path == DEFAULT_OA_CALLBACK_PATH:
                self._send_json(
                    200,
                    {
                        "ok": True,
                        "service": "dingtalk_finance_oa_callback",
                        "path": DEFAULT_OA_CALLBACK_PATH,
                    },
                )
                return
            if parsed.path == "/login":
                if self._authorized():
                    self._send_redirect("/")
                    return
                self._send_login()
                return
            if parsed.path == "/auth/dingtalk/start":
                self._start_dingtalk_oauth()
                return
            if parsed.path == "/auth/dingtalk/callback":
                self._handle_dingtalk_callback()
                return
            if parsed.path == "/logout":
                cookie = SimpleCookie()
                cookie.load(self.headers.get("Cookie", ""))
                if FINANCE_TOOL_SESSION_COOKIE in cookie:
                    session_id = unquote(cookie[FINANCE_TOOL_SESSION_COOKIE].value)
                    with auth_lock:
                        sessions.pop(session_id, None)
                self._send_redirect(
                    "/login",
                    headers={
                        "Set-Cookie": [
                            (
                                f"{FINANCE_TOOL_SESSION_COOKIE}=; Max-Age=0; "
                                "HttpOnly; SameSite=Lax; Path=/"
                            ),
                            (
                                f"{FINANCE_TOOL_FALLBACK_COOKIE}=; Max-Age=0; "
                                "HttpOnly; SameSite=Lax; Path=/"
                            ),
                        ]
                    },
                )
                return
            if parsed.path == "/api/session":
                self._send_json(
                    200,
                    {
                        "ok": True,
                        "authenticated": bool(self._current_user()),
                        "user": self._current_user() or {},
                        "auth_mode": finance_tool_auth_mode(config),
                    },
                )
                return
            if not self._authorized():
                if parsed.path == "/":
                    self._send_redirect("/login?next=/")
                    return
                self._send_json(401, {"ok": False, "error": "Unauthorized"})
                return
            if parsed.path == "/":
                self._send_html()
                return
            if parsed.path == "/api/status":
                self._send_json(
                    200,
                    build_finance_tool_status(
                        config,
                        snapshot_auto_scan_state(),
                        snapshot_stream_state(),
                    ),
                )
                return
            if parsed.path.startswith("/output/"):
                self._serve_output_file(parsed.path.removeprefix("/output/"))
                return
            self._send_json(404, {"ok": False, "error": "Not found"})

        def do_POST(self) -> None:
            """Handle scan, export, and OA automation callback submissions."""
            parsed = urlparse(self.path)
            try:
                if parsed.path == DEFAULT_OA_CALLBACK_PATH:
                    callback_token = config.oa_callback_token
                    if callback_token:
                        supplied = (
                            self.headers.get("X-Dingtalk-Finance-Token")
                            or parse_qs(parsed.query).get("token", [""])[0]
                        )
                        if supplied != callback_token:
                            self._send_json(
                                401,
                                {"ok": False, "error": "Unauthorized"},
                            )
                            return
                    length = int(self.headers.get("Content-Length") or "0")
                    raw_body = self.rfile.read(length) if length else b"{}"
                    payload = json.loads(raw_body.decode("utf-8") or "{}")
                    if not isinstance(payload, dict):
                        raise ValueError("Callback JSON body must be an object")
                    self._send_json(
                        200,
                        build_oa_callback_review_response(
                            config,
                            payload,
                            ledger_path,
                        ),
                    )
                    return
                form = self._read_form()
                if parsed.path == "/login":
                    if not config.finance_tool_token:
                        self._send_login(error="请使用钉钉账号授权登录")
                        return
                    if self._authorized(form):
                        cookie_value = quote(config.finance_tool_token, safe="")
                        self._send_redirect(
                            "/",
                            headers={
                                "Set-Cookie": (
                                    f"{FINANCE_TOOL_FALLBACK_COOKIE}="
                                    f"{cookie_value}; HttpOnly; SameSite=Lax; Path=/"
                                )
                            },
                        )
                        return
                    self._send_login(error="访问口令不正确")
                    return
                if not self._authorized(form):
                    self._send_redirect("/login?next=/")
                    return
                if parsed.path == "/scan":
                    current_user = self._current_user()
                    reviewer_user_id = (
                        str(current_user.get("user_id", ""))
                        if current_user
                        else config.finance_reviewer_user_id
                    )
                    reviewer_union_id = (
                        str(current_user.get("union_id", ""))
                        if current_user
                        else config.finance_reviewer_union_id
                    )
                    if not reviewer_user_id:
                        self._send_html(
                            error=(
                                f"请使用{config.finance_reviewer_name}钉钉账号登录后再扫描，"
                                "或在本地配置 DINGTALK_FINANCE_REVIEWER_USER_ID"
                            )
                        )
                        return
                    terms = [
                        value
                        for value in [
                            form.get("employee", "").strip(),
                            form.get("amount", "").strip(),
                            form.get("keyword", "").strip(),
                        ]
                        if value
                    ]
                    limit = int(form.get("limit", "20") or "20")
                    limit = max(1, min(limit, 100))
                    summary = run_reviewer_queue_once(
                        config,
                        reviewer_user_id=reviewer_user_id,
                        reviewer_union_id=reviewer_union_id,
                        ledger_path=ledger_path,
                        terms=terms,
                        limit=limit,
                    )
                    self._send_html(
                        message=(
                            f"{config.finance_reviewer_name}待审扫描完成："
                            f"检查 {summary['inspected_count']} 条，"
                            f"生成 {summary['reviewed_count']} 条报告。"
                        )
                    )
                    return
                if parsed.path == "/export":
                    instance_id = form.get("instance_id", "").strip() or None
                    workbook_path = run_export(config, instance_id)
                    href = output_file_href(config, str(workbook_path))
                    self._send_html(message=f"Excel 已生成：{href}")
                    return
                if parsed.path == "/summary":
                    workbook_path = build_review_summary_workbook(config)
                    href = output_file_href(config, str(workbook_path))
                    self._send_html(message=f"初审汇总已生成：{href}")
                    return
                self._send_json(404, {"ok": False, "error": "Not found"})
            except Exception as exc:
                self._send_html(error=str(exc))

        def log_message(self, format_string: str, *args: Any) -> None:
            """Write access logs to stderr.

            Args:
                format_string: BaseHTTPRequestHandler log format string.
                *args: Format values.
            """
            sys.stderr.write(f"[dingtalk-finance-tool] {format_string % args}\n")

    config.output_dir.mkdir(parents=True, exist_ok=True)
    server = ThreadingHTTPServer((host, port), FinanceToolHandler)
    print(f"Serving DingTalk finance NAS tool on http://{host}:{port}/")
    try:
        server.serve_forever()
    finally:
        auto_scan_stop.set()


def main() -> int:
    """Run the DingTalk finance bot CLI.

    Returns:
        Process exit code.
    """
    parser = argparse.ArgumentParser(description="DingTalk finance reimbursement bot")
    parser.add_argument(
        "command",
        choices=[
            "capabilities",
            "probe",
            "export",
            "users",
            "review",
            "poll-once",
            "reviewer-queue",
            "serve-callback",
            "serve-tool",
        ],
    )
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG_PATH)
    parser.add_argument("--instance-id", help="Export one DingTalk approval instance")
    parser.add_argument(
        "--ledger",
        type=Path,
        default=DEFAULT_DEDUPE_LEDGER_PATH,
        help="Local duplicate invoice ledger JSON for review",
    )
    parser.add_argument("--employee", help="Probe filter term, for example 蔡挺")
    parser.add_argument("--amount", help="Probe filter term, for example 1315.98")
    parser.add_argument("--reviewer-user-id", help="Finance reviewer DingTalk userId")
    parser.add_argument(
        "--reviewer-union-id",
        help="Finance reviewer DingTalk unionId for diagnostics",
    )
    parser.add_argument("--keyword", action="append", default=[], help="Probe keyword")
    parser.add_argument(
        "--comment-keyword",
        action="append",
        default=[],
        help="Probe comment text keyword",
    )
    parser.add_argument("--limit", type=int, default=20, help="Max details to inspect")
    parser.add_argument(
        "--host",
        default="",
        help="Server host for serve-tool or serve-callback",
    )
    parser.add_argument(
        "--port",
        type=int,
        default=0,
        help="Server port for serve-tool or serve-callback",
    )
    parser.add_argument(
        "--path",
        default=DEFAULT_OA_CALLBACK_PATH,
        help="Callback HTTP path for serve-callback",
    )
    parser.add_argument(
        "--callback-token",
        default="",
        help="Optional shared token for OA automation HTTP requests",
    )
    parser.add_argument(
        "--skip-reviewed",
        action="store_true",
        help="Skip poll-once approvals that already have a local review report",
    )
    parser.add_argument(
        "--auto-scan-interval",
        type=int,
        default=0,
        help="Seconds between background scans for serve-tool; 0 disables it",
    )
    parser.add_argument(
        "--auto-scan-limit",
        type=int,
        default=20,
        help="Max approval details per serve-tool background scan",
    )
    parser.add_argument(
        "--auto-scan-skip-reviewed",
        action="store_true",
        help="Skip already reviewed approvals in serve-tool background scans",
    )
    parser.add_argument(
        "--stream-events",
        action="store_true",
        help="Start DingTalk event Stream listener with serve-tool",
    )
    args = parser.parse_args()

    try:
        config = load_config(args.config)
        ledger_path = (
            args.ledger if args.ledger.is_absolute() else SCRIPT_DIR / args.ledger
        )
        if args.command == "capabilities":
            report = run_capability_probe(config, ledger_path=ledger_path)
            output = sys.stdout if report["ok"] else sys.stderr
            print(json.dumps(report, ensure_ascii=False, indent=2), file=output)
            return 0 if report["ok"] else 1
        if args.command == "serve-tool":
            run_finance_tool_server(
                config,
                host=args.host or "0.0.0.0",
                port=args.port or 6192,
                ledger_path=ledger_path,
                auto_scan_interval=args.auto_scan_interval,
                auto_scan_limit=args.auto_scan_limit,
                auto_scan_skip_reviewed=args.auto_scan_skip_reviewed,
                stream_events=args.stream_events or config.finance_stream_enabled,
            )
            return 0
        if args.command == "serve-callback":
            callback_token = args.callback_token or config.oa_callback_token
            run_oa_callback_server(
                config,
                host=args.host or "127.0.0.1",
                port=args.port or 6193,
                path=args.path,
                ledger_path=ledger_path,
                callback_token=callback_token,
            )
            return 0
        if args.command == "review":
            if not args.instance_id:
                raise ValueError("review requires --instance-id")
            report = run_initial_review(config, args.instance_id, ledger_path)
            print(json.dumps(report, ensure_ascii=False, indent=2))
            return 0
        if args.command == "poll-once":
            terms = [
                term
                for term in [
                    args.employee,
                    args.amount,
                    *args.keyword,
                    *args.comment_keyword,
                ]
                if term
            ]
            report = run_poll_once(
                config,
                ledger_path=ledger_path,
                terms=terms,
                limit=args.limit,
                skip_reviewed=args.skip_reviewed,
            )
            print(json.dumps(report, ensure_ascii=False, indent=2))
            return 0
        if args.command == "reviewer-queue":
            reviewer_user_id = args.reviewer_user_id or config.finance_reviewer_user_id
            if not reviewer_user_id:
                raise ValueError(
                    "reviewer-queue requires --reviewer-user-id or "
                    "DINGTALK_FINANCE_REVIEWER_USER_ID"
                )
            terms = [
                term
                for term in [
                    args.employee,
                    args.amount,
                    *args.keyword,
                    *args.comment_keyword,
                ]
                if term
            ]
            report = run_reviewer_queue_once(
                config,
                reviewer_user_id=reviewer_user_id,
                reviewer_union_id=(
                    args.reviewer_union_id or config.finance_reviewer_union_id
                ),
                ledger_path=ledger_path,
                terms=terms,
                limit=args.limit,
                skip_reviewed=args.skip_reviewed,
            )
            print(json.dumps(report, ensure_ascii=False, indent=2))
            return 0
        if args.command in {"probe", "users"}:
            client = DingTalkFinanceClient(config)
            token = client.get_access_token()
            ids = (
                [args.instance_id]
                if args.instance_id
                else client.list_approval_instance_ids(token)
            )
            terms = [
                term
                for term in [
                    args.employee,
                    args.amount,
                    *args.keyword,
                    *args.comment_keyword,
                ]
                if term
            ]
            if args.command == "users":
                matched_users: list[dict[str, Any]] = []
                inspected = 0
                for current_instance_id in ids[: args.limit]:
                    detail = client.get_approval_detail(token, current_instance_id)
                    inspected += 1
                    if terms and not detail_matches_terms(detail, terms):
                        continue
                    matched_users.append(
                        {
                            "instance_id": current_instance_id,
                            "users": resolve_approval_users(client, token, detail),
                        }
                    )
                print(
                    json.dumps(
                        {
                            "ok": True,
                            "inspected_count": inspected,
                            "matched_count": len(matched_users),
                            "matched": matched_users,
                        },
                        ensure_ascii=False,
                        indent=2,
                    )
                )
                return 0
            if terms or args.instance_id:
                stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                report = probe_approval_details(
                    client,
                    token,
                    ids,
                    config.output_dir / f"probe_{stamp}",
                    terms=terms,
                    limit=args.limit,
                )
                print(json.dumps(report, ensure_ascii=False, indent=2))
                return 0
            print(
                json.dumps(
                    {"ok": True, "count": len(ids), "ids": ids[:20]},
                    ensure_ascii=False,
                    indent=2,
                )
            )
            return 0
        workbook_path = run_export(config, args.instance_id)
        print(
            json.dumps(
                {"ok": True, "workbook": str(workbook_path)},
                ensure_ascii=False,
                indent=2,
            )
        )
        return 0
    except Exception as exc:
        print(
            json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False, indent=2),
            file=sys.stderr,
        )
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
