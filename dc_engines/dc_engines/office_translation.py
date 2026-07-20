"""Translate Office files through Qwen-MT while preserving editable structure."""

from __future__ import annotations

import hashlib
import json
import re
import unicodedata
import uuid
import zipfile
from copy import copy
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from dc_engines.office_document_delivery import generate_office_deliverables

MAX_TRANSLATION_CHARACTERS = 500_000
MAX_TRANSLATION_UNITS = 20_000
MAX_BATCH_CHARACTERS = 8_000
SUPPORTED_SOURCE_SUFFIXES = {".docx", ".xlsx", ".txt", ".md", ".markdown", ".pdf"}
SOURCE_DELIVERY_FORMATS = {
    ".docx": {"docx", "pdf", "txt", "feishu_doc"},
    ".xlsx": {"xlsx"},
    ".txt": {"docx", "pdf", "txt", "feishu_doc"},
    ".md": {"docx", "pdf", "txt", "feishu_doc"},
    ".markdown": {"docx", "pdf", "txt", "feishu_doc"},
    ".pdf": {"docx", "pdf", "txt", "feishu_doc"},
}
_LANGUAGE_NAMES = {
    "auto",
    "Chinese",
    "English",
    "Japanese",
    "Korean",
    "French",
    "German",
    "Spanish",
    "Portuguese",
    "Russian",
    "Arabic",
    "Thai",
    "Vietnamese",
    "Indonesian",
    "Italian",
}


class TranslationError(ValueError):
    """A translation failure with a stable code and user recovery path.

    Args:
        code: Stable machine-readable failure code.
        message: Concise user-visible failure detail.
        recovery: Concrete action that can recover the task.
        locations: Optional source locations affected by a partial failure.
    """

    def __init__(
        self,
        code: str,
        message: str,
        recovery: str,
        *,
        locations: list[str] | None = None,
    ) -> None:
        super().__init__(message)
        self.code = code
        self.recovery = recovery
        self.locations = locations or []


@dataclass(slots=True)
class TranslationOptions:
    """Validated controls for one document translation run."""

    source_language: str
    target_language: str
    direction_mode: str
    translation_style: str
    output_mode: str
    professional_domain: str
    layout_mode: str
    output_format: str
    model_choice: str
    terms: list[dict[str, str]] = field(default_factory=list)
    translation_memory: list[dict[str, str]] = field(default_factory=list)


@dataclass(slots=True)
class TranslationUnit:
    """One independently traceable translatable source fragment."""

    unit_id: str
    text: str
    location: str


@dataclass(slots=True)
class TranslationResult:
    """Artifacts and trace data produced by one completed translation."""

    run_id: str
    model: str
    artifacts: dict[str, Path]
    manifest_path: Path
    delivery_text: str


def _parse_pairs(value: str, *, field_name: str) -> list[dict[str, str]]:
    """Parse compact glossary or translation-memory pairs.

    Args:
        value: Newline or semicolon separated ``source=target`` pairs.
        field_name: User-visible field name used in validation errors.

    Returns:
        Ordered Qwen-MT source-target dictionaries.

    Raises:
        TranslationError: If a non-empty row has no usable pair separator.
    """
    pairs: list[dict[str, str]] = []
    for raw_row in re.split(r"[\n；;]+", str(value or "")):
        row = raw_row.strip()
        if not row:
            continue
        match = re.match(r"^(.+?)(?:\t|=>|->|=|＝)(.+)$", row)
        if not match or not match.group(1).strip() or not match.group(2).strip():
            raise TranslationError(
                "invalid_terminology",
                f"{field_name}格式无效：{row[:60]}",
                "请按“原文=译文”逐行填写后重试。",
            )
        pairs.append(
            {
                "source": match.group(1).strip()[:200],
                "target": match.group(2).strip()[:200],
            }
        )
        if len(pairs) > 100:
            raise TranslationError(
                "terminology_too_large",
                f"{field_name}最多支持 100 组词条",
                "请拆分词表，保留本次文件实际使用的词条后重试。",
            )
    return pairs


def build_translation_options(task_data: dict[str, Any]) -> TranslationOptions:
    """Validate workbench fields and build deterministic translation controls.

    Args:
        task_data: Sanitized file-workbench task fields.

    Returns:
        Validated translation options.

    Raises:
        TranslationError: If a required or enumerated control is invalid.
    """
    source_language = str(task_data.get("source_language") or "auto").strip()
    target_language = str(task_data.get("target_language") or "").strip()
    if not target_language:
        raise TranslationError(
            "target_language_required",
            "请选择目标语言",
            "返回第 2 步选择目标语言后重新执行。",
        )
    if (
        source_language not in _LANGUAGE_NAMES
        or target_language not in _LANGUAGE_NAMES - {"auto"}
    ):
        raise TranslationError(
            "unsupported_language",
            "源语言或目标语言不受支持",
            "请选择工作台语言列表中的语种后重试。",
        )
    direction_mode = str(task_data.get("direction_mode") or "one_way")
    if direction_mode not in {"one_way", "bidirectional"}:
        raise TranslationError(
            "invalid_direction",
            "翻译方向不受支持",
            "请选择单向翻译或双向互译。",
        )
    if direction_mode == "bidirectional" and source_language == "auto":
        raise TranslationError(
            "language_pair_required",
            "双向互译需要明确两种语言",
            "为双向互译选择明确的源语言和目标语言后重试。",
        )
    if source_language == target_language:
        raise TranslationError(
            "same_language_pair",
            "源语言和目标语言不能相同",
            "请选择不同的源语言和目标语言。",
        )
    translation_style = str(task_data.get("translation_style") or "faithful")
    if translation_style not in {"faithful", "professional"}:
        raise TranslationError(
            "invalid_style",
            "翻译方式不受支持",
            "请选择忠实翻译或专业润色。",
        )
    output_mode = str(task_data.get("output_mode") or "target_only")
    if output_mode not in {"target_only", "bilingual"}:
        raise TranslationError(
            "invalid_output_mode",
            "译文呈现方式不受支持",
            "请选择仅译文或中外文对照。",
        )
    if direction_mode == "bidirectional":
        output_mode = "bilingual"
    domain = str(task_data.get("professional_domain") or "general").strip()
    if domain not in {
        "general",
        "legal",
        "finance",
        "technology",
        "medical",
        "marketing",
        "manufacturing",
    }:
        raise TranslationError(
            "invalid_domain",
            "专业领域不受支持",
            "请选择通用、法律、金融、科技、医疗、营销或制造领域。",
        )
    layout_mode = str(task_data.get("preserve_layout") or "preserve")
    if layout_mode not in {"preserve", "business_clean"}:
        raise TranslationError(
            "invalid_layout",
            "译文排版方式不受支持",
            "请选择保留原版面或商务简洁排版。",
        )
    output_format = str(task_data.get("output_format") or "").strip()
    if output_format not in {"docx", "xlsx", "pdf", "txt", "feishu_doc"}:
        raise TranslationError(
            "unsupported_delivery",
            "请选择可生成的交付格式",
            "返回第 3 步重新选择交付文件。",
        )
    model_choice = str(task_data.get("model_choice") or "auto")
    if model_choice not in {"auto", "fast", "balanced", "quality_first"}:
        raise TranslationError(
            "invalid_model_choice",
            "模型模式不受支持",
            "请选择智能、快速、均衡或质量优先。",
        )
    return TranslationOptions(
        source_language=source_language,
        target_language=target_language,
        direction_mode=direction_mode,
        translation_style=translation_style,
        output_mode=output_mode,
        professional_domain=domain,
        layout_mode=layout_mode,
        output_format=output_format,
        model_choice=model_choice,
        terms=_parse_pairs(str(task_data.get("glossary") or ""), field_name="术语表"),
        translation_memory=_parse_pairs(
            str(task_data.get("translation_memory") or ""),
            field_name="参考译法",
        ),
    )


def choose_translation_model(options: TranslationOptions) -> str:
    """Choose the real AIHubMix Qwen-MT model for one run.

    Args:
        options: Validated translation controls.

    Returns:
        AIHubMix model ID without a provider prefix.
    """
    quality_required = (
        options.translation_style == "professional"
        or options.professional_domain != "general"
        or bool(options.terms)
        or bool(options.translation_memory)
        or options.direction_mode == "bidirectional"
        or options.model_choice == "quality_first"
    )
    return "qwen-mt-plus" if quality_required else "qwen-mt-turbo"


def _contains_language_text(value: str) -> bool:
    """Return whether a fragment contains letters worth translating.

    Args:
        value: Source fragment.

    Returns:
        True for Unicode letter-bearing text, otherwise False.
    """
    return any(unicodedata.category(character).startswith("L") for character in value)


def _extract_document(
    source_path: Path,
) -> tuple[list[TranslationUnit], Any, list[str]]:
    """Extract editable translation units and retain the source object.

    Args:
        source_path: Local supported source document.

    Returns:
        Translation units, the editable source object, and extraction warnings.

    Raises:
        TranslationError: If the source is encrypted, scanned, corrupt, or partial.
    """
    suffix = source_path.suffix.lower()
    units: list[TranslationUnit] = []
    warnings: list[str] = []
    try:
        if suffix == ".docx":
            with zipfile.ZipFile(source_path) as archive:
                names = set(archive.namelist())
                if "word/vbaProject.bin" in names:
                    raise TranslationError(
                        "unsupported_format",
                        "暂不处理带宏的 Word 文件",
                        "请另存为不含宏的 .docx 后重试。",
                    )
            document = Document(source_path)
            paragraphs: list[Any] = list(document.paragraphs)
            tables = list(document.tables)
            for section in document.sections:
                paragraphs.extend(section.header.paragraphs)
                paragraphs.extend(section.footer.paragraphs)
                tables.extend(section.header.tables)
                tables.extend(section.footer.tables)
            table_index = 0
            while table_index < len(tables):
                table = tables[table_index]
                for row in table.rows:
                    for cell in row.cells:
                        paragraphs.extend(cell.paragraphs)
                        tables.extend(cell.tables)
                table_index += 1
            seen_paragraphs: set[int] = set()
            paragraph_units: dict[str, Any] = {}
            for index, paragraph in enumerate(paragraphs):
                paragraph_identity = id(paragraph._p)
                if paragraph_identity in seen_paragraphs:
                    continue
                seen_paragraphs.add(paragraph_identity)
                text = paragraph.text.strip()
                if text and _contains_language_text(text):
                    units.append(
                        TranslationUnit(
                            unit_id=f"docx:{len(units)}",
                            text=text,
                            location=f"Word 段落 {index + 1}",
                        )
                    )
                    paragraph_units[units[-1].unit_id] = paragraph
            return (
                units,
                {"document": document, "paragraphs": paragraph_units},
                warnings,
            )

        if suffix == ".xlsx":
            with zipfile.ZipFile(source_path) as archive:
                risky_prefixes = ("xl/pivotTables/", "xl/externalLinks/")
                risky_files = {"xl/vbaProject.bin", "xl/connections.xml"}
                names = set(archive.namelist())
                if any(
                    name.startswith(risky_prefixes) for name in names
                ) or names.intersection(risky_files):
                    raise TranslationError(
                        "unsupported_workbook_feature",
                        "Excel 含数据透视表、外部链接、连接或宏，当前无法安全回写",
                        "请另存一份仅保留值、公式和基础格式的 .xlsx 后重试。",
                    )
            workbook = load_workbook(source_path, data_only=False, read_only=False)
            cell_units: dict[str, tuple[str, str]] = {}
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        value = cell.value
                        if (
                            isinstance(value, str)
                            and not value.startswith("=")
                            and _contains_language_text(value)
                        ):
                            units.append(
                                TranslationUnit(
                                    unit_id=f"xlsx:{len(units)}",
                                    text=value,
                                    location=f"工作表 {sheet.title}!{cell.coordinate}",
                                )
                            )
                            cell_units[units[-1].unit_id] = (
                                sheet.title,
                                cell.coordinate,
                            )
            return (
                units,
                {"workbook": workbook, "cells": cell_units},
                warnings,
            )

        if suffix in {".txt", ".md", ".markdown"}:
            source_text = ""
            for encoding in ("utf-8-sig", "utf-16", "gb18030"):
                try:
                    source_text = source_path.read_text(encoding=encoding)
                    break
                except UnicodeError:
                    continue
            if not source_text:
                raise TranslationError(
                    "text_encoding",
                    "文本文件为空或编码无法识别",
                    "请另存为 UTF-8 文本后重试。",
                )
            lines = source_text.splitlines(keepends=True)
            for index, line in enumerate(lines):
                bare = line.rstrip("\r\n")
                if bare.strip() and _contains_language_text(bare):
                    units.append(
                        TranslationUnit(
                            unit_id=f"text:{index}",
                            text=bare,
                            location=f"第 {index + 1} 行",
                        )
                    )
            return units, lines, warnings

        if suffix == ".pdf":
            reader = PdfReader(str(source_path))
            if reader.is_encrypted:
                try:
                    unlocked = bool(reader.decrypt(""))
                except Exception:
                    unlocked = False
                if not unlocked:
                    raise TranslationError(
                        "password_protected",
                        "PDF 已加密，无法读取正文",
                        "请移除密码保护或导出未加密 PDF 后重试。",
                    )
            pages: list[str] = []
            failed_pages: list[str] = []
            for index, page in enumerate(reader.pages):
                try:
                    text = str(page.extract_text() or "").strip()
                except Exception:
                    text = ""
                    failed_pages.append(f"第 {index + 1} 页")
                if not text:
                    try:
                        has_images = bool(list(page.images))
                    except Exception:
                        has_images = False
                    if has_images and f"第 {index + 1} 页" not in failed_pages:
                        failed_pages.append(f"第 {index + 1} 页")
                pages.append(text)
                if text and _contains_language_text(text):
                    units.append(
                        TranslationUnit(
                            unit_id=f"pdf:{index}",
                            text=text,
                            location=f"第 {index + 1} 页",
                        )
                    )
            if failed_pages:
                raise TranslationError(
                    "partial_pdf_failure",
                    "PDF 部分页读取失败",
                    "请重新导出失败页，或拆分 PDF 后重新上传。",
                    locations=failed_pages,
                )
            if not units:
                raise TranslationError(
                    "scanned_pdf",
                    "PDF 没有可提取文字，可能是扫描件或图片型 PDF",
                    "请先完成 OCR，或上传可搜索文字的 PDF；图片翻译将在独立能力中提供。",
                )
            return units, pages, warnings
    except TranslationError:
        raise
    except (OSError, ValueError, zipfile.BadZipFile) as exc:
        raise TranslationError(
            "corrupt_document",
            "文件损坏或内容与扩展名不一致",
            "请用原应用重新打开并另存后再上传。",
        ) from exc
    raise TranslationError(
        "unsupported_format",
        f"暂不支持 {suffix or '无扩展名'} 文件翻译",
        "第一期请上传 .docx、.xlsx、.txt、.md 或可搜索文字的 .pdf。",
    )


def _translation_options_payload(
    options: TranslationOptions,
    *,
    source_language: str | None = None,
    target_language: str | None = None,
) -> dict[str, Any]:
    """Build the provider-native Qwen-MT translation options.

    Args:
        options: Validated workbench controls.
        source_language: Optional per-batch direction override.
        target_language: Optional per-batch direction override.

    Returns:
        Qwen-MT ``translation_options`` payload.
    """
    payload: dict[str, Any] = {
        "source_lang": source_language or options.source_language,
        "target_lang": target_language or options.target_language,
    }
    if options.translation_style == "professional":
        payload["domains"] = (
            options.professional_domain
            if options.professional_domain != "general"
            else "general professional business writing"
        )
    elif options.professional_domain != "general":
        payload["domains"] = options.professional_domain
    if options.terms:
        payload["terms"] = options.terms
    if options.translation_memory:
        payload["tm_list"] = options.translation_memory
    return payload


async def _bidirectional_pair_for_text(
    provider: Any,
    text: str,
    options: TranslationOptions,
) -> tuple[str, str]:
    """Infer one side of an explicitly selected bidirectional pair.

    Args:
        provider: AstrBot AIHubMix-compatible provider used for ambiguous pairs.
        text: Translation unit text.
        options: Validated explicit language pair.

    Returns:
        Source and target language for the unit.

    Raises:
        TranslationError: If the unit direction cannot be classified safely.
    """
    source = options.source_language
    target = options.target_language

    def script_count(language: str) -> int | None:
        if language == "Chinese":
            return sum("\u3400" <= character <= "\u9fff" for character in text)
        if language == "Japanese":
            return sum(
                "\u3040" <= character <= "\u30ff" or "\u31f0" <= character <= "\u31ff"
                for character in text
            )
        if language == "Korean":
            return sum("\uac00" <= character <= "\ud7af" for character in text)
        if language == "Russian":
            return sum("\u0400" <= character <= "\u052f" for character in text)
        if language == "Arabic":
            return sum("\u0600" <= character <= "\u06ff" for character in text)
        if language == "Thai":
            return sum("\u0e00" <= character <= "\u0e7f" for character in text)
        if language in {
            "English",
            "French",
            "German",
            "Spanish",
            "Portuguese",
            "Vietnamese",
            "Indonesian",
            "Italian",
        }:
            return sum(
                "LATIN" in unicodedata.name(character, "")
                and unicodedata.category(character).startswith("L")
                for character in text
            )
        return None

    source_count = script_count(source)
    target_count = script_count(target)
    if source_count is not None and target_count is not None:
        if source_count and not target_count:
            return (source, target)
        if target_count and not source_count:
            return (target, source)
        if source_count and target_count and source_count != target_count:
            total = source_count + target_count
            minority_ratio = min(source_count, target_count) / total
            if minority_ratio < 0.2:
                return (
                    (source, target)
                    if source_count > target_count
                    else (target, source)
                )
            raise TranslationError(
                "language_direction_failed",
                "同一翻译单元严重混合了语言对两侧内容",
                "请把混合段落拆开，或改用单向翻译后重试。",
            )

    classifier_prompt = (
        "Classify the language of the text using exactly one allowed label. "
        f"Allowed labels: {source}, {target}, unknown. "
        'Return JSON only: {"language":"<label>"}.\n\nText:\n' + text[:4000]
    )
    try:
        response = await provider.text_chat(
            prompt=classifier_prompt,
            contexts=[],
            model="qwen3.6-flash",
            request_max_retries=2,
        )
    except Exception as exc:
        raise TranslationError(
            "language_direction_failed",
            "双向互译语言方向识别失败",
            "请手动拆分为两个单向翻译任务后重试。",
        ) from exc
    completion = str(getattr(response, "completion_text", "") or "").strip()
    detected_language = "unknown"
    try:
        payload = json.loads(completion)
        detected_language = str(payload.get("language") or "unknown")
    except (AttributeError, json.JSONDecodeError):
        if completion in {source, target, "unknown"}:
            detected_language = completion
    if detected_language == source:
        return (source, target)
    if detected_language == target:
        return (target, source)
    raise TranslationError(
        "language_direction_failed",
        "部分内容无法判断双向互译方向或含第三种语言",
        "请改用单向翻译，或把两种语言内容拆成两个文件分别翻译。",
    )


async def _translate_units(
    provider: Any,
    units: list[TranslationUnit],
    options: TranslationOptions,
    model: str,
) -> tuple[dict[str, str], dict[str, tuple[str, str]]]:
    """Translate bounded batches and validate every returned unit marker.

    Args:
        provider: AstrBot AIHubMix-compatible chat provider.
        units: Ordered source units.
        options: Validated translation controls.
        model: Qwen-MT model ID.

    Returns:
        Translated text and detected source-target directions keyed by stable unit ID.

    Raises:
        TranslationError: If the provider fails or returns an incomplete batch.
    """
    work_units: list[TranslationUnit] = []
    parent_fragments: dict[str, list[str]] = {}
    fragment_limit = MAX_BATCH_CHARACTERS - 512
    for unit in units:
        if len(unit.text) <= fragment_limit:
            work_units.append(unit)
            parent_fragments[unit.unit_id] = [unit.unit_id]
            continue
        parent_fragments[unit.unit_id] = []
        start = 0
        part_index = 0
        while start < len(unit.text):
            end = min(len(unit.text), start + fragment_limit)
            if end < len(unit.text):
                window = unit.text[start:end]
                boundaries = [
                    window.rfind(separator)
                    for separator in ("\n", "。", "！", "？", ". ", "! ", "? ", "; ")
                ]
                boundary = max(boundaries)
                if boundary >= fragment_limit // 2:
                    end = start + boundary + 1
            fragment_id = f"{unit.unit_id}:part:{part_index}"
            work_units.append(
                TranslationUnit(
                    unit_id=fragment_id,
                    text=unit.text[start:end],
                    location=f"{unit.location}（分片 {part_index + 1}）",
                )
            )
            parent_fragments[unit.unit_id].append(fragment_id)
            start = end
            part_index += 1

    directional_groups: dict[tuple[str, str], list[TranslationUnit]] = {}
    fragment_directions: dict[str, tuple[str, str]] = {}
    for unit in work_units:
        if options.direction_mode == "bidirectional":
            direction = await _bidirectional_pair_for_text(provider, unit.text, options)
        else:
            direction = (options.source_language, options.target_language)
        fragment_directions[unit.unit_id] = direction
        directional_groups.setdefault(direction, []).append(unit)

    translations: dict[str, str] = {}
    for (
        source_language,
        target_language,
    ), directional_units in directional_groups.items():
        batches: list[list[TranslationUnit]] = []
        current_batch: list[TranslationUnit] = []
        current_size = 0
        for unit in directional_units:
            marker_size = len(unit.unit_id) * 2 + 64
            unit_size = len(unit.text) + marker_size
            if current_batch and current_size + unit_size > MAX_BATCH_CHARACTERS:
                batches.append(current_batch)
                current_batch = []
                current_size = 0
            current_batch.append(unit)
            current_size += unit_size
        if current_batch:
            batches.append(current_batch)

        for batch in batches:
            prompt = "\n".join(
                f"<<<DC_TRANSLATION_UNIT:{unit.unit_id}>>>\n{unit.text}\n"
                f"<<<DC_TRANSLATION_END:{unit.unit_id}>>>"
                for unit in batch
            )
            try:
                response = await provider.text_chat(
                    prompt=prompt,
                    contexts=[],
                    model=model,
                    translation_options=_translation_options_payload(
                        options,
                        source_language=source_language,
                        target_language=target_language,
                    ),
                    request_max_retries=3,
                )
            except Exception as exc:
                raise TranslationError(
                    "translation_provider_failed",
                    "翻译模型调用失败",
                    "任务未生成不完整文件；请检查 AIHubMix Provider 后重新执行。",
                    locations=[unit.location for unit in batch],
                ) from exc
            completion = str(getattr(response, "completion_text", "") or "")
            for unit in batch:
                pattern = re.compile(
                    rf"<<<DC_TRANSLATION_UNIT:{re.escape(unit.unit_id)}>>>\s*"
                    rf"(.*?)\s*<<<DC_TRANSLATION_END:{re.escape(unit.unit_id)}>>>",
                    re.DOTALL,
                )
                match = pattern.search(completion)
                if match and match.group(1).strip():
                    translations[unit.unit_id] = match.group(1).strip()

            missing = [unit for unit in batch if unit.unit_id not in translations]
            for unit in missing:
                try:
                    retry = await provider.text_chat(
                        prompt=unit.text,
                        contexts=[],
                        model=model,
                        translation_options=_translation_options_payload(
                            options,
                            source_language=source_language,
                            target_language=target_language,
                        ),
                        request_max_retries=2,
                    )
                    translated = str(
                        getattr(retry, "completion_text", "") or ""
                    ).strip()
                except Exception:
                    translated = ""
                if translated:
                    translations[unit.unit_id] = translated
            unresolved = [
                unit.location for unit in batch if unit.unit_id not in translations
            ]
            if unresolved:
                raise TranslationError(
                    "partial_translation_failure",
                    "部分内容翻译失败，未生成不完整交付文件",
                    "请按提示位置拆分源文件后重新生成。",
                    locations=unresolved,
                )
    assembled: dict[str, str] = {}
    detected_directions: dict[str, tuple[str, str]] = {}
    joiner = "" if options.target_language in {"Chinese", "Japanese"} else " "
    for unit in units:
        fragment_ids = parent_fragments[unit.unit_id]
        assembled[unit.unit_id] = joiner.join(
            translations[fragment_id] for fragment_id in fragment_ids
        )
        directions = {fragment_directions[fragment_id] for fragment_id in fragment_ids}
        if len(directions) != 1:
            raise TranslationError(
                "language_direction_failed",
                "同一超长翻译单元被识别为多个语言方向",
                "请把该段落或单元格按语言拆开后重新生成。",
                locations=[unit.location],
            )
        detected_directions[unit.unit_id] = next(iter(directions))
    return assembled, detected_directions


def _render_translation(
    source_path: Path,
    source_object: Any,
    units: list[TranslationUnit],
    translations: dict[str, str],
    options: TranslationOptions,
    output_dir: Path,
) -> tuple[dict[str, Path], str]:
    """Render translated content into a truthful selected delivery format.

    Args:
        source_path: Original local document path.
        source_object: Editable object retained during extraction.
        units: Ordered source units.
        translations: Validated translated text by unit ID.
        options: Validated translation controls.
        output_dir: Runtime run directory.

    Returns:
        Local artifacts and text suitable for a native Feishu document.

    Raises:
        TranslationError: If the selected delivery cannot be rendered.
    """
    suffix = source_path.suffix.lower()
    safe_stem = re.sub(r"[^\w\-.\u3400-\u9fff]+", "_", source_path.stem).strip("._")
    safe_stem = safe_stem[:80] or "translated"
    delivery_text = "\n\n".join(
        (
            f"{unit.text}\n{translations[unit.unit_id]}"
            if options.output_mode == "bilingual"
            else translations[unit.unit_id]
        )
        for unit in units
    )
    artifacts: dict[str, Path] = {}
    if options.output_format == "feishu_doc":
        return artifacts, delivery_text

    if (
        suffix == ".docx"
        and options.output_format == "docx"
        and options.layout_mode == "preserve"
    ):
        document = source_object["document"]
        unit_by_id = {unit.unit_id: unit for unit in units}
        for unit_id, paragraph in source_object["paragraphs"].items():
            source = unit_by_id[unit_id].text
            translated = translations[unit_id]
            rendered = (
                f"{source}\n{translated}"
                if options.output_mode == "bilingual"
                else translated
            )
            text_nodes = paragraph._p.xpath(".//w:t")
            if text_nodes:
                text_nodes[0].text = rendered
                for text_node in text_nodes[1:]:
                    text_node.text = ""
            else:
                paragraph.add_run(rendered)
        output_path = output_dir / f"{safe_stem}_{options.target_language}.docx"
        document.save(output_path)
        artifacts["docx"] = output_path
        return artifacts, delivery_text

    if suffix == ".xlsx" and options.output_format == "xlsx":
        workbook = source_object["workbook"]
        unit_by_id = {unit.unit_id: unit for unit in units}
        for unit_id, (sheet_name, coordinate) in source_object["cells"].items():
            cell = workbook[sheet_name][coordinate]
            source = unit_by_id[unit_id].text
            cell.value = (
                f"{source}\n{translations[unit_id]}"
                if options.output_mode == "bilingual"
                else translations[unit_id]
            )
            if options.output_mode == "bilingual":
                alignment = copy(cell.alignment)
                alignment.wrap_text = True
                cell.alignment = alignment
        output_path = output_dir / f"{safe_stem}_{options.target_language}.xlsx"
        workbook.save(output_path)
        artifacts["xlsx"] = output_path
        return artifacts, delivery_text

    if options.output_format == "txt":
        if suffix in {".txt", ".md", ".markdown"}:
            lines = list(source_object)
            unit_by_index = {int(unit.unit_id.split(":", 1)[1]): unit for unit in units}
            rendered_lines: list[str] = []
            for index, line in enumerate(lines):
                unit = unit_by_index.get(index)
                if unit is None:
                    rendered_lines.append(line)
                    continue
                newline = (
                    "\r\n"
                    if line.endswith("\r\n")
                    else "\n"
                    if line.endswith("\n")
                    else ""
                )
                rendered = translations[unit.unit_id]
                if options.output_mode == "bilingual":
                    rendered = f"{unit.text}{newline or chr(10)}{rendered}"
                rendered_lines.append(rendered + newline)
            delivery_text = "".join(rendered_lines)
        output_path = output_dir / f"{safe_stem}_{options.target_language}.txt"
        output_path.write_text(delivery_text, encoding="utf-8")
        artifacts["txt"] = output_path
        return artifacts, delivery_text

    try:
        artifacts = generate_office_deliverables(
            "file",
            {
                "file_goal": f"{source_path.name} → {options.target_language}",
                "operation": "translate",
                "output_format": options.output_format,
            },
            delivery_text,
            output_dir=output_dir,
        )
    except (OSError, TypeError, ValueError) as exc:
        raise TranslationError(
            "delivery_generation_failed",
            "译文已完成，但交付文件生成失败",
            "保留当前任务记录，改选另一种可用交付格式后重新生成。",
        ) from exc
    return artifacts, delivery_text


async def translate_office_document(
    provider: Any,
    source_path: Path,
    task_data: dict[str, Any],
    *,
    output_dir: Path,
    run_id: str | None = None,
) -> TranslationResult:
    """Translate one supported document and write its trace manifest.

    Args:
        provider: Configured AstrBot AIHubMix provider.
        source_path: Local original document.
        task_data: Sanitized translation task settings.
        output_dir: Runtime directory dedicated to this run.
        run_id: Optional stable regeneration identifier.

    Returns:
        Real artifact paths, delivery text, and manifest path.

    Raises:
        TranslationError: If validation, extraction, translation, or rendering fails.
    """
    source_path = Path(source_path)
    options = build_translation_options(task_data)
    suffix = source_path.suffix.lower()
    if suffix not in SUPPORTED_SOURCE_SUFFIXES:
        raise TranslationError(
            "unsupported_format",
            f"暂不支持 {suffix or '无扩展名'} 文件翻译",
            "第一期请上传 .docx、.xlsx、.txt、.md 或可搜索文字的 .pdf。",
        )
    if options.output_format not in SOURCE_DELIVERY_FORMATS[suffix]:
        recovery = (
            "Excel 翻译请选择 Excel (.xlsx)，以保留工作表、公式和格式。"
            if suffix == ".xlsx"
            else "返回第 3 步选择当前源文件支持的交付格式。"
        )
        raise TranslationError(
            "unsupported_delivery",
            "所选交付格式不能安全保留当前源文件结构",
            recovery,
        )
    if not source_path.is_file():
        raise TranslationError(
            "source_missing",
            "原文件不存在或已过期",
            "请重新上传原文件后再次生成。",
        )
    if source_path.stat().st_size > 80 * 1024 * 1024:
        raise TranslationError(
            "source_too_large",
            "单个翻译文件不能超过 80 MB",
            "请压缩图片或拆分文档后分批翻译。",
        )

    run_id = run_id or uuid.uuid4().hex
    output_dir.mkdir(parents=True, exist_ok=True)
    manifest_path = output_dir / "manifest.json"
    source_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()
    model = choose_translation_model(options)
    regeneration_of = str(task_data.get("regeneration_of") or "")
    if not regeneration_of:
        previous_runs: list[tuple[str, str]] = []
        for previous_manifest_path in output_dir.parent.glob("*/manifest.json"):
            if previous_manifest_path == manifest_path:
                continue
            try:
                previous_manifest = json.loads(
                    previous_manifest_path.read_text(encoding="utf-8")
                )
            except (OSError, TypeError, json.JSONDecodeError):
                continue
            if (
                previous_manifest.get("status") == "completed"
                and previous_manifest.get("source", {}).get("sha256") == source_hash
                and previous_manifest.get("translation", {}).get("target_language")
                == options.target_language
            ):
                previous_runs.append(
                    (
                        str(
                            previous_manifest.get("completed_at")
                            or previous_manifest.get("created_at")
                            or ""
                        ),
                        str(previous_manifest.get("run_id") or ""),
                    )
                )
        if previous_runs:
            regeneration_of = max(previous_runs)[1]
    manifest: dict[str, Any] = {
        "schema_version": 1,
        "run_id": run_id,
        "created_at": datetime.now(UTC).isoformat(),
        "status": "running",
        "source": {
            "name": source_path.name,
            "path": str(source_path),
            "sha256": source_hash,
            "size_bytes": source_path.stat().st_size,
        },
        "translation": {
            "source_language": options.source_language,
            "target_language": options.target_language,
            "direction_mode": options.direction_mode,
            "style": options.translation_style,
            "domain": options.professional_domain,
            "output_mode": options.output_mode,
            "layout_mode": options.layout_mode,
            "output_format": options.output_format,
            "model": f"aihubmix/{model}",
            "term_count": len(options.terms),
            "translation_memory_count": len(options.translation_memory),
            "glossary_sha256": hashlib.sha256(
                json.dumps(options.terms, ensure_ascii=False, sort_keys=True).encode()
            ).hexdigest(),
            "translation_memory_sha256": hashlib.sha256(
                json.dumps(
                    options.translation_memory,
                    ensure_ascii=False,
                    sort_keys=True,
                ).encode()
            ).hexdigest(),
        },
        "regeneration_of": regeneration_of,
        "workspace_token": str(task_data.get("workspace_token") or ""),
        "artifacts": {},
        "warnings": [],
    }
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    try:
        units, source_object, warnings = _extract_document(source_path)
        if not units:
            raise TranslationError(
                "no_translatable_text",
                "文件中没有可翻译文字",
                "请确认文件包含正文文字，而不是仅有数字、公式或空白内容。",
            )
        character_count = sum(len(unit.text) for unit in units)
        if (
            len(units) > MAX_TRANSLATION_UNITS
            or character_count > MAX_TRANSLATION_CHARACTERS
        ):
            raise TranslationError(
                "source_too_large",
                "文件正文超过第一期单任务处理上限",
                "请将文档拆分为不超过 50 万字符的文件后分批翻译。",
            )
        manifest["unit_count"] = len(units)
        manifest["character_count"] = character_count
        manifest["units"] = [
            {
                "id": unit.unit_id,
                "location": unit.location,
                "source_sha256": hashlib.sha256(unit.text.encode()).hexdigest(),
            }
            for unit in units
        ]
        manifest_path.write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        translations, detected_directions = await _translate_units(
            provider, units, options, model
        )
        for unit_row in manifest["units"]:
            source_language, target_language = detected_directions[unit_row["id"]]
            unit_row["source_language"] = source_language
            unit_row["target_language"] = target_language
        artifacts, delivery_text = _render_translation(
            source_path,
            source_object,
            units,
            translations,
            options,
            output_dir,
        )
        manifest["status"] = "completed"
        manifest["completed_at"] = datetime.now(UTC).isoformat()
        manifest["warnings"] = warnings
        manifest["artifacts"] = {key: str(path) for key, path in artifacts.items()}
        manifest_path.write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        return TranslationResult(
            run_id=run_id,
            model=f"aihubmix/{model}",
            artifacts=artifacts,
            manifest_path=manifest_path,
            delivery_text=delivery_text,
        )
    except TranslationError as exc:
        manifest["status"] = "failed"
        manifest["failed_at"] = datetime.now(UTC).isoformat()
        manifest["error"] = {
            "code": exc.code,
            "message": str(exc),
            "recovery": exc.recovery,
            "locations": exc.locations,
        }
        manifest_path.write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        raise


__all__ = [
    "SOURCE_DELIVERY_FORMATS",
    "TranslationError",
    "TranslationOptions",
    "TranslationResult",
    "build_translation_options",
    "choose_translation_model",
    "translate_office_document",
]
