"""Generate traceable internal and sanitized market quotation files."""

from __future__ import annotations

import json
import re
from decimal import Decimal
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont as PdfTTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from dc_engines.material_quotation import MAX_QUOTATION_IMAGES_PER_ITEM


def generate_material_quotation_deliverables(
    task_data: dict[str, Any],
    *,
    output_dir: Path,
    version: int,
    image_root: Path | None = None,
) -> dict[str, Path]:
    """Generate one internal workbook and two sanitized market documents.

    The workbook is an auditable preparation-team calculation surface. The
    Word and PDF files deliberately contain only the market-facing estimate
    summary, so they can be shared without leaking price sources or margins.

    Args:
        task_data: Server-normalized material quotation fields.
        output_dir: Runtime directory that will contain the generated files.
        version: Positive revision number included in filenames and documents.
        image_root: Mounted NAS root used to resolve validated item image paths.

    Returns:
        Paths keyed by ``internal_xlsx``, ``market_docx``, and ``market_pdf``.

    Raises:
        ValueError: If the normalized quotation has no valid line items.
        OSError: If the output directory or files cannot be written.
    """
    try:
        items = json.loads(str(task_data.get("quotation_items") or "[]"))
    except json.JSONDecodeError as exc:
        raise ValueError("报价物料明细格式无效") from exc
    if (
        not isinstance(items, list)
        or not items
        or not all(isinstance(item, dict) for item in items)
    ):
        raise ValueError("报价物料明细为空")

    output_dir.mkdir(parents=True, exist_ok=True)
    revision = max(1, int(version))
    project_name = str(task_data.get("project_name") or "未命名项目").strip()
    safe_project = re.sub(r"[^\w\u4e00-\u9fff-]+", "-", project_name).strip("-_")
    safe_project = safe_project[:48] or "material-estimate"
    workbook_path = output_dir / f"{safe_project}-成本清单-V{revision}.xlsx"
    document_path = output_dir / f"{safe_project}-市场估价-V{revision}.docx"
    pdf_path = output_dir / f"{safe_project}-市场估价-V{revision}.pdf"

    workbook = Workbook()
    cost_sheet = workbook.active
    cost_sheet.title = "成本清单"
    trace_sheet = workbook.create_sheet("测算溯源")
    excel_font_name = "Arial Unicode MS"
    thin_border = Border(
        left=Side(style="thin", color="D8DEE8"),
        right=Side(style="thin", color="D8DEE8"),
        top=Side(style="thin", color="D8DEE8"),
        bottom=Side(style="thin", color="D8DEE8"),
    )

    trace_sheet.sheet_view.showGridLines = False
    trace_sheet.freeze_panes = "A6"
    trace_sheet.merge_cells("A1:J1")
    trace_sheet["A1"] = f"{project_name}｜筹备组物料测算 V{revision}"
    trace_sheet["A1"].font = Font(
        name=excel_font_name, size=18, bold=True, color="FFFFFF"
    )
    trace_sheet["A1"].fill = PatternFill("solid", fgColor="163A5F")
    trace_sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    trace_sheet.row_dimensions[1].height = 34
    trace_sheet.merge_cells("A2:J2")
    trace_sheet["A2"] = "内部资料｜含历史价格来源与测算参数，禁止对外发送"
    trace_sheet["A2"].font = Font(
        name=excel_font_name, size=10, bold=True, color="8A3B12"
    )
    trace_sheet["A2"].fill = PatternFill("solid", fgColor="FFF2E8")
    trace_sheet["A2"].alignment = Alignment(horizontal="left", vertical="center")
    trace_sheet.merge_cells("A3:E3")
    trace_sheet.merge_cells("F3:J3")
    trace_sheet["A3"] = f"客户 / 品牌：{task_data.get('client_name') or '未填写'}"
    delivery_date = str(task_data.get("delivery_date") or "未填写")
    trace_sheet["F3"] = f"交付日期：{delivery_date}"
    for cell in (trace_sheet["A3"], trace_sheet["F3"]):
        cell.font = Font(name=excel_font_name, size=10, color="3F4B5A")
        cell.alignment = Alignment(vertical="center")

    trace_headers = [
        "序号",
        "物料",
        "规格 / 工艺",
        "数量",
        "单位",
        "单价",
        "小计",
        "供应方",
        "历史来源",
        "价格状态",
    ]
    for column, header in enumerate(trace_headers, start=1):
        cell = trace_sheet.cell(row=5, column=column, value=header)
        cell.font = Font(name=excel_font_name, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2D6A7A")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    trace_sheet.row_dimensions[5].height = 26

    for item_index, item in enumerate(items, start=1):
        row = 5 + item_index
        unit_price = str(item.get("unit_price") or "").strip()
        values: list[Any] = [
            item_index,
            str(item.get("item_name") or ""),
            str(item.get("specification") or ""),
            Decimal(str(item.get("quantity") or "0")),
            str(item.get("unit") or "项"),
            Decimal(unit_price) if unit_price else None,
            f"=D{row}*F{row}" if unit_price else None,
            str(item.get("supplier") or ""),
            str(item.get("source_path") or ""),
            str(item.get("price_status") or "待询价"),
        ]
        for column, value in enumerate(values, start=1):
            cell = trace_sheet.cell(row=row, column=column, value=value)
            cell.font = Font(name=excel_font_name, size=9, color="202A35")
            cell.fill = PatternFill(
                "solid", fgColor="F7FAFC" if item_index % 2 == 0 else "FFFFFF"
            )
            cell.border = thin_border
            cell.alignment = Alignment(
                horizontal="center" if column in {1, 4, 5, 6, 7, 10} else "left",
                vertical="center",
                wrap_text=column in {2, 3, 8, 9, 10},
            )
            if column in {6, 7}:
                cell.number_format = "¥#,##0.00;[Red]-¥#,##0.00"
        trace_sheet.row_dimensions[row].height = 34

    first_item_row = 6
    last_item_row = 5 + len(items)
    summary_start = max(last_item_row + 2, 9)
    summary_rows = {
        summary_start: (
            "已定价物料小计",
            f"=SUM(G{first_item_row}:G{last_item_row})",
        ),
        summary_start + 1: (
            "损耗率",
            Decimal(str(task_data.get("loss_rate") or "0")) / 100,
        ),
        summary_start + 2: (
            "损耗金额",
            f"=B{summary_start}*B{summary_start + 1}",
        ),
        summary_start + 3: (
            "运输费",
            Decimal(str(task_data.get("transport_fee") or "0")),
        ),
        summary_start + 4: (
            "安装费",
            Decimal(str(task_data.get("installation_fee") or "0")),
        ),
        summary_start + 5: (
            "加急费",
            Decimal(str(task_data.get("rush_fee") or "0")),
        ),
        summary_start + 6: (
            "成本基数",
            f"=SUM(B{summary_start},B{summary_start + 2}:B{summary_start + 5})",
        ),
        summary_start + 7: (
            "利润率",
            Decimal(str(task_data.get("profit_rate") or "0")) / 100,
        ),
        summary_start + 8: (
            "利润额",
            f"=B{summary_start + 6}*B{summary_start + 7}",
        ),
        summary_start + 9: (
            "税率",
            Decimal(str(task_data.get("tax_rate") or "0")) / 100,
        ),
        summary_start + 10: (
            "税额",
            f"=(B{summary_start + 6}+B{summary_start + 8})*B{summary_start + 9}",
        ),
        summary_start + 11: (
            "含税估价",
            f"=B{summary_start + 6}+B{summary_start + 8}+B{summary_start + 10}",
        ),
    }
    for row, (label, value) in summary_rows.items():
        trace_sheet.cell(row=row, column=1, value=label)
        trace_sheet.cell(row=row, column=2, value=value)
        for column in range(1, 3):
            cell = trace_sheet.cell(row=row, column=column)
            cell.font = Font(
                name=excel_font_name,
                size=10,
                bold=row == max(summary_rows),
                color="FFFFFF" if row == max(summary_rows) else "24364B",
            )
            cell.fill = PatternFill(
                "solid", fgColor="163A5F" if row == max(summary_rows) else "EEF3F7"
            )
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="right", vertical="center")
        if label in {"损耗率", "利润率", "税率"}:
            trace_sheet.cell(row=row, column=2).number_format = "0.00%"
        else:
            trace_sheet.cell(
                row=row, column=2
            ).number_format = "¥#,##0.00;[Red]-¥#,##0.00"

    trace_widths = [8, 20, 26, 11, 9, 14, 15, 20, 34, 15]
    for column, width in enumerate(trace_widths, start=1):
        trace_sheet.column_dimensions[get_column_letter(column)].width = width
    trace_sheet.auto_filter.ref = f"A5:J{last_item_row}"
    trace_sheet.print_title_rows = "1:5"
    trace_sheet.page_setup.orientation = "landscape"
    trace_sheet.page_setup.fitToWidth = 1
    trace_sheet.page_setup.fitToHeight = 0
    trace_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    trace_sheet.oddFooter.center.text = "筹备组内部测算｜请勿外传"

    # Keep the first worksheet close to the colleague-provided cost-list layout.
    cost_sheet.sheet_view.showGridLines = False
    cost_sheet.freeze_panes = "A4"
    cost_sheet.merge_cells("A1:I1")
    cost_sheet["A1"] = f"{project_name}成本清单"
    cost_sheet["A1"].font = Font(name=excel_font_name, size=18, bold=True)
    cost_sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    cost_sheet.row_dimensions[1].height = 32
    cost_sheet.merge_cells("A2:I2")
    cost_sheet["A2"] = "方案"
    cost_sheet["A2"].font = Font(name=excel_font_name, size=11, bold=True)
    cost_sheet["A2"].alignment = Alignment(horizontal="center", vertical="center")
    cost_sheet.row_dimensions[2].height = 22
    cost_headers = [
        "序号",
        "礼品 / 物料组合",
        "规格",
        "数量",
        "单位",
        "单价",
        "小计",
        "图示",
        "备注",
    ]
    cost_border = Border(
        left=Side(style="thin", color="1F1F1F"),
        right=Side(style="thin", color="1F1F1F"),
        top=Side(style="thin", color="1F1F1F"),
        bottom=Side(style="thin", color="1F1F1F"),
    )
    for column, header in enumerate(cost_headers, start=1):
        cell = cost_sheet.cell(row=3, column=column, value=header)
        cell.font = Font(name=excel_font_name, size=10, bold=True)
        cell.fill = PatternFill("solid", fgColor="BFBFBF")
        cell.border = cost_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    cost_sheet.row_dimensions[3].height = 30

    first_cost_row = 4
    for item_index, item in enumerate(items, start=1):
        row = first_cost_row + item_index - 1
        unit_price = str(item.get("unit_price") or "").strip()
        cost_values: list[Any] = [
            item_index,
            str(item.get("item_name") or ""),
            str(item.get("specification") or "—"),
            Decimal(str(item.get("quantity") or "0")),
            str(item.get("unit") or "项"),
            Decimal(unit_price) if unit_price else None,
            f"=D{row}*F{row}" if unit_price else None,
            None,
            str(item.get("remark") or "") or ("待询价" if not unit_price else ""),
        ]
        for column, value in enumerate(cost_values, start=1):
            cell = cost_sheet.cell(row=row, column=column, value=value)
            cell.font = Font(name=excel_font_name, size=9)
            cell.border = cost_border
            cell.alignment = Alignment(
                horizontal="left" if column in {2, 3, 9} else "center",
                vertical="center",
                wrap_text=column in {2, 3, 9},
            )
            if column in {6, 7}:
                cell.number_format = "¥#,##0.00;[Red]-¥#,##0.00"
        raw_image_paths = item.get("image_paths")
        image_paths = (
            [str(path or "").strip() for path in raw_image_paths]
            if isinstance(raw_image_paths, list)
            else []
        )
        legacy_image_path = str(item.get("image_path") or "").strip()
        if legacy_image_path and legacy_image_path not in image_paths:
            image_paths.insert(0, legacy_image_path)
        image_paths = list(dict.fromkeys(path for path in image_paths if path))
        if len(image_paths) > MAX_QUOTATION_IMAGES_PER_ITEM:
            raise ValueError(f"每项物料最多写入 {MAX_QUOTATION_IMAGES_PER_ITEM} 张图片")
        if image_paths:
            if image_root is None:
                raise ValueError("报价图片缺少 NAS 根目录")
            resolved_root = image_root.resolve()
            image_columns = (
                1 if len(image_paths) == 1 else 2 if len(image_paths) <= 4 else 3
            )
            image_rows = (len(image_paths) + image_columns - 1) // image_columns
            image_gap = 2
            slot_width = (110 - image_gap * (image_columns - 1)) / image_columns
            slot_height = (64 - image_gap * (image_rows - 1)) / image_rows
            for image_index, image_path in enumerate(image_paths):
                image_file = (resolved_root / image_path).resolve()
                try:
                    image_file.relative_to(resolved_root)
                except ValueError as exc:
                    raise ValueError("报价图片路径无效") from exc
                if not image_file.is_file():
                    raise ValueError("报价图片不存在，请重新上传")
                try:
                    excel_image = ExcelImage(image_file)
                except (OSError, ValueError) as exc:
                    raise ValueError("报价图片无法写入 Excel") from exc
                scale = min(
                    slot_width / excel_image.width,
                    slot_height / excel_image.height,
                    1,
                )
                display_width = max(1, round(excel_image.width * scale))
                display_height = max(1, round(excel_image.height * scale))
                image_column = image_index % image_columns
                image_row = image_index // image_columns
                x_offset = round(
                    image_column * (slot_width + image_gap)
                    + (slot_width - display_width) / 2
                )
                y_offset = round(
                    image_row * (slot_height + image_gap)
                    + (slot_height - display_height) / 2
                )
                excel_image.width = display_width
                excel_image.height = display_height
                excel_image.anchor = OneCellAnchor(
                    _from=AnchorMarker(
                        col=7,
                        colOff=pixels_to_EMU(x_offset),
                        row=row - 1,
                        rowOff=pixels_to_EMU(y_offset),
                    ),
                    ext=XDRPositiveSize2D(
                        cx=pixels_to_EMU(display_width),
                        cy=pixels_to_EMU(display_height),
                    ),
                )
                cost_sheet.add_image(excel_image)
            cost_sheet.row_dimensions[row].height = 52
        else:
            cost_sheet.row_dimensions[row].height = 42

    loss_rate_text = format(
        Decimal(str(task_data.get("loss_rate") or "0")).normalize(), "f"
    )
    profit_rate_text = format(
        Decimal(str(task_data.get("profit_rate") or "0")).normalize(), "f"
    )
    tax_rate_text = format(
        Decimal(str(task_data.get("tax_rate") or "0")).normalize(), "f"
    )
    fee_rows = [
        ("损耗费", summary_start + 2, f"损耗率 {loss_rate_text}%"),
        ("运输费", summary_start + 3, "项目运输"),
        ("安装费", summary_start + 4, "现场安装"),
        ("加急费", summary_start + 5, "加急交付"),
        ("利润额", summary_start + 8, f"利润率 {profit_rate_text}%"),
        ("税费", summary_start + 10, f"税率 {tax_rate_text}%"),
    ]
    last_cost_item_row = first_cost_row + len(items) - 1
    for fee_index, (label, trace_row, remark) in enumerate(fee_rows, start=1):
        row = last_cost_item_row + fee_index
        fee_values: list[Any] = [
            len(items) + fee_index,
            label,
            "—",
            1,
            "项",
            f"='测算溯源'!B{trace_row}",
            f"=D{row}*F{row}",
            None,
            remark,
        ]
        for column, value in enumerate(fee_values, start=1):
            cell = cost_sheet.cell(row=row, column=column, value=value)
            cell.font = Font(name=excel_font_name, size=9)
            cell.border = cost_border
            cell.alignment = Alignment(
                horizontal="left" if column in {2, 3, 9} else "center",
                vertical="center",
                wrap_text=column in {2, 3, 9},
            )
            if column in {6, 7}:
                cell.number_format = "¥#,##0.00;[Red]-¥#,##0.00"
        cost_sheet.row_dimensions[row].height = 26

    last_fee_row = last_cost_item_row + len(fee_rows)
    total_row = last_fee_row + 1
    total_values: list[Any] = [
        "—",
        "合计",
        None,
        "—",
        None,
        "—",
        f"=SUM(G{first_cost_row}:G{last_fee_row})",
        "—",
        str(task_data.get("quotation_notes") or "内部资料｜请勿外传"),
    ]
    for column, value in enumerate(total_values, start=1):
        cell = cost_sheet.cell(row=total_row, column=column, value=value)
        cell.font = Font(name=excel_font_name, size=10, bold=column in {2, 7})
        cell.fill = PatternFill("solid", fgColor="E7E6E6")
        cell.border = cost_border
        cell.alignment = Alignment(
            horizontal="left" if column == 9 else "center",
            vertical="center",
            wrap_text=column == 9,
        )
        if column == 7:
            cell.number_format = "¥#,##0.00;[Red]-¥#,##0.00"
    cost_sheet.row_dimensions[total_row].height = 30

    cost_widths = [7, 25, 24, 11, 10, 14, 15, 18, 32]
    for column, width in enumerate(cost_widths, start=1):
        cost_sheet.column_dimensions[get_column_letter(column)].width = width
    cost_sheet.print_title_rows = "1:3"
    cost_sheet.print_area = f"A1:I{total_row}"
    cost_sheet.page_setup.orientation = "landscape"
    cost_sheet.page_setup.paperSize = cost_sheet.PAPERSIZE_A4
    cost_sheet.page_setup.fitToWidth = 1
    cost_sheet.page_setup.fitToHeight = 0
    cost_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    cost_sheet.oddFooter.center.text = "筹备组内部资料｜请勿外传"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(workbook_path)

    amount = Decimal(str(task_data.get("grand_total") or "0")).quantize(Decimal("0.01"))
    client_name = str(task_data.get("client_name") or "未填写")
    validity_days = str(task_data.get("validity_days") or "15")
    estimate_status = str(task_data.get("quotation_status") or "估价方案")
    safe_rows = [
        ("项目", project_name),
        ("客户 / 品牌", client_name),
        ("估价金额", f"¥{amount:,.2f}"),
        ("有效期", f"{validity_days} 天"),
        ("状态", estimate_status),
    ]

    document = Document()
    word_font_name = "Arial Unicode MS"
    normal_style = document.styles["Normal"]
    normal_style.font.name = word_font_name
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), word_font_name)
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("市场估价结果")
    title_run.bold = True
    title_run.font.size = Pt(22)
    title_run.font.color.rgb = RGBColor(22, 58, 95)
    title_run.font.name = word_font_name
    title_run._element.rPr.rFonts.set(qn("w:eastAsia"), word_font_name)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(f"ESTIMATE · V{revision}")
    subtitle_run.font.size = Pt(9)
    subtitle_run.font.color.rgb = RGBColor(99, 115, 129)
    table = document.add_table(rows=0, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(4.2)
    table.columns[1].width = Cm(11.2)
    for index, (label, value) in enumerate(safe_rows):
        cells = table.add_row().cells
        cells[0].width = Cm(4.2)
        cells[1].width = Cm(11.2)
        cells[0].text = label
        cells[1].text = value
        for cell_index, cell in enumerate(cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "EEF3F7" if cell_index == 0 else "FFFFFF")
            cell._tc.get_or_add_tcPr().append(shade)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = word_font_name
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), word_font_name)
                    run.font.size = Pt(10.5)
                    run.bold = cell_index == 0 or index == 2
                    if index == 2 and cell_index == 1:
                        run.font.size = Pt(15)
                        run.font.color.rgb = RGBColor(22, 58, 95)
    document.add_paragraph()
    note = document.add_paragraph("本结果用于市场方案估价，正式金额以最终复核为准。")
    note.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in note.runs:
        run.font.name = word_font_name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), word_font_name)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(99, 115, 129)
    document.save(document_path)

    pdf_font_name = "STSong-Light"
    font_candidates = (
        Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
        Path("/Library/Fonts/Arial Unicode.ttf"),
        Path("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("C:/Windows/Fonts/msyh.ttc"),
    )
    for font_path in font_candidates:
        if not font_path.is_file():
            continue
        try:
            pdfmetrics.registerFont(PdfTTFont("QuotationChinese", str(font_path)))
            pdf_font_name = "QuotationChinese"
            break
        except Exception:  # noqa: BLE001
            continue
    if pdf_font_name == "STSong-Light":
        try:
            pdfmetrics.registerFont(UnicodeCIDFont(pdf_font_name))
        except KeyError:
            pass
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ChineseTitle",
        parent=styles["Title"],
        fontName=pdf_font_name,
        fontSize=22,
        leading=30,
        textColor=colors.HexColor("#163A5F"),
        alignment=TA_CENTER,
        spaceAfter=3 * mm,
    )
    subtitle_style = ParagraphStyle(
        "ChineseSubtitle",
        parent=styles["Normal"],
        fontName=pdf_font_name,
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#637381"),
        alignment=TA_CENTER,
        spaceAfter=9 * mm,
    )
    cell_style = ParagraphStyle(
        "ChineseCell",
        parent=styles["Normal"],
        fontName=pdf_font_name,
        fontSize=10.5,
        leading=16,
        textColor=colors.HexColor("#202A35"),
        alignment=TA_LEFT,
    )
    note_style = ParagraphStyle(
        "ChineseNote",
        parent=cell_style,
        fontSize=9,
        textColor=colors.HexColor("#637381"),
        leading=14,
    )
    pdf_document = SimpleDocTemplate(
        str(pdf_path),
        pagesize=A4,
        rightMargin=22 * mm,
        leftMargin=22 * mm,
        topMargin=24 * mm,
        bottomMargin=24 * mm,
        title=f"{project_name} 市场估价",
        author="巅池文化",
    )
    pdf_table = Table(
        [
            [Paragraph(label, cell_style), Paragraph(value, cell_style)]
            for label, value in safe_rows
        ],
        colWidths=[42 * mm, 112 * mm],
        rowHeights=[13 * mm] * len(safe_rows),
    )
    pdf_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EEF3F7")),
                ("BACKGROUND", (1, 0), (1, -1), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D8DEE8")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 12),
                ("RIGHTPADDING", (0, 0), (-1, -1), 12),
            ]
        )
    )
    pdf_document.build(
        [
            Paragraph("市场估价结果", title_style),
            Paragraph(f"ESTIMATE · V{revision}", subtitle_style),
            pdf_table,
            Spacer(1, 8 * mm),
            Paragraph("本结果用于市场方案估价，正式金额以最终复核为准。", note_style),
        ]
    )

    return {
        "internal_xlsx": workbook_path,
        "market_docx": document_path,
        "market_pdf": pdf_path,
    }


__all__ = ["generate_material_quotation_deliverables"]
