from __future__ import annotations

import importlib.util
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from dc_engines.office_document_delivery import generate_office_deliverables
from openpyxl import load_workbook
from pypdf import PdfReader

from astrbot.api.provider import LLMResponse
from astrbot.core.message.components import File, Plain
from astrbot.core.message.message_event_result import MessageEventResult

ROOT = Path(__file__).resolve().parents[1]
PLUGIN_PATH = ROOT / "data" / "plugins" / "office_document_delivery_plugin" / "main.py"
SPEC = importlib.util.spec_from_file_location(
    "office_document_delivery_plugin_test", PLUGIN_PATH
)
assert SPEC and SPEC.loader
PLUGIN_MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(PLUGIN_MODULE)
OfficeDocumentDeliveryPlugin = PLUGIN_MODULE.OfficeDocumentDeliveryPlugin


class _Event:
    def __init__(
        self, extras: dict[str, object], result: MessageEventResult | None = None
    ):
        self.extras = extras
        self.result = result

    def get_extra(self, key: str):
        return self.extras.get(key)

    def set_extra(self, key: str, value: object) -> None:
        self.extras[key] = value

    def get_result(self) -> MessageEventResult | None:
        return self.result

    def set_result(self, result: MessageEventResult) -> None:
        self.result = result

    def get_platform_id(self) -> str:
        return "巅池-Agent小助手"


def test_research_office_bundle_creates_real_word_excel_and_pdf(tmp_path: Path) -> None:
    result_text = """# 供应商交付能力比较

结论：乙方在时效与售后方面更稳健。

| 供应商 | 交付周期 | 售后响应 | 结论 |
| --- | --- | --- | --- |
| 甲方 | 10 天 | 48 小时 | 成本较低 |
| 乙方 | 7 天 | 12 小时 | 综合优先 |

## 证据与限制
- 数据来自本次上传的供应商方案。
- 报价有效期仍需确认。
"""

    artifacts = generate_office_deliverables(
        "research",
        {
            "research_question": "供应商交付能力比较",
            "analysis_mode": "comparison",
            "delivery_format": "office_bundle",
        },
        result_text,
        output_dir=tmp_path,
    )

    assert set(artifacts) == {"docx", "xlsx", "pdf"}
    assert all(
        path.is_file() and path.stat().st_size > 500 for path in artifacts.values()
    )
    assert artifacts["docx"].read_bytes().startswith(b"PK")
    workbook = load_workbook(artifacts["xlsx"], data_only=True)
    assert workbook.sheetnames == ["分析结果", "任务信息"]
    assert any(
        cell.value == "供应商"
        for row in workbook["分析结果"].iter_rows()
        for cell in row
    )
    pdf = PdfReader(str(artifacts["pdf"]))
    assert len(pdf.pages) >= 1
    assert "供应商交付能力比较" in "".join(
        page.extract_text() or "" for page in pdf.pages
    )


def test_file_delivery_generates_only_selected_text_file(tmp_path: Path) -> None:
    artifacts = generate_office_deliverables(
        "file",
        {
            "file_goal": "提取合同风险",
            "operation": "extract",
            "output_format": "txt",
        },
        "# 风险清单\n- 付款条件不明确\n- 违约责任缺失",
        output_dir=tmp_path,
    )

    assert set(artifacts) == {"txt"}
    assert "付款条件不明确" in artifacts["txt"].read_text(encoding="utf-8")


def test_result_card_and_feishu_doc_do_not_create_fake_local_files(
    tmp_path: Path,
) -> None:
    card = generate_office_deliverables(
        "research",
        {"research_question": "结论", "delivery_format": "result_card"},
        "直接结论",
        output_dir=tmp_path / "card",
    )
    cloud = generate_office_deliverables(
        "file",
        {"file_goal": "改写", "output_format": "feishu_doc"},
        "改写结果",
        output_dir=tmp_path / "cloud",
    )

    assert card == {}
    assert cloud == {}
    assert not (tmp_path / "card").exists()
    assert not (tmp_path / "cloud").exists()


@pytest.mark.asyncio
async def test_plugin_retains_llm_result_then_appends_file_after_card_consumption(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    context = SimpleNamespace()
    plugin = OfficeDocumentDeliveryPlugin(context)
    event = _Event(
        {
            "assistant_workbench_task_type": "file",
            "assistant_workbench_task_data": {
                "file_goal": "总结会议纪要",
                "operation": "summarize",
                "output_format": "docx",
            },
        },
        MessageEventResult(chain=[]),
    )
    monkeypatch.setattr(PLUGIN_MODULE, "get_astrbot_data_path", lambda: str(tmp_path))

    await plugin.retain_office_result(
        event,
        LLMResponse(role="assistant", completion_text="# 会议结论\n- 确认下周交付"),
    )
    await plugin.append_office_deliverables(event)

    assert event.extras["office_delivery_text"].startswith("# 会议结论")
    assert event.result is not None
    files = [
        component for component in event.result.chain if isinstance(component, File)
    ]
    assert len(files) == 1
    assert files[0].name.endswith(".docx")
    assert Path(files[0].file_).is_file()


@pytest.mark.asyncio
async def test_plugin_creates_native_feishu_cloud_document() -> None:
    create = AsyncMock(
        return_value=SimpleNamespace(
            success=lambda: True,
            data=SimpleNamespace(document=SimpleNamespace(document_id="doccn_office")),
            msg="",
        )
    )
    write = AsyncMock(return_value=SimpleNamespace(success=lambda: True, msg=""))
    client = SimpleNamespace(
        docx=SimpleNamespace(
            v1=SimpleNamespace(
                document=SimpleNamespace(acreate=create),
                document_block_children=SimpleNamespace(acreate=write),
            )
        )
    )
    context = SimpleNamespace(
        get_platform_inst=lambda _platform_id: SimpleNamespace(lark_api=client)
    )
    plugin = OfficeDocumentDeliveryPlugin(context)
    event = _Event(
        {
            "assistant_workbench_task_type": "research",
            "assistant_workbench_task_data": {
                "research_question": "市场趋势判断",
                "delivery_format": "feishu_doc",
            },
            "office_delivery_text": "# 结论\n市场需求保持增长。",
        },
        MessageEventResult(chain=[]),
    )

    await plugin.append_office_deliverables(event)

    assert create.await_count == 1
    assert write.await_count == 1
    assert event.result is not None
    assert any(
        isinstance(component, Plain)
        and "https://feishu.cn/docx/doccn_office" in component.text
        for component in event.result.chain
    )
