"""Employee supplier-price intake backed by Obsidian and the NAS index."""

from __future__ import annotations

import csv
import hashlib
import json
import re
import sqlite3
import uuid
from datetime import date, datetime, timezone
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from nas_sync.dc_memory_indexer import chunk_text, init_nas_db

MONEY = Decimal("0.01")
MAX_SUPPLIER_PRICE_ITEMS = 50
MAX_EVIDENCE_BYTES = 20 * 1024 * 1024
ALLOWED_CATEGORIES = {"搭建类", "礼品类", "印刷类", "制作类", "其他类"}


def _rows_from_matrix(matrix: list[list[Any]]) -> list[dict[str, str]]:
    """Extract price rows from one spreadsheet-like matrix.

    Args:
        matrix: Rows containing arbitrary cell values from CSV or Excel.

    Returns:
        At most fifty normalized item, specification, unit, and price rows.
    """
    header_row = -1
    item_index = -1
    price_index = -1
    unit_index = -1
    specification_index = -1
    header_unit = ""
    for row_index, raw_row in enumerate(matrix[:20]):
        headers = [re.sub(r"\s+", "", str(cell or "")) for cell in raw_row]
        item_candidate = next(
            (
                index
                for index, header in enumerate(headers)
                if any(
                    marker in header
                    for marker in ("物料", "材料", "项目", "产品", "品名")
                )
            ),
            -1,
        )
        price_candidate = next(
            (
                index
                for index, header in enumerate(headers)
                if any(marker in header for marker in ("单价", "价格", "报价"))
            ),
            -1,
        )
        if item_candidate < 0 or price_candidate < 0:
            continue
        header_row = row_index
        item_index = item_candidate
        price_index = price_candidate
        unit_index = next(
            (
                index
                for index, header in enumerate(headers)
                if "单位" in header and index != price_index
            ),
            -1,
        )
        specification_index = next(
            (
                index
                for index, header in enumerate(headers)
                if any(
                    marker in header
                    for marker in ("规格", "材质", "工艺", "说明", "备注")
                )
            ),
            -1,
        )
        unit_match = re.search(r"元[/／]([^）)]+)", headers[price_index])
        header_unit = unit_match.group(1).strip() if unit_match else ""
        break
    if header_row < 0:
        return []

    rows: list[dict[str, str]] = []
    for raw_row in matrix[header_row + 1 :]:
        if max(item_index, price_index) >= len(raw_row):
            continue
        item_name = str(raw_row[item_index] or "").strip()
        price_text = str(raw_row[price_index] or "").strip()
        price_match = re.search(r"(?<![\d.])([0-9][0-9,]*(?:\.[0-9]+)?)", price_text)
        if not item_name or price_match is None:
            continue
        try:
            unit_price = Decimal(price_match.group(1).replace(",", ""))
        except InvalidOperation:
            continue
        if unit_price <= 0:
            continue
        specification = (
            str(raw_row[specification_index] or "").strip()
            if 0 <= specification_index < len(raw_row)
            else ""
        )
        unit = (
            str(raw_row[unit_index] or "").strip()
            if 0 <= unit_index < len(raw_row)
            else header_unit
        )
        if not unit:
            cell_unit_match = re.search(r"元[/／]([\u4e00-\u9fff㎡²]+)", price_text)
            unit = cell_unit_match.group(1) if cell_unit_match else ""
        if unit in {"-", "待确认", "待定", "未知"}:
            unit = ""
        rows.append(
            {
                "item_name": item_name[:80],
                "specification": "" if specification == "-" else specification[:160],
                "unit": unit[:16],
                "unit_price": str(unit_price.quantize(MONEY, ROUND_HALF_UP)),
            }
        )
        if len(rows) >= MAX_SUPPLIER_PRICE_ITEMS:
            break
    return rows


def parse_supplier_price_attachment(
    filename: str,
    content: bytes,
) -> list[dict[str, str]]:
    """Parse structured price rows from an uploaded CSV or Excel workbook.

    Args:
        filename: Original source filename used to select the parser.
        content: Complete source bytes from the employee upload.

    Returns:
        Structured price candidates, or an empty list for unstructured files.

    Raises:
        ValueError: If a supported spreadsheet file is malformed.
    """
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        decoded = ""
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                decoded = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if not decoded:
            raise ValueError("CSV 编码无法识别")
        return _rows_from_matrix(list(csv.reader(StringIO(decoded))))
    if suffix in {".xlsx", ".xlsm"}:
        try:
            workbook = load_workbook(
                BytesIO(content),
                read_only=True,
                data_only=True,
            )
        except Exception as exc:  # noqa: BLE001
            raise ValueError("Excel 报价资料无法读取") from exc
        rows: list[dict[str, str]] = []
        try:
            for worksheet in workbook.worksheets:
                matrix = [
                    list(row[:20])
                    for row in worksheet.iter_rows(
                        min_row=1,
                        max_row=min(worksheet.max_row or 1, 500),
                        values_only=True,
                    )
                ]
                rows.extend(_rows_from_matrix(matrix))
                if len(rows) >= MAX_SUPPLIER_PRICE_ITEMS:
                    break
        finally:
            workbook.close()
        return rows[:MAX_SUPPLIER_PRICE_ITEMS]
    return []


def save_supplier_price_submission(
    submission: dict[str, Any],
    *,
    submitted_by: str,
    vault_path: str | Path,
    nas_path: str | Path,
    db_path: str | Path,
    attachment_name: str = "",
    attachment_bytes: bytes = b"",
) -> dict[str, Any]:
    """Archive employee-entered supplier prices and index them immediately.

    Args:
        submission: Supplier metadata and optional manual item rows.
        submitted_by: Employee display name or stable sender identifier.
        vault_path: Obsidian vault root used for the supplier note.
        nas_path: Mounted NAS root used as the authoritative source archive.
        db_path: NAS memory SQLite index updated for immediate lookup.
        attachment_name: Optional original quote filename kept as evidence.
        attachment_bytes: Optional original quote bytes.

    Returns:
        Saved relative note path, item count, review state, and evidence name.

    Raises:
        ValueError: If supplier fields, rows, units, prices, or attachment fail
            validation.
    """
    supplier_name = re.sub(
        r"\s+", " ", str(submission.get("supplier_name") or "")
    ).strip()[:80]
    if len(supplier_name) < 2:
        raise ValueError("请填写合作公司名称")
    category = str(submission.get("category") or "其他类").strip()
    if category not in ALLOWED_CATEGORIES:
        category = "其他类"
    quote_date = str(submission.get("quote_date") or date.today().isoformat()).strip()
    try:
        quote_date = date.fromisoformat(quote_date).isoformat()
    except ValueError as exc:
        raise ValueError("报价日期格式无效") from exc
    notes = str(submission.get("notes") or "").strip()[:800]
    raw_items = submission.get("items") or []
    if not isinstance(raw_items, list):
        raise ValueError("合作公司价格明细格式无效")
    raw_items = [
        item
        for item in raw_items
        if isinstance(item, dict)
        and any(
            str(item.get(key) or "").strip()
            for key in ("item_name", "unit", "unit_price")
        )
    ]
    if not raw_items and attachment_name and attachment_bytes:
        raw_items = parse_supplier_price_attachment(
            attachment_name,
            attachment_bytes,
        )
    if not raw_items:
        raise ValueError("未识别出价格，请在下方至少录入一项")
    if len(raw_items) > MAX_SUPPLIER_PRICE_ITEMS:
        raise ValueError(f"一次最多录入 {MAX_SUPPLIER_PRICE_ITEMS} 项价格")

    items: list[dict[str, str]] = []
    for index, raw_item in enumerate(raw_items, start=1):
        item_name = str(raw_item.get("item_name") or "").strip()[:80]
        unit = str(raw_item.get("unit") or "").strip()[:16]
        if not item_name:
            raise ValueError(f"请填写第 {index} 项物料名称")
        if not unit or unit in {"待确认", "待定", "未知", "-"}:
            raise ValueError(f"请填写第 {index} 项真实计价单位")
        try:
            unit_price = Decimal(str(raw_item.get("unit_price") or ""))
        except InvalidOperation as exc:
            raise ValueError(f"第 {index} 项单价无效") from exc
        if unit_price <= 0 or unit_price > Decimal("100000000"):
            raise ValueError(f"第 {index} 项单价必须大于 0")
        items.append(
            {
                "item_name": item_name,
                "specification": str(raw_item.get("specification") or "").strip()[:160],
                "unit": unit,
                "unit_price": str(unit_price.quantize(MONEY, ROUND_HALF_UP)),
            }
        )

    if len(attachment_bytes) > MAX_EVIDENCE_BYTES:
        raise ValueError("原始报价资料不能超过 20 MB")
    nas_root = Path(nas_path)
    if not nas_root.is_dir():
        raise ValueError("NAS 当前未挂载，价格未保存，请稍后重试")
    supplier_part = re.sub(r"[\\/:*?\"<>|]+", "_", supplier_name).strip(" .")[:60]
    category_part = re.sub(r"[\\/:*?\"<>|]+", "_", category).strip(" .")[:30]
    submission_id = uuid.uuid4().hex[:8]
    relative_dir = Path("20_Operations") / "供应商库" / "员工录入" / category_part
    note_filename = f"{supplier_part}_报价_{quote_date}_{submission_id}.md"
    relative_path = relative_dir / note_filename
    note_path = Path(vault_path) / relative_path
    note_path.parent.mkdir(parents=True, exist_ok=True)
    nas_relative_dir = (
        Path("projects")
        / "供应商价格库"
        / category_part
        / supplier_part
        / f"{quote_date}_{submission_id}"
    )
    nas_archive_dir = nas_root / nas_relative_dir
    nas_archive_dir.mkdir(parents=True, exist_ok=True)
    nas_note_path = nas_archive_dir / "报价记录.md"

    saved_attachment_name = ""
    if attachment_name and attachment_bytes:
        original_name = Path(attachment_name).name
        stem = re.sub(r"[\\/:*?\"<>|]+", "_", Path(original_name).stem).strip(" .")[:80]
        suffix = Path(original_name).suffix.lower()[:10]
        saved_attachment_name = f"{submission_id}_{stem or '原始报价'}{suffix}"
        (nas_archive_dir / saved_attachment_name).write_bytes(attachment_bytes)

    quoted_title = json.dumps(f"{supplier_name}报价", ensure_ascii=False)
    quoted_supplier = json.dumps(supplier_name, ensure_ascii=False)
    quoted_submitter = json.dumps(
        str(submitted_by or "飞书员工")[:120], ensure_ascii=False
    )
    lines = [
        "---",
        f"title: {quoted_title}",
        "tags: [供应商, 员工录入, 历史报价]",
        f"supplier: {quoted_supplier}",
        f"quote_date: {quote_date}",
        f"category: {category}",
        f"submitted_by: {quoted_submitter}",
        "review_status: need_review",
        f"nas_archive: {json.dumps(nas_relative_dir.as_posix(), ensure_ascii=False)}",
        "---",
        "",
        f"# {supplier_name}",
        "",
        "> [!warning] 员工录入 · 待复核",
        "> 本页用于内部历史估价，正式采购前仍需向合作公司复核。",
        "",
        "| 物料 | 规格/说明 | 单位 | 单价（元） |",
        "| --- | --- | --- | ---: |",
    ]
    for item in items:
        cells = [
            str(item[key]).replace("|", "\\|").replace("\n", " ")
            for key in ("item_name", "specification", "unit", "unit_price")
        ]
        lines.append(f"| {' | '.join(cells)} |")
    if notes:
        lines.extend(["", "## 录入备注", "", notes])
    if saved_attachment_name:
        lines.extend(
            [
                "",
                "## 原始报价资料",
                "",
                f"- NAS 归档：{nas_relative_dir.as_posix()}/{saved_attachment_name}",
            ]
        )
    lines.extend(
        ["", f"录入时间：{datetime.now(timezone.utc).isoformat(timespec='seconds')}"]
    )
    note_content = "\n".join(lines) + "\n"
    nas_note_path.write_text(note_content, encoding="utf-8")
    note_path.write_text(note_content, encoding="utf-8")

    source_hash = hashlib.sha256(note_content.encode("utf-8")).hexdigest()
    document_key = hashlib.sha1(
        f"{relative_path.as_posix()}:{source_hash}".encode()
    ).hexdigest()
    indexed_at = (
        datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")
    )
    metadata = {
        "source": "employee_supplier_price_intake",
        "supplier": supplier_name,
        "category": category,
        "quote_date": quote_date,
        "submitted_by": str(submitted_by or "飞书员工")[:120],
        "attachment": saved_attachment_name,
        "item_count": len(items),
        "nas_archive_path": nas_relative_dir.as_posix(),
    }
    chunks = chunk_text(note_content)
    database = Path(db_path)
    database.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(database) as connection:
        init_nas_db(connection)
        connection.execute(
            """
            INSERT OR REPLACE INTO documents (
                doc_key, rel_path, source_path, archive_path, sha256, file_size,
                parser, project_id, project_name, doc_type, initiator, owner,
                departments_json, participants_json, confidence, review_status,
                title, summary, tags_json, indexed_at, metadata_json
            ) VALUES (?, ?, ?, ?, ?, ?, 'md', 'supplier-prices', '供应商价格库',
                '供应商报价', ?, ?, '[]', '[]', 0.9, 'need_review', ?, ?, ?, ?, ?)
            """,
            (
                document_key,
                relative_path.as_posix(),
                str(note_path),
                str(nas_note_path),
                source_hash,
                len(note_content.encode("utf-8")),
                str(submitted_by or "飞书员工")[:120],
                str(submitted_by or "飞书员工")[:120],
                supplier_name,
                f"{supplier_name} {quote_date} 共 {len(items)} 项员工录入报价",
                json.dumps(["供应商", "员工录入", "历史报价"], ensure_ascii=False),
                indexed_at,
                json.dumps(metadata, ensure_ascii=False, sort_keys=True),
            ),
        )
        connection.execute("DELETE FROM chunks_fts WHERE doc_key = ?", (document_key,))
        connection.execute("DELETE FROM chunks WHERE doc_key = ?", (document_key,))
        for index, chunk in enumerate(chunks):
            chunk_id = f"{document_key}:{index:04d}"
            connection.execute(
                "INSERT INTO chunks (chunk_id, doc_key, chunk_index, text) VALUES (?, ?, ?, ?)",
                (chunk_id, document_key, index, chunk),
            )
            connection.execute(
                "INSERT INTO chunks_fts (chunk_id, doc_key, title, text) VALUES (?, ?, ?, ?)",
                (chunk_id, document_key, supplier_name, chunk),
            )
        review_id = hashlib.sha1(f"{document_key}:supplier-price".encode()).hexdigest()
        connection.execute(
            """
            INSERT OR REPLACE INTO review_queue (
                review_id, doc_key, reason, status, payload_json, created_at
            ) VALUES (?, ?, ?, 'open', ?, ?)
            """,
            (
                review_id,
                document_key,
                "员工新增合作公司报价待复核",
                json.dumps(metadata, ensure_ascii=False, sort_keys=True),
                indexed_at,
            ),
        )

    return {
        "relative_path": relative_path.as_posix(),
        "item_count": len(items),
        "review_status": "待复核",
        "attachment_filename": saved_attachment_name,
        "nas_archive_path": nas_relative_dir.as_posix(),
        "document_key": document_key,
    }


__all__ = ["parse_supplier_price_attachment", "save_supplier_price_submission"]
